"""計画書中心の画面の共通処理（期・段階の状態・完了の進め方・連絡帳・ファイル登録）"""
import csv
import io
from datetime import date

from django.db import transaction
from django.utils import timezone

from support_plans.models import SupportPlan

from .models import DOMAINS, GOAL_CATEGORY_LABELS, GOAL_FIELDS, ContactNote, Interview

# 段階（タブ）と、support_plans のステップ番号の対応
STAGES = [
    ('monitoring', 'モニタリング'),
    ('interview',  '面談記録'),
    ('draft',      '個別支援計画書【原案】'),
    ('meeting',    'スタッフ会議'),
    ('plan',       '個別支援計画書'),
]
STAGE_LABELS = dict(STAGES)
STAGE_STEP = {'interview': SupportPlan.STEP_ASSESSMENT, 'draft': SupportPlan.STEP_DRAFT,
              'meeting': SupportPlan.STEP_MEETING, 'plan': SupportPlan.STEP_CONSENT}


# ---------------------------------------------------------------- 期
def period_plans(beneficiary):
    """利用者の計画を作った順に（1期, 2期, …）"""
    return list(beneficiary.support_plans.order_by('created_at', 'pk'))


def period_number(plan):
    plans = period_plans(plan.beneficiary)
    for i, p in enumerate(plans, 1):
        if p.pk == plan.pk:
            return i
    return len(plans) + 1


def latest_plan(beneficiary):
    plans = period_plans(beneficiary)
    return plans[-1] if plans else None


def create_period(beneficiary, user):
    """次の期の計画を作る（前の期があれば引き継ぐ）"""
    prev = latest_plan(beneficiary)
    n = len(period_plans(beneficiary)) + 1
    plan = SupportPlan.objects.create(
        facility=beneficiary.facility, beneficiary=beneficiary, title=f'第{n}期 個別支援計画',
        created_by=user, manager=user if user.is_child_dev_manager else None,
        predecessor=prev if (prev is not None and not hasattr(prev, 'successor')) else None,
    )
    Interview.objects.create(plan=plan)
    return plan


def next_update_date(plan):
    """次回更新日＝支援期間の終了（面談記録、なければ原案の計画期間）"""
    if plan is None:
        return None
    iv = getattr(plan, 'interview', None)
    if iv and iv.period_end:
        return iv.period_end
    d = getattr(plan, 'draft', None)
    return d.period_end if d else None


# ---------------------------------------------------------------- 段階の状態
def _state(done, started):
    return 'done' if done else ('progress' if started else 'none')


def stage_status(plan):
    """
    各段階の状態。{'interview': {'state': 'none'|'progress'|'done', 'date': date|None, 'label': '未作成'|'作成中'|'完了'}, ...}
    plan が None（期がまだない）ならすべて none。
    """
    labels = {'none': '未作成', 'progress': '作成中', 'done': '完了'}
    out = {}
    if plan is None:
        for k, _ in STAGES:
            out[k] = {'state': 'none', 'date': None, 'label': '--' if k == 'monitoring' else labels['none']}
        return out
    iv = getattr(plan, 'interview', None)
    iv_started = bool(iv and (iv.interview_date or iv.period_start or iv.home_parent or iv.future_wishes))
    out['interview'] = {'state': _state(bool(iv and iv.is_completed), iv_started),
                        'date': iv.completed_at.date() if iv and iv.completed_at else iv.interview_date if iv else None}
    d = getattr(plan, 'draft', None)
    d_started = bool(d and (d.period_start or d.policy)) or plan.goals.exists()
    out['draft'] = {'state': _state(bool(d and d.is_completed), d_started), 'date': d.completed_at.date() if d and d.completed_at else None}
    m = getattr(plan, 'meeting', None)
    m_started = bool(m and (m.meeting_date or m.opinions))
    out['meeting'] = {'state': _state(bool(m and m.is_completed), m_started), 'date': m.completed_at.date() if m and m.completed_at else m.meeting_date if m else None}
    c = getattr(plan, 'consent', None)
    c_started = bool(c and (c.explained_date or c.service_start_date))
    out['plan'] = {'state': _state(bool(c and c.is_completed), c_started), 'date': c.completed_at.date() if c and c.completed_at else None}
    last = plan.monitoring_records.order_by('-date').first()
    out['monitoring'] = {'state': 'done' if last else 'none', 'date': last.date if last else None,
                         'label': last.date.strftime('%Y/%m/%d') if last else '--'}
    for k in ('interview', 'draft', 'meeting', 'plan'):
        out[k]['label'] = labels[out[k]['state']]
    return out


def complete_stage(plan, stage, user):
    """段階を完了にし、計画のステップを進める（標準画面の5ステップと整合させる）"""
    n = STAGE_STEP[stage]
    step = plan.get_step(n)
    step.completed_at = timezone.now()
    step.completed_by = user
    step.save(update_fields=['completed_at', 'completed_by', 'updated_at'])
    if plan.current_step <= n:
        plan.current_step = n + 1
        if plan.current_step == SupportPlan.STEP_MONITORING:
            plan.status = SupportPlan.STATUS_ACTIVE
            if plan.predecessor_id and plan.predecessor.status != SupportPlan.STATUS_CLOSED:
                plan.predecessor.status = SupportPlan.STATUS_CLOSED
                plan.predecessor.closed_at = timezone.now()
                plan.predecessor.save(update_fields=['status', 'closed_at'])
        plan.save(update_fields=['current_step', 'status', 'updated_at'])


def reopen_stage(plan, stage):
    n = STAGE_STEP[stage]
    step = plan.get_step(n)
    step.completed_at = None
    step.completed_by = None
    step.save(update_fields=['completed_at', 'completed_by', 'updated_at'])
    if plan.current_step > n:
        plan.current_step = n
        if plan.status == SupportPlan.STATUS_ACTIVE:
            plan.status = SupportPlan.STATUS_IN_PROGRESS
        plan.save(update_fields=['current_step', 'status', 'updated_at'])


def sync_interview_to_assessment(interview):
    """面談記録の内容を標準画面のアセスメント（ステップ1）にも写す（標準画面から見ても中身が揃うように）"""
    a = interview.plan.get_step(SupportPlan.STEP_ASSESSMENT)
    a.interview_date = interview.interview_date
    a.interviewed_with = a.interviewed_with or '保護者'
    a.interviewer = interview.author
    a.condition = '\n'.join(x for x in (interview.home_parent, interview.school_parent, interview.social_parent) if x)
    a.environment = '\n'.join(x for x in (interview.home_staff, interview.school_staff, interview.social_staff) if x)
    a.wishes = '\n'.join(x for x in (interview.future_wishes, interview.other_wishes) if x)
    a.save()


# ---------------------------------------------------------------- 目標（支援内容）
def goal_extra(goal):
    """目標の追加項目（form_extra）を、画面用に既定値つきで返す"""
    e = dict(goal.form_extra or {})
    for key, _, _ in GOAL_FIELDS:
        e.setdefault(key, '')
    e.setdefault('category', '')
    e.setdefault('priority', '')
    e['domains'] = [d for d in (e.get('domains') or []) if d in dict(DOMAINS)]
    e['category_label'] = GOAL_CATEGORY_LABELS.get(e['category'], '')
    e['content'] = goal.content
    e['support'] = goal.support_content
    return e


def save_goals(plan, post):
    """
    画面の目標の行をまとめて保存する。行は goal_<idx>_<field>。idx は既存の pk か new-N。
    削除された行（画面から消えた行）は削除する。
    """
    from support_plans.models import PlanGoal
    idxs = post.getlist('goal_idx')
    keep = []
    order = 0
    for idx in idxs:
        def g(field):
            return (post.get(f'goal_{idx}_{field}') or '').strip()
        content = g('content')
        support = g('support')
        extra = {
            'category': g('category') if g('category') in GOAL_CATEGORY_LABELS else '',
            'priority': g('priority')[:2],
            'staff': g('staff')[:70], 'notes': g('notes')[:210], 'target': g('target')[:210],
            'timing': g('timing')[:12], 'eval_timing': g('eval_timing')[:12],
            'domains': [d for d in post.getlist(f'goal_{idx}_domains') if d in dict(DOMAINS)],
        }
        if not any([content, support, extra['staff'], extra['notes'], extra['target']]):
            continue
        if idx.startswith('new'):
            goal = PlanGoal(plan=plan, goal_type=PlanGoal.TYPE_SHORT)
        else:
            goal = plan.goals.filter(pk=idx).first()
            if goal is None:
                continue
        goal.content = content[:200] or '（未入力）'
        goal.support_content = support[:210]
        goal.order = order
        goal.form_extra = {**(goal.form_extra or {}), **extra}
        goal.save()
        keep.append(goal.pk)
        order += 1
    plan.goals.exclude(pk__in=keep).delete()
    return len(keep)


# ---------------------------------------------------------------- 連絡帳
def post_staff_note(facility, guardian, author, body):
    """職員の記入を連絡帳に積み、LINE 連携があれば送る"""
    note = ContactNote.objects.create(facility=facility, guardian=guardian, sender=ContactNote.FROM_STAFF,
                                      author=author, body=body)
    if facility.use_line and guardian.line_linked and guardian.line_user_id:
        from line_integration.sending import push_text
        ok, error = push_text(facility, guardian.line_user_id, body)
        note.line_sent = ok
        note.line_error = (error or '')[:200]
        note.save(update_fields=['line_sent', 'line_error'])
    return note


def record_guardian_line(facility, guardian, text):
    """保護者からの LINE を連絡帳に積む（未読）"""
    return ContactNote.objects.create(facility=facility, guardian=guardian, sender=ContactNote.FROM_GUARDIAN,
                                      body=text, is_read=False)


def note_threads(facility):
    """保護者ごとの連絡帳（最新の1件と未読数）。やりとりがある保護者だけ"""
    from beneficiaries.models import Guardian
    rows = []
    guardians = Guardian.objects.filter(beneficiary__facility=facility).select_related('beneficiary')
    notes = ContactNote.objects.filter(facility=facility).order_by('-created_at')
    latest = {}
    unread = {}
    for n in notes:
        latest.setdefault(n.guardian_id, n)
        if n.sender == ContactNote.FROM_GUARDIAN and not n.is_read:
            unread[n.guardian_id] = unread.get(n.guardian_id, 0) + 1
    for g in guardians:
        if g.pk in latest:
            rows.append({'guardian': g, 'latest': latest[g.pk], 'unread': unread.get(g.pk, 0)})
    rows.sort(key=lambda r: r['latest'].created_at, reverse=True)
    return rows


def unread_note_count(facility):
    return ContactNote.objects.filter(facility=facility, sender=ContactNote.FROM_GUARDIAN, is_read=False).count()


# ---------------------------------------------------------------- ファイルから登録
IMPORT_COLUMNS = [
    ('last_name', '姓', True), ('first_name', '名', True),
    ('last_name_kana', 'せい（ふりがな）', False), ('first_name_kana', 'めい（ふりがな）', False),
    ('date_of_birth', '生年月日', True), ('gender', '性別', False),
    ('certificate_number', '受給者番号', False), ('certificate_expiry', '受給者証期限', False),
    ('admission_date', '入所日', False), ('postal_code', '郵便番号', False), ('address', '住所', False),
    ('mobile_phone', '携帯電話番号', False), ('home_phone', '自宅電話番号', False),
    ('school_name', '通学学校名', False), ('grade', '学年', False),
    ('guardian_last_name', '保護者 姓', False), ('guardian_first_name', '保護者 名', False),
    ('guardian_relation', '保護者 続柄', False), ('guardian_phone', '保護者 電話番号', False), ('guardian_email', '保護者 メール', False),
]
IMPORT_HEADERS = [label for _, label, _ in IMPORT_COLUMNS]


def import_template_csv():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(IMPORT_HEADERS)
    w.writerow(['山田', '太郎', 'やまだ', 'たろう', '2017-04-01', '男', '2600001234', '2027-03-31', '2026-04-01',
                '6008216', '京都市下京区○○町1-1', '090-0000-0000', '', '○○小学校', '小3', '山田', '花子', '母', '090-0000-0001', ''])
    return buf.getvalue()


def _parse_date(v):
    v = (v or '').strip().replace('/', '-').replace('.', '-')
    if not v:
        return None
    for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
        try:
            from datetime import datetime
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    m = __import__('re').match(r'^(\d{4})年(\d{1,2})月(\d{1,2})日$', v)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def read_import_rows(uploaded):
    """CSV（UTF-8／Shift_JIS）か Excel（.xlsx）を読み、見出し→値の dict のリストにする"""
    name = (uploaded.name or '').lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(uploaded, read_only=True, data_only=True)
        ws = wb.active
        rows = [[('' if c is None else (c.strftime('%Y-%m-%d') if hasattr(c, 'strftime') else str(c))) for c in r]
                for r in ws.iter_rows(values_only=True)]
    else:
        raw = uploaded.read()
        text = None
        for enc in ('utf-8-sig', 'cp932', 'utf-8'):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError('文字コードを読み取れません（UTF-8 か Shift_JIS の CSV にしてください）')
        rows = list(csv.reader(io.StringIO(text)))
    rows = [r for r in rows if any((c or '').strip() for c in r)]
    if not rows:
        return []
    header = [(h or '').strip() for h in rows[0]]
    label_to_key = {label: key for key, label, _ in IMPORT_COLUMNS}
    keys = [label_to_key.get(h, h) for h in header]
    out = []
    for r in rows[1:]:
        d = {}
        for i, k in enumerate(keys):
            if k:
                d[k] = (r[i] if i < len(r) else '').strip()
        out.append(d)
    return out


GRADE_BY_LABEL = {}


def _grade_key(label):
    from beneficiaries.models import Beneficiary
    if not GRADE_BY_LABEL:
        GRADE_BY_LABEL.update({l: k for k, l in Beneficiary.GRADE_CHOICES if k})
    return GRADE_BY_LABEL.get((label or '').strip(), '')


@transaction.atomic
def import_beneficiaries(facility, rows):
    """行を利用者として登録する。戻り値 (登録数, エラーの一覧[(行番号, 理由)])"""
    from beneficiaries.models import Beneficiary, Guardian, RecipientCertificate
    created = 0
    errors = []
    gender_map = {'男': 'male', '女': 'female', 'male': 'male', 'female': 'female', 'その他': 'other'}
    relation_map = {'父': 'father', '母': 'mother', 'father': 'father', 'mother': 'mother'}
    for i, r in enumerate(rows, 2):
        last, first = r.get('last_name', ''), r.get('first_name', '')
        dob = _parse_date(r.get('date_of_birth'))
        if not (last and first):
            errors.append((i, '姓・名が空です'))
            continue
        if dob is None:
            errors.append((i, '生年月日を読み取れません（例：2017-04-01）'))
            continue
        b = Beneficiary.objects.create(
            facility=facility, last_name=last[:50], first_name=first[:50],
            last_name_kana=r.get('last_name_kana', '')[:50], first_name_kana=r.get('first_name_kana', '')[:50],
            date_of_birth=dob, gender=gender_map.get(r.get('gender', ''), 'male'),
            admission_date=_parse_date(r.get('admission_date')),
            postal_code=r.get('postal_code', '')[:8], address=r.get('address', '')[:200],
            mobile_phone=r.get('mobile_phone', '')[:20], home_phone=r.get('home_phone', '')[:20],
            school_name=r.get('school_name', '')[:100], grade=_grade_key(r.get('grade')),
            has_prior_records=True,
        )
        cert_no, expiry = r.get('certificate_number', ''), _parse_date(r.get('certificate_expiry'))
        if cert_no or expiry:
            RecipientCertificate.objects.create(
                beneficiary=b, certificate_number=cert_no[:20] or '未入力', granted_days=0, monthly_cap=0,
                valid_from=date.today(), valid_until=expiry or date.today(),
            )
        gl, gf = r.get('guardian_last_name', ''), r.get('guardian_first_name', '')
        if gl or gf:
            Guardian.objects.create(
                beneficiary=b, last_name=gl[:50] or last, first_name=gf[:50] or '保護者',
                relation=relation_map.get(r.get('guardian_relation', ''), 'other'),
                phone=r.get('guardian_phone', '')[:20], email=r.get('guardian_email', '')[:254], is_primary=True,
            )
        created += 1
    return created, errors
