"""時間枠の予約・月予約利用希望・月間予定表（発達支援ルーム　ゆあーず）のテスト"""
import datetime
from unittest import mock

from django.test import TestCase, override_settings
from django.urls import reverse

from accounts.models import StaffAccount
from beneficiaries.models import Beneficiary
from config.jp_holidays import holidays, is_weekend_or_holiday
from facilities.models import Facility

from . import monthly, services
from .models import Customer, DayStaff, MonthlyRequest, RequestScan, Reservation, ReservationNotice


def child(facility, last, first='子'):
    return Beneficiary.objects.create(facility=facility, last_name=last, first_name=first,
                                      last_name_kana=last, date_of_birth=datetime.date(2019, 4, 1))


def ryoiku():
    f = Facility.objects.create(name='発達支援ルーム　ゆあーず', use_reservation=True, use_therapy_record=True,
                                layout=Facility.LAYOUT_RYOIKU)
    s = services.get_setting(f)
    s.slot_mode, s.slot_capacity = True, 3
    s.weekday_first_hour, s.weekday_last_hour = 10, 18
    s.holiday_first_hour, s.holiday_last_hour = 9, 17
    s.break_hours, s.closed_weekdays = [12], [0, 3]
    s.save()
    return f, s


class HolidayTests(TestCase):
    def test_known_holidays(self):
        h = holidays(2026)
        self.assertEqual(h[datetime.date(2026, 1, 1)], '元日')
        self.assertEqual(h[datetime.date(2026, 1, 12)], '成人の日')      # 第2月曜
        self.assertEqual(h[datetime.date(2026, 3, 20)], '春分の日')
        self.assertEqual(h[datetime.date(2026, 9, 21)], '敬老の日')
        self.assertEqual(h[datetime.date(2026, 9, 23)], '秋分の日')
        self.assertEqual(h[datetime.date(2026, 9, 22)], '国民の休日')     # 敬老の日と秋分の日の間
        self.assertEqual(h[datetime.date(2026, 5, 6)], '振替休日')        # 5/3 が日曜
        self.assertTrue(is_weekend_or_holiday(datetime.date(2026, 10, 3)))    # 土
        self.assertTrue(is_weekend_or_holiday(datetime.date(2026, 11, 3)))    # 文化の日（火）
        self.assertFalse(is_weekend_or_holiday(datetime.date(2026, 10, 2)))   # 金


class SlotSettingTests(TestCase):
    def setUp(self):
        self.f, self.s = ryoiku()

    def test_slot_hours_by_day_type(self):
        self.assertEqual(self.s.slot_hours(datetime.date(2026, 10, 2)), [10, 11, 13, 14, 15, 16, 17, 18])   # 金
        self.assertEqual(self.s.slot_hours(datetime.date(2026, 10, 3)), [9, 10, 11, 13, 14, 15, 16, 17])    # 土
        self.assertEqual(self.s.slot_hours(datetime.date(2026, 11, 3)), [9, 10, 11, 13, 14, 15, 16, 17])    # 祝（火）
        self.assertEqual(self.s.slot_hours(datetime.date(2026, 10, 5)), [])     # 月曜はお休み
        self.assertEqual(self.s.all_slot_hours(), [9, 10, 11, 13, 14, 15, 16, 17, 18])
        self.assertEqual(self.s.slot_capacity_of(datetime.date(2026, 10, 2)), 24)
        self.assertIn('平日 10:00〜18:00枠', self.s.hours_text)
        self.assertEqual(self.s.closed_weekdays_text, '月曜日・木曜日')


class SlotReservationTests(TestCase):
    def setUp(self):
        self.f, self.s = ryoiku()
        self.day = datetime.date(2026, 10, 2)    # 金
        self.kids = [child(self.f, n) for n in ('青木', '井上', '上田', '江口')]

    def test_requires_time_and_counts_per_slot(self):
        with self.assertRaises(services.ReservationError):
            services.create_reservation(self.f, self.kids[0], self.day)
        with self.assertRaises(services.ReservationError):
            services.create_reservation(self.f, self.kids[0], self.day, start_time='9')    # 平日に9時の枠は無い
        for k in self.kids[:3]:
            res, notice = services.create_reservation(self.f, k, self.day, start_time='10')
            self.assertEqual((res.status, res.start_time), (Reservation.STATUS_CONFIRMED, datetime.time(10, 0)))
        self.assertIn('10:00', notice.body)
        st = services.day_state(self.f, self.day)
        slot10 = next(x for x in st['slots'] if x['hour'] == 10)
        self.assertEqual((slot10['confirmed'], slot10['full'], st['confirmed'], st['capacity']), (3, True, 3, 24))
        # 4人目は同じ枠ではキャンセル待ち、別の枠なら確定
        res4, _ = services.create_reservation(self.f, self.kids[3], self.day, start_time='10:00')
        self.assertEqual(res4.status, Reservation.STATUS_WAITLIST)
        services.cancel_reservation(res4, notify=False)
        res4b, _ = services.create_reservation(self.f, self.kids[3], self.day, start_time=11)
        self.assertEqual(res4b.status, Reservation.STATUS_CONFIRMED)

    def test_waitlist_promotes_within_same_slot(self):
        rs = [services.create_reservation(self.f, k, self.day, start_time=10)[0] for k in self.kids[:3]]
        waiting, _ = services.create_reservation(self.f, self.kids[3], self.day, start_time=10)
        self.assertEqual(waiting.status, Reservation.STATUS_WAITLIST)
        # 別の枠が空いても繰り上げない（同じ枠が空いたとき）
        promoted = services.cancel_reservation(rs[0])
        self.assertEqual([p.pk for p in promoted], [waiting.pk])

    def test_move_to_other_slot_same_day(self):
        res, _ = services.create_reservation(self.f, self.kids[0], self.day, start_time=10)
        services.move_reservation(res, self.day, start_time='14')
        res.refresh_from_db()
        self.assertEqual((res.date, res.start_time, res.status), (self.day, datetime.time(14, 0), 'confirmed'))
        # 満枠の枠には移せない → キャンセル待ち
        for k in self.kids[1:4]:
            services.create_reservation(self.f, k, self.day, start_time=15)
        services.move_reservation(res, self.day, start_time=15)
        res.refresh_from_db()
        self.assertEqual(res.status, Reservation.STATUS_WAITLIST)

    def test_closed_weekday_has_no_capacity(self):
        monday = datetime.date(2026, 10, 5)
        self.assertEqual(services.day_state(self.f, monday)['capacity'], 0)
        with self.assertRaises(services.ReservationError):
            services.create_reservation(self.f, self.kids[0], monday, start_time=10)


class MonthlyRequestTests(TestCase):
    def setUp(self):
        self.f, self.s = ryoiku()
        self.kids = [child(self.f, n) for n in ('青木', '井上', '上田', '江口', '大野')]
        self.customer = Customer.objects.create(facility=self.f, name='青木 母', line_user_id='U-aoki')
        self.customer.children.add(self.kids[0])

    def test_wish_helpers(self):
        req = monthly.save_request(self.f, self.kids[0], 2026, 10, 4,
                                   {'2026-10-02': 'all', '2026-10-03': [9, 10, 12], '2026-10-05': 'all'})
        self.assertEqual(req.wish_hours(datetime.date(2026, 10, 2), self.s), [10, 11, 13, 14, 15, 16, 17, 18])
        self.assertEqual(req.wish_hours(datetime.date(2026, 10, 3), self.s), [9, 10])   # 12時は枠なし
        self.assertEqual(req.wish_hours(datetime.date(2026, 10, 5), self.s), [])        # 月曜はお休み
        self.assertEqual(req.slot_count(self.s), 10)
        grid = monthly.request_grid(self.f, 2026, 10, self.s, request=req)
        self.assertEqual(grid['hours'], [9, 10, 11, 13, 14, 15, 16, 17, 18])
        row2 = next(r for r in grid['rows'] if r['date'] == datetime.date(2026, 10, 2))
        self.assertTrue(row2['all_wished'])
        self.assertFalse(row2['cells'][0]['active'])          # 平日の9時は網掛け
        row5 = next(r for r in grid['rows'] if r['date'] == datetime.date(2026, 10, 5))
        self.assertTrue(row5['closed'])

    def test_assign_respects_capacity_and_spreads(self):
        # 5人全員が「土曜の 9・10 時」だけ希望、各3回 → 土曜は 10/3,10,17,24,31 の5日 × 2枠 × 3人
        saturdays = ['2026-10-03', '2026-10-10', '2026-10-17', '2026-10-24', '2026-10-31']
        for k in self.kids:
            monthly.save_request(self.f, k, 2026, 10, 3, {d: [9, 10] for d in saturdays})
        result = monthly.assign_month(self.f, 2026, 10, self.s)
        self.assertEqual(len(result.made), 15)
        self.assertEqual(result.short, [])
        for k in self.kids:
            mine = Reservation.objects.filter(beneficiary=k, status='confirmed')
            self.assertEqual(mine.count(), 3)
            self.assertEqual(len({r.date for r in mine}), 3)          # 同じ日に2回は入れない
        # どの枠も3人まで
        from collections import Counter
        c = Counter((r.date, r.start_time) for r in Reservation.objects.filter(status='confirmed'))
        self.assertTrue(all(n <= 3 for n in c.values()))
        # 通知は利用者（の連絡先）ごとに1通、対象日をまとめて
        notes = ReservationNotice.objects.filter(customer=self.customer)
        self.assertEqual(notes.count(), 1)
        self.assertIn('ご利用日が決まりました', notes.first().body)
        # もう一度回しても増えない（希望回数に達している）
        again = monthly.assign_month(self.f, 2026, 10, self.s)
        self.assertEqual(len(again.made), 0)

    def test_assign_reports_shortage(self):
        # 4人が同じ1枠だけ希望 → 3人しか入らない
        for k in self.kids[:4]:
            monthly.save_request(self.f, k, 2026, 10, 1, {'2026-10-03': [9]})
        result = monthly.assign_month(self.f, 2026, 10, self.s)
        self.assertEqual(len(result.made), 3)
        self.assertEqual(len(result.short), 1)
        self.assertIn('あと1回', result.summary)

    def test_month_schedule_grid(self):
        monthly.save_request(self.f, self.kids[0], 2026, 10, 1, {'2026-10-03': [9]})
        monthly.assign_month(self.f, 2026, 10, self.s)
        sched = monthly.month_schedule(self.f, 2026, 10, self.s)
        self.assertEqual(sched['weeks'][0]['days'][0]['date'].weekday(), 6)      # 日曜はじまり
        sat = next(d for w in sched['weeks'] for d in w['days'] if d['date'] == datetime.date(2026, 10, 3))
        slot9 = sat['slots'][0]
        self.assertEqual((slot9['hour'], slot9['active'], len(slot9['boxes'])), (9, True, 3))
        self.assertEqual(slot9['boxes'][0].beneficiary, self.kids[0])
        self.assertIsNone(slot9['boxes'][1])
        mon = next(d for w in sched['weeks'] for d in w['days'] if d['date'] == datetime.date(2026, 10, 5))
        self.assertTrue(mon['closed'])
        rows = monthly.request_rows(self.f, 2026, 10, self.s)
        aoki = next(r for r in rows if r['beneficiary'] == self.kids[0])
        self.assertEqual((aoki['desired'], aoki['confirmed'], aoki['remaining']), (1, 1, 0))


    def test_all_day_wishes_fill_from_morning(self):
        # 4人が 10/3（土）を「終日」で1回 → 9時に3人、あふれた1人は次に早い10時
        for k in self.kids[:4]:
            monthly.save_request(self.f, k, 2026, 10, 1, {'2026-10-03': 'all'})
        monthly.assign_month(self.f, 2026, 10, self.s)
        hours = sorted(r.start_time.hour for r in Reservation.objects.filter(status='confirmed'))
        self.assertEqual(hours, [9, 9, 9, 10])

    def test_time_specified_wishes_go_before_all_day(self):
        # 井上・上田・江口は 9時だけ希望（ほかの日に1回確定ずみ）、青木は終日。
        # 回数の少ない青木が先に回っても、9時の枠を先に取らない
        sat, sun = datetime.date(2026, 10, 3), datetime.date(2026, 10, 18)
        for k in self.kids[1:4]:
            services.create_reservation(self.f, k, sun, start_time=datetime.time(9, 0), notify=False)
            monthly.save_request(self.f, k, 2026, 10, 2, {'2026-10-03': [9], '2026-10-18': [9]})
        monthly.save_request(self.f, self.kids[0], 2026, 10, 1, {'2026-10-03': 'all'})
        result = monthly.assign_month(self.f, 2026, 10, self.s)
        self.assertEqual(result.short, [])
        at9 = set(Reservation.objects.filter(date=sat, start_time=datetime.time(9, 0)).values_list('beneficiary', flat=True))
        self.assertEqual(at9, {k.pk for k in self.kids[1:4]})
        self.assertEqual(Reservation.objects.get(beneficiary=self.kids[0]).start_time, datetime.time(10, 0))

    def test_time_specified_wish_never_placed_outside_marked_hours(self):
        monthly.save_request(self.f, self.kids[0], 2026, 10, 1, {'2026-10-03': [14]})
        monthly.assign_month(self.f, 2026, 10, self.s)
        self.assertEqual(Reservation.objects.get(beneficiary=self.kids[0]).start_time, datetime.time(14, 0))

    def _res(self, kid, day, hour):
        res, _ = services.create_reservation(self.f, kid, day, start_time=datetime.time(hour, 0), notify=False)
        return res

    def test_swap_two_children(self):
        d1, d2 = datetime.date(2026, 10, 3), datetime.date(2026, 10, 10)
        a, b = self._res(self.kids[0], d1, 9), self._res(self.kids[1], d2, 14)
        before = ReservationNotice.objects.count()
        monthly.swap_reservations(a, b)
        a.refresh_from_db(), b.refresh_from_db()
        self.assertEqual((a.date, a.start_time.hour), (d2, 14))
        self.assertEqual((b.date, b.start_time.hour), (d1, 9))
        self.assertEqual(ReservationNotice.objects.count(), before)      # お知らせは積まない
        # 同じ日の中での入れ替え（時刻だけ交換）
        c = self._res(self.kids[2], d1, 10)
        monthly.swap_reservations(b, c)
        b.refresh_from_db(), c.refresh_from_db()
        self.assertEqual((b.start_time.hour, c.start_time.hour), (10, 9))

    def test_swap_refuses_double_booking_and_waitlist(self):
        d1, d2 = datetime.date(2026, 10, 3), datetime.date(2026, 10, 10)
        a, b = self._res(self.kids[0], d1, 9), self._res(self.kids[1], d2, 9)
        self._res(self.kids[0], d2, 14)                     # 青木は 10/10 にも予約がある
        with self.assertRaises(services.ReservationError):
            monthly.swap_reservations(a, b)
        a.refresh_from_db()
        self.assertEqual(a.date, d1)
        b.status = Reservation.STATUS_WAITLIST
        b.save()
        with self.assertRaises(services.ReservationError):
            monthly.swap_reservations(a, b)

    def test_move_on_schedule(self):
        d1, d2 = datetime.date(2026, 10, 3), datetime.date(2026, 10, 10)
        a = self._res(self.kids[0], d1, 9)
        for k in self.kids[1:4]:
            self._res(k, d2, 9)
        with self.assertRaises(services.ReservationError):   # 満員の枠へは移さない
            monthly.move_on_schedule(a, d2, 9)
        with self.assertRaises(services.ReservationError):   # 枠の無い時刻（休憩）
            monthly.move_on_schedule(a, d2, 12)
        before = ReservationNotice.objects.count()
        a = monthly.move_on_schedule(a, d2, 10)
        self.assertEqual((a.date, a.start_time.hour, a.status), (d2, 10, 'confirmed'))
        self.assertEqual(ReservationNotice.objects.count(), before)

    def test_swap_rewrites_pending_month_notice(self):
        monthly.save_request(self.f, self.kids[0], 2026, 10, 1, {'2026-10-03': [9]})
        monthly.assign_month(self.f, 2026, 10, self.s)
        notice = ReservationNotice.objects.get(customer=self.customer)
        self.assertIn('10月3日(土) 9:00', notice.body)
        a = Reservation.objects.get(beneficiary=self.kids[0])
        b = self._res(self.kids[1], datetime.date(2026, 10, 10), 14)
        monthly.swap_reservations(a, b)
        monthly.refresh_month_notice(self.f, self.kids[0], 2026, 10, self.s)
        notice.refresh_from_db()
        self.assertIn('10月10日(土) 14:00', notice.body)
        self.assertNotIn('10月3日', notice.body)
        # 送信ずみなら書き直さない
        notice.status = ReservationNotice.STATUS_SENT
        notice.save()
        monthly.swap_reservations(a, b)
        self.assertIsNone(monthly.refresh_month_notice(self.f, self.kids[0], 2026, 10, self.s))


class NgDaysTests(TestCase):
    """「来られない日」だけを書いた用紙：それ以外の日は終日可能として扱う"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.kid = child(self.f, '青木')
        self.other = child(self.f, '井上')

    def test_model_helpers(self):
        req = monthly.save_request(self.f, self.kid, 2026, 10, 2, {'2026-10-02': 'all'},
                                   wish_mode=MonthlyRequest.WISH_NG, ng_dates=['2026-10-04', '2026-10-11', '2026-10-04', 'bad'])
        self.assertTrue(req.is_ng_mode)
        self.assertEqual(req.wishes, {'2026-10-02': 'all'})              # 来られる日の○は残す
        self.assertEqual(req.ng_days(), [datetime.date(2026, 10, 4), datetime.date(2026, 10, 11)])
        self.assertEqual(req.wish_of(datetime.date(2026, 10, 2)), 'all')
        self.assertEqual(req.wish_of(datetime.date(2026, 10, 3)), 'any')  # 印の無い日はどの枠でも可
        self.assertIsNone(req.wish_of(datetime.date(2026, 10, 4)))
        self.assertEqual(req.marked_days(), [datetime.date(2026, 10, 2)])
        self.assertEqual(req.slot_count(self.s), len(self.s.slot_hours(datetime.date(2026, 10, 2))))
        self.assertIn('○', monthly.wish_summary(req, self.s))
        # 来られない日に付いた○は捨てる
        req2 = monthly.save_request(self.f, self.other, 2026, 10, 2, {'2026-10-04': [10], '2026-10-06': [10]},
                                    wish_mode=MonthlyRequest.WISH_NG, ng_dates=['2026-10-04'])
        self.assertEqual(req2.wishes, {'2026-10-06': [10]})
        self.assertIsNone(req.wish_of(datetime.date(2026, 11, 3)))        # ほかの月
        self.assertEqual(len(req.wished_days()), 31 - 2)
        self.assertEqual(req.wish_hours(datetime.date(2026, 10, 5), self.s), [])   # 月曜はお休み
        self.assertEqual(req.wish_hours(datetime.date(2026, 10, 3), self.s), [9, 10, 11, 13, 14, 15, 16, 17])
        grid = monthly.request_grid(self.f, 2026, 10, self.s, request=req)
        self.assertTrue(grid['ng_mode'])
        self.assertTrue(next(r for r in grid['rows'] if r['date'].day == 4)['ng'])
        self.assertFalse(next(r for r in grid['rows'] if r['date'].day == 3)['ng'])
        self.assertIn('来られない日 2 日', monthly.wish_summary(req, self.s))
        # ふつうの書き方に戻すと来られない日は消える
        req = monthly.save_request(self.f, self.kid, 2026, 10, 2, {'2026-10-02': 'all'})
        self.assertFalse(req.is_ng_mode)
        self.assertEqual((req.ng_dates, req.wishes), ([], {'2026-10-02': 'all'}))

    def test_assign_avoids_ng_days(self):
        ng = ['2026-10-02', '2026-10-03', '2026-10-04', '2026-10-06', '2026-10-07']   # 1・5・8 は休業日（木・月・木）
        req = monthly.save_request(self.f, self.kid, 2026, 10, 2, {}, wish_mode='ng', ng_dates=ng)
        result = monthly.assign_month(self.f, 2026, 10, self.s, notify=False)
        made = sorted(r.date for r in result.made)
        self.assertEqual(len(made), 2)
        for d in made:
            self.assertNotIn(d.isoformat(), ng)
            self.assertTrue(self.s.slot_hours(d))
        self.assertEqual(made[0], datetime.date(2026, 10, 9))              # 来られない日と休業日を飛ばした最初の日
        self.assertEqual(result.short, [])

    def test_staff_screen_saves_ng_mode(self):
        StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        url = reverse('reservations:monthly_request_edit', args=[2026, 10, self.kid.pk])
        res = self.client.get(url)
        self.assertContains(res, '来られない日を書く')
        self.assertContains(res, 'name="ng_2026-10-03"')
        res = self.client.post(url, {'desired_count': '2', 'wish_mode': 'ng', 'ng_2026-10-04': '1', 'ng_2026-10-11': '1',
                                     'h_2026-10-02': ['10']}, follow=True)
        self.assertContains(res, '来られない日 2 日・○ 1 枠')
        req = MonthlyRequest.objects.get(beneficiary=self.kid)
        self.assertEqual((req.wish_mode, req.ng_dates, req.wishes), ('ng', ['2026-10-04', '2026-10-11'], {'2026-10-02': [10]}))
        res = self.client.get(url)
        self.assertContains(res, 'id="wishModeNg" value="ng" autocomplete="off" checked')
        self.assertContains(res, 'id="ng_2026-10-04" value="1" autocomplete="off" checked')
        res = self.client.get(reverse('reservations:monthly_requests', args=[2026, 10]))
        self.assertContains(res, '来られない日 2 日・○ 1 枠')
        self.assertContains(res, 'NG児童')                                   # 来られない日のある子は NG 児童
        self.assertContains(res, '来られない日：10/4・10/11')
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(res, 'NG児童：来られない日 10/4・10/11')
        res = self.client.get(reverse('reservations:monthly_schedule_pdf', args=[2026, 10]) + '?fmt=html')
        self.assertContains(res, '<b>NG</b>')

    def test_assign_prefers_marked_hours_in_ng_mode(self):
        """来られない日の書き方でも、○のある日・時刻を先に使い、足りないぶんだけ印の無い日から取る"""
        req = monthly.save_request(self.f, self.kid, 2026, 10, 2, {'2026-10-10': [14], '2026-10-17': 'all'},
                                   wish_mode='ng', ng_dates=['2026-10-02', '2026-10-03'])
        result = monthly.assign_month(self.f, 2026, 10, self.s, notify=False)
        made = sorted((r.date, r.hour) for r in result.made)
        self.assertEqual(made, [(datetime.date(2026, 10, 10), 14), (datetime.date(2026, 10, 17), 9)])
        # 3回目は印の無い日（来られない日・休業日・すでに入った日以外）から
        req.desired_count = 3
        req.save()
        result = monthly.assign_month(self.f, 2026, 10, self.s, notify=False)
        self.assertEqual(len(result.made), 1)
        d = result.made[0].date
        self.assertNotIn(d, (datetime.date(2026, 10, 2), datetime.date(2026, 10, 3), datetime.date(2026, 10, 10), datetime.date(2026, 10, 17)))
        self.assertTrue(self.s.slot_hours(d))
        self.assertEqual(result.short, [])

    def test_customer_page_saves_ng_mode(self):
        customer = Customer.objects.create(facility=self.f, name='青木 母')
        customer.children.add(self.kid)
        url = reverse('reservations_public:customer', args=[customer.token])
        self.assertContains(self.client.get(url), '来られない日を選ぶ')
        today = datetime.date.today()
        y, m = monthly.next_month(today.year, today.month)
        day = next(d for d in monthly.month_days(y, m) if self.s.slot_hours(d))
        res = self.client.post(url, {'action': 'wish', 'beneficiary': self.kid.pk, 'wish_year': y, 'wish_month': m,
                                     'desired_count': '2', 'wish_mode': 'ng', f'ng_{day.isoformat()}': '1'}, follow=True)
        self.assertContains(res, '来られない日 1 日')
        req = MonthlyRequest.objects.get(beneficiary=self.kid, year=y, month=m)
        self.assertEqual((req.wish_mode, req.ng_dates), ('ng', [day.isoformat()]))

    def test_scan_normalize_ng_mode(self):
        from . import scan
        data = {'name': '青木 子', 'year': 2026, 'month': 10, 'desired_count': 2,
                'days': [{'day': 3, 'all_day': False, 'hours': [10]}, {'day': 4, 'all_day': True, 'hours': []}],
                'mode': 'ng', 'ng_days': [4, 11, 12, 17, 18, 24, 31, 40], 'note': '', 'unreadable': ''}
        out = scan.normalize(data, self.f, 2026, 10, self.s)
        self.assertEqual(out['wish_mode'], 'ng')
        self.assertEqual(out['ng_dates'], ['2026-10-04', '2026-10-11', '2026-10-12', '2026-10-17', '2026-10-18', '2026-10-24', '2026-10-31'])
        self.assertEqual(out['wishes'], {'2026-10-03': [10]})          # 来られる日の○は残し、来られない日（4日）の○は捨てる
        self.assertEqual(out['slot_count'], 1)
        req = scan.as_request_data(out, self.f, 2026, 10)
        self.assertTrue(req.is_ng_mode)
        self.assertEqual(req.wish_of(datetime.date(2026, 10, 3)), [10])
        self.assertEqual(req.wish_of(datetime.date(2026, 10, 6)), 'any')
        self.assertIsNone(req.wish_of(datetime.date(2026, 10, 4)))
        ok = scan.normalize(dict(READ, mode='ok', ng_days=[4]), self.f, 2026, 10, self.s)
        self.assertEqual((ok['wish_mode'], ok['ng_dates']), ('ok', []))
        self.assertIn('ng_days', scan.SCHEMA['properties'])


class DailyLogTests(TestCase):
    """業務日誌（1日ごとの予定表・A4 に4日）と「その日の担当」"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_ADMIN, display_name='大坂')
        self.client.login(username='ryo', password='pw12345678')
        self.kids = [child(self.f, n) for n in ('青木', '井上', '上田')]
        for kid, (d, h) in zip(self.kids, ((2, 11), (2, 10), (3, 9))):
            services.create_reservation(self.f, kid, datetime.date(2026, 10, d), start_time=datetime.time(h, 0), notify=False)

    def test_pages_and_rows(self):
        from .models import DayStaff
        DayStaff.objects.create(facility=self.f, date=datetime.date(2026, 10, 2), text='pm大坂')
        pages = monthly.daily_log_pages(self.f, datetime.date(2026, 10, 1), datetime.date(2026, 10, 31), self.s)
        days = [d for page in pages for d in page]
        self.assertEqual(len(pages), 6)                                   # 10月の営業日 22 日 → 4日ずつ 6 枚
        self.assertTrue(all(len(p) <= 4 for p in pages))
        self.assertEqual([d['date'].day for d in days[:4]], [2, 3, 4, 6])   # 1日(木)・5日(月) はお休みで出ない
        d2 = days[0]
        self.assertEqual((d2['weekday'], d2['staff']), ('金曜日', 'pm大坂'))
        self.assertEqual(len(d2['rows']), monthly.DAILY_LOG_ROWS)
        self.assertEqual([(r['time'], r['name']) for r in d2['rows'][:3]], [('10:00', '井上 子'), ('11:00', '青木 子'), ('', '')])
        self.assertEqual(days[1]['rows'][0], {'time': '9:00', 'name': '上田 子', 'att': ''})
        self.assertEqual(days[1]['staff'], '')

    def test_staff_saved_from_schedule_and_printed(self):
        url = reverse('reservations:monthly_schedule', args=[2026, 10])
        res = self.client.get(url)
        self.assertContains(res, 'name="staff_2026-10-02"')
        self.assertNotContains(res, 'name="staff_2026-10-05"')              # お休みの日は欄が無い
        self.assertContains(res, '業務日誌（4日/枚）')
        res = self.client.post(reverse('reservations:monthly_day_staff', args=[2026, 10]),
                               {'staff_2026-10-02': ' pm大坂 ', 'staff_2026-10-03': '終日土田', 'staff_2026-10-04': ''}, follow=True)
        self.assertContains(res, '直した日 2 日')
        self.assertContains(res, 'value="pm大坂"')
        self.assertContains(res, 'title="その日の担当">pm大坂')
        res = self.client.post(reverse('reservations:monthly_day_staff', args=[2026, 10]), {'staff_2026-10-03': ''}, follow=True)
        self.assertContains(res, '直した日 1 日')
        from .models import DayStaff
        self.assertEqual(list(DayStaff.objects.filter(facility=self.f).values_list('date', 'text')), [(datetime.date(2026, 10, 2), 'pm大坂')])
        # 業務日誌
        res = self.client.get(reverse('reservations:daily_log_pdf', args=[2026, 10]) + '?fmt=html')
        self.assertContains(res, '業務日誌')
        self.assertContains(res, '2026年10月2日')
        self.assertContains(res, 'pm大坂')
        self.assertContains(res, '井上 子')
        self.assertContains(res, '提供<br>形態')
        self.assertEqual(res.content.decode().count('class="pb"'), 5)      # 6 枚
        res = self.client.get(reservations_url := reverse('reservations:daily_log_pdf', args=[2026, 10]) + '?fmt=html&from=2026-10-03&to=2026-10-02')
        self.assertContains(res, '2026年10月2日')
        self.assertContains(res, '2026年10月3日')
        self.assertNotContains(res, '2026年10月6日')
        self.assertEqual(res.content.decode().count('class="pb"'), 0)
        # 候補：入れた担当と職員の表示名
        self.assertEqual(monthly.staff_suggestions(self.f), ['pm大坂', '大坂'])

    def test_printed_pages_fit_the_paper(self):
        """印刷したときに用紙からはみ出さない：業務日誌は A4 1 枚に 4 日、月間予定表は A3・A4 とも 1 枚（5 週・6 週の月）"""
        try:
            import weasyprint  # noqa: F401
            import pymupdf
        except (ImportError, OSError):
            self.skipTest('WeasyPrint か pymupdf が無い')

        def pages(url):
            res = self.client.get(url)
            self.assertEqual(res['Content-Type'], 'application/pdf')
            with pymupdf.open(stream=res.content, filetype='pdf') as doc:
                return len(doc), round(doc[0].rect.width), round(doc[0].rect.height)

        self.assertEqual(pages(reverse('reservations:daily_log_pdf', args=[2026, 10])), (6, 595, 842))   # 22 日 → 6 枚
        for year, month in ((2026, 10), (2026, 8)):            # 10月は 5 週、8月は 6 週
            url = reverse('reservations:monthly_schedule_pdf', args=[year, month])
            self.assertEqual(pages(url), (1, 1191, 842))                       # A3 横
            self.assertEqual(pages(url + '?paper=a4'), (1, 842, 595))          # A4 横
        # 画面のボタンは PDF をブラウザで開く（inline）。見本の画面の「印刷する（PDF）」も条件（A4）を引き継ぐ
        res = self.client.get(reverse('reservations:daily_log_pdf', args=[2026, 10]) + '?inline=1')
        self.assertTrue(res['Content-Disposition'].startswith('inline;'))
        res = self.client.get(reverse('reservations:monthly_schedule_pdf', args=[2026, 10]) + '?fmt=html&paper=a4')
        self.assertContains(res, '?paper=a4&amp;inline=1')
        self.assertContains(res, '印刷する（PDF）')
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(res, 'yotei/nisshi/?inline=1')
        self.assertContains(res, 'yotei/pdf/?paper=a4&inline=1')

    def test_pdf_uses_bundled_japanese_font(self):
        """PDF はサーバーのフォントに頼らず、同梱の IPAPゴシックを埋め込む（文字化け対策）"""
        try:
            import weasyprint  # noqa: F401
        except (ImportError, OSError):
            self.skipTest('WeasyPrint が無い')
        from config.pdf import font_css, font_path
        self.assertIsNotNone(font_path())
        self.assertIn('@font-face', font_css())
        self.assertIn('!important', font_css())
        res = self.client.get(reverse('reservations:daily_log_pdf', args=[2026, 10]))
        self.assertEqual(res['Content-Type'], 'application/pdf')
        # フォント名は圧縮されたストリームの中にあるので、ほどいてから探す
        import re
        import zlib
        parts = [res.content]
        for m in re.finditer(rb'stream\r?\n(.*?)\r?\nendstream', res.content, re.S):
            try:
                parts.append(zlib.decompress(m.group(1)))
            except zlib.error:
                pass
        fonts = set(re.findall(rb'/BaseFont\s*/[A-Z]+\+([A-Za-z0-9\-]+)', b''.join(parts)))
        self.assertTrue(fonts and all(f.startswith(b'IPAPGothic') for f in fonts), fonts)

    def test_only_for_ryoiku_layout(self):
        """ゆあーず以外（標準の画面の事業所）には、担当の欄・業務日誌・来られない日の書き方を出さない"""
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertNotContains(res, '業務日誌')
        self.assertNotContains(res, 'name="staff_2026-10-02"')
        self.assertEqual(self.client.get(reverse('reservations:daily_log_pdf', args=[2026, 10]) + '?fmt=html').status_code, 404)
        self.assertEqual(self.client.post(reverse('reservations:monthly_day_staff', args=[2026, 10]), {'staff_2026-10-02': 'x'}).status_code, 404)
        url = reverse('reservations:monthly_request_edit', args=[2026, 10, self.kids[0].pk])
        res = self.client.get(url)
        self.assertNotContains(res, '来られない日を書く')
        res = self.client.post(url, {'desired_count': '1', 'wish_mode': 'ng', 'ng_2026-10-04': '1', 'h_2026-10-02': ['10']})
        req = MonthlyRequest.objects.get(beneficiary=self.kids[0])
        self.assertEqual((req.wish_mode, req.wishes), ('ok', {'2026-10-02': [10]}))   # ng は受け付けない
        res = self.client.get(reverse('reservations:monthly_request_form', args=[2026, 10]) + '?fmt=html')
        self.assertNotContains(res, 'ダメな日')
        from . import scan
        out = scan.normalize({'name': '', 'year': 2026, 'month': 10, 'desired_count': 1, 'days': [], 'mode': 'ng',
                              'ng_days': [4], 'note': '', 'unreadable': ''}, self.f, 2026, 10, self.s)
        self.assertEqual((out['wish_mode'], out['ng_dates']), ('ok', []))
        customer = Customer.objects.create(facility=self.f, name='青木 母')
        customer.children.add(self.kids[0])
        self.assertNotContains(self.client.get(reverse('reservations_public:customer', args=[customer.token])), '来られない日を選ぶ')

    def test_other_facility_data_not_shown(self):
        g = Facility.objects.create(name='ほか', use_reservation=True)
        from .models import DayStaff
        DayStaff.objects.create(facility=g, date=datetime.date(2026, 10, 2), text='よそ')
        res = self.client.get(reverse('reservations:daily_log_pdf', args=[2026, 10]) + '?fmt=html')
        self.assertNotContains(res, 'よそ')


class AttendanceTests(TestCase):
    """実績（来た・欠席・キャンセル）：月間予定表の「実績を入れる」・日の画面・業務日誌の実績欄・一覧の「来」"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_ADMIN, display_name='大坂')
        self.client.login(username='ryo', password='pw12345678')
        self.kid = child(self.f, '青木')
        self.res, _ = services.create_reservation(self.f, self.kid, datetime.date(2026, 10, 2),
                                                  start_time=datetime.time(10, 0), notify=False)

    def test_set_from_schedule_and_print(self):
        url = reverse('reservations:monthly_schedule', args=[2026, 10])
        res = self.client.get(url)
        self.assertContains(res, 'id="ms-att-toggle"')
        self.assertContains(res, '<div title="実績が「来た」の回数">来</div>')
        res = self.client.post(reverse('reservations:monthly_attendance', args=[2026, 10]),
                               {'reservation': self.res.pk, 'value': 'attended'}, follow=True)
        self.assertContains(res, '実績を「来た」にしました')
        self.assertContains(res, 'class="ms-att attended">○</span>青木 子')
        self.assertContains(res, 'title="来た">来1')
        self.res.refresh_from_db()
        self.assertEqual((self.res.attendance, self.res.attendance_mark), ('attended', '○'))
        rows = monthly.request_rows(self.f, 2026, 10, self.s)
        self.assertEqual(next(r for r in rows if r['beneficiary'] == self.kid)['attended'], 1)
        # 業務日誌の「実績」欄
        res = self.client.get(reverse('reservations:daily_log_pdf', args=[2026, 10]) + '?fmt=html')
        self.assertContains(res, '<td class="n">青木 子</td><td class="k">1・2</td><td class="j">○</td>')
        self.assertContains(res, '実績：○ 来た／欠 欠席／取消 キャンセル')
        # 欠席 → 取消 → 未入力
        self.client.post(reverse('reservations:monthly_attendance', args=[2026, 10]), {'reservation': self.res.pk, 'value': 'absent'})
        self.res.refresh_from_db(); self.assertEqual(self.res.attendance_mark, '欠')
        self.client.post(reverse('reservations:monthly_attendance', args=[2026, 10]), {'reservation': self.res.pk, 'value': 'cancelled'})
        self.res.refresh_from_db(); self.assertEqual(self.res.attendance_mark, '取消')
        res = self.client.post(reverse('reservations:monthly_attendance', args=[2026, 10]), {'reservation': self.res.pk, 'value': ''}, follow=True)
        self.assertContains(res, '「未入力」')
        self.res.refresh_from_db(); self.assertEqual(self.res.attendance, '')
        res = self.client.post(reverse('reservations:monthly_attendance', args=[2026, 10]), {'reservation': self.res.pk, 'value': 'x'}, follow=True)
        self.assertContains(res, '実績の値が正しくありません')

    def test_set_from_day_page(self):
        url = reverse('reservations:day', args=[2026, 10, 2])
        res = self.client.get(url)
        self.assertContains(res, 'name="action" value="attendance"')
        res = self.client.post(url, {'action': 'attendance', 'reservation': self.res.pk, 'value': 'absent'}, follow=True)
        self.assertContains(res, '実績を「欠席」にしました')
        self.res.refresh_from_db()
        self.assertEqual(self.res.attendance, 'absent')

    def test_not_for_other_layouts_or_facilities(self):
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        self.assertEqual(self.client.post(reverse('reservations:monthly_attendance', args=[2026, 10]),
                                          {'reservation': self.res.pk, 'value': 'attended'}).status_code, 404)
        res = self.client.get(reverse('reservations:day', args=[2026, 10, 2]))
        self.assertNotContains(res, 'value="attendance"')
        self.client.post(reverse('reservations:day', args=[2026, 10, 2]),
                         {'action': 'attendance', 'reservation': self.res.pk, 'value': 'attended'})
        self.res.refresh_from_db(); self.assertEqual(self.res.attendance, '')
        self.assertNotContains(self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10])), 'id="ms-att-toggle"')
        # ほかの事業所の予約は触れない
        self.f.layout = Facility.LAYOUT_RYOIKU
        self.f.save(update_fields=['layout'])
        g, _ = ryoiku()
        other, _ = services.create_reservation(g, child(g, '他'), datetime.date(2026, 10, 2), start_time=datetime.time(10, 0), notify=False)
        self.assertEqual(self.client.post(reverse('reservations:monthly_attendance', args=[2026, 10]),
                                          {'reservation': other.pk, 'value': 'attended'}).status_code, 404)


class DailyBoardTests(TestCase):
    """きょうの予定（スマートフォン向け）：担当・時刻順の予定・実績・前後の営業日"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_ADMIN, display_name='大坂')
        self.client.login(username='ryo', password='pw12345678')
        self.kids = [child(self.f, n) for n in ('青木', '井上', '上田')]
        self.day = datetime.date(2026, 10, 2)    # 金
        self.r1, _ = services.create_reservation(self.f, self.kids[0], self.day, start_time=datetime.time(11, 0), notify=False, note='送迎あり')
        self.r2, _ = services.create_reservation(self.f, self.kids[1], self.day, start_time=datetime.time(10, 0), notify=False)
        monthly.set_attendance(self.r2, 'attended')

    def test_board(self):
        b = monthly.day_board(self.f, self.day, self.s)
        self.assertEqual([r.display_name for r in b['reservations']], ['井上 子', '青木 子'])
        self.assertEqual((b['total'], b['attended'], b['pending'], b['weekday'], b['closed']), (2, 1, 1, '金', False))
        self.assertEqual((b['prev'], b['next']), (datetime.date(2026, 9, 30), datetime.date(2026, 10, 3)))   # 木曜はお休み
        self.assertEqual(monthly.day_board(self.f, datetime.date(2026, 10, 5), self.s)['closed'], True)      # 月曜

    def test_page_and_actions(self):
        url = reverse('reservations:daily_board')
        res = self.client.get(url + '?d=2026-10-02')
        self.assertContains(res, '10月2日')
        self.assertContains(res, '（まだ入っていません）')
        self.assertContains(res, '送迎あり')
        self.assertContains(res, '?d=2026-09-30')
        self.assertContains(res, '?d=2026-10-03')
        self.assertContains(res, 'from=2026-10-02&to=2026-10-02')
        # 担当
        res = self.client.post(url + '?d=2026-10-02', {'action': 'staff', 'text': ' 終日土田 '}, follow=True)
        self.assertContains(res, 'その日の担当を保存しました')
        self.assertContains(res, '<div class="db-staff" id="db-staff-text">終日土田</div>')
        # 実績
        res = self.client.post(url + '?d=2026-10-02', {'action': 'attendance', 'reservation': self.r1.pk, 'value': 'absent'}, follow=True)
        self.assertContains(res, '実績を「欠席」にしました')
        self.assertContains(res, '欠席 1')
        self.r1.refresh_from_db(); self.assertEqual(self.r1.attendance, 'absent')
        res = self.client.post(url + '?d=2026-10-02', {'action': 'staff', 'text': ''}, follow=True)
        self.assertContains(res, 'その日の担当を空にしました')
        # 日付が変でもきょうを出す
        self.assertEqual(self.client.get(url + '?d=xx').status_code, 200)
        # お休みの日
        self.assertContains(self.client.get(url + '?d=2026-10-05'), 'お休みの日です')

    def test_only_for_ryoiku(self):
        self.assertContains(self.client.get(reverse('facilities:dashboard')), 'きょうの予定と担当')
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        self.assertEqual(self.client.get(reverse('reservations:daily_board')).status_code, 404)
        self.assertNotContains(self.client.get(reverse('facilities:dashboard')), 'きょうの予定と担当')


class StaffShiftTests(TestCase):
    """職員のシフト（曜日×時間帯）から「その日の担当」を入れる"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_ADMIN, display_name='大坂')
        self.client.login(username='ryo', password='pw12345678')

    def test_text_and_fill(self):
        from .models import DayStaff, StaffShift
        a = StaffShift.objects.create(facility=self.f, name='土田', weekdays=[1, 2, 5], part=StaffShift.PART_ALL)
        b = StaffShift.objects.create(facility=self.f, name='大坂', weekdays=[2, 5], part=StaffShift.PART_PM)
        StaffShift.objects.create(facility=self.f, name='休', weekdays=[2], part=StaffShift.PART_AM, is_active=False)
        self.assertEqual((a.label, b.label, a.weekdays_text), ('終日土田', 'pm大坂', '火・水・土'))
        shifts = list(StaffShift.objects.filter(facility=self.f, is_active=True))
        self.assertEqual(StaffShift.text_for(datetime.date(2026, 10, 7), shifts), '終日土田 pm大坂')   # 水
        self.assertEqual(StaffShift.text_for(datetime.date(2026, 10, 6), shifts), '終日土田')         # 火
        self.assertEqual(StaffShift.text_for(datetime.date(2026, 10, 4), shifts), '')                 # 日
        DayStaff.objects.create(facility=self.f, date=datetime.date(2026, 10, 6), text='手入力')
        n = monthly.fill_day_staff_from_shifts(self.f, datetime.date(2026, 10, 1), datetime.date(2026, 10, 31), self.s)
        staff = DayStaff.for_range(self.f, datetime.date(2026, 10, 1), datetime.date(2026, 10, 31))
        self.assertEqual(staff[datetime.date(2026, 10, 6)], '手入力')                                  # 入力済みは残す
        self.assertEqual(staff[datetime.date(2026, 10, 7)], '終日土田 pm大坂')
        self.assertEqual(staff[datetime.date(2026, 10, 3)], '終日土田 pm大坂')                          # 土
        self.assertNotIn(datetime.date(2026, 10, 5), staff)                                            # 月曜はお休み
        self.assertNotIn(datetime.date(2026, 10, 4), staff)                                            # 日曜は誰もいない
        self.assertEqual(n, len(staff) - 1)
        n = monthly.fill_day_staff_from_shifts(self.f, datetime.date(2026, 10, 1), datetime.date(2026, 10, 31), self.s, overwrite=True)
        self.assertEqual(DayStaff.objects.get(facility=self.f, date=datetime.date(2026, 10, 6)).text, '終日土田')

    def test_pages(self):
        from .models import DayStaff, StaffShift
        url = reverse('reservations:staff_shifts')
        res = self.client.get(url)
        self.assertContains(res, 'まだシフトがありません')
        res = self.client.post(url, {'action': 'add', 'name': ' 土田 ', 'weekdays': ['2', '5', 'x'], 'part': 'all'}, follow=True)
        self.assertContains(res, '終日土田（水・土）を登録しました')
        res = self.client.post(url, {'action': 'add', 'name': '', 'weekdays': ['2'], 'part': 'all'}, follow=True)
        self.assertContains(res, '名前・曜日・時間帯を入れてください')
        shift = StaffShift.objects.get(facility=self.f)
        res = self.client.post(url, {'action': 'edit', 'shift': shift.pk, 'name': '大坂', 'weekdays': ['5'], 'part': 'pm'}, follow=True)
        self.assertContains(res, 'pm大坂（土）に直しました')
        # 月間予定表から入れる
        sched = reverse('reservations:monthly_schedule', args=[2026, 10])
        res = self.client.get(sched)
        self.assertContains(res, 'シフトから入れる')
        self.assertContains(res, 'シフトの登録（1）')
        res = self.client.post(reverse('reservations:staff_shift_fill', args=[2026, 10]), {}, follow=True)
        self.assertContains(res, 'シフトから 10月の担当を入れました（5 日）')     # 10月の土曜は 3・10・17・24・31
        self.assertEqual(DayStaff.objects.filter(facility=self.f, text='pm大坂').count(), 5)
        self.assertContains(res, 'value="pm大坂"')
        res = self.client.post(url, {'action': 'toggle', 'shift': shift.pk}, follow=True)
        self.assertContains(res, '休止にしました')
        res = self.client.post(reverse('reservations:staff_shift_fill', args=[2026, 10]), {}, follow=True)
        self.assertContains(res, 'シフトがまだ登録されていません')
        res = self.client.post(url, {'action': 'delete', 'shift': shift.pk}, follow=True)
        self.assertContains(res, 'のシフトを消しました')
        self.assertFalse(StaffShift.objects.exists())
        self.assertEqual(DayStaff.objects.filter(facility=self.f).count(), 5)       # 入れた担当は残る

    def test_only_for_ryoiku_and_own_facility(self):
        from .models import StaffShift
        g, _ = ryoiku()
        other = StaffShift.objects.create(facility=g, name='よそ', weekdays=[2], part='all')
        url = reverse('reservations:staff_shifts')
        self.assertNotContains(self.client.get(url), 'よそ')
        self.assertEqual(self.client.post(url, {'action': 'delete', 'shift': other.pk}).status_code, 404)
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.post(reverse('reservations:staff_shift_fill', args=[2026, 10]), {}).status_code, 404)


class SheetImportTests(TestCase):
    """月予約利用希望を Excel・CSV から取り込む（ゆあーず）"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.aoki = child(self.f, '青木')
        self.inoue = child(self.f, '井上', '太郎')

    def xlsx(self, rows):
        import io
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = Workbook(); ws = wb.active
        for r in rows:
            ws.append(r)
        buf = io.BytesIO(); wb.save(buf)
        return SimpleUploadedFile('kibou.xlsx', buf.getvalue(),
                                  content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_template_has_names_and_days(self):
        from openpyxl import load_workbook
        import io
        res = self.client.get(reverse('reservations:monthly_request_sheet_template', args=[2026, 10]))
        self.assertEqual(res['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        ws = load_workbook(io.BytesIO(res.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][:4], ('氏名', '希望回数', '書き方', 1))
        self.assertEqual(rows[0][3 + 30], 31)
        self.assertEqual(rows[1][3], '木')                                  # 2026-10-01 は木曜
        self.assertEqual({rows[2][0], rows[3][0]}, {'青木 子', '井上 太郎'})
        # 標準の型では出ない
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        self.assertEqual(self.client.get(reverse('reservations:monthly_request_sheet_template', args=[2026, 10])).status_code, 404)

    def test_parse_ok_and_ng_modes(self):
        from . import sheet
        head = ['氏名', '希望回数', '書き方'] + list(range(1, 32))
        row_aoki = ['青木', '3', '○'] + [''] * 31
        row_aoki[3 + 1] = '○'          # 2日：終日
        row_aoki[3 + 2] = '10,11'      # 3日：10時・11時
        row_aoki[3 + 4] = '○'          # 5日：月曜はお休み → 入れない
        row_aoki[3 + 5] = '9'          # 6日（火）：9時は枠なし → 印だけなので終日
        row_inoue = ['井上 太郎', '2回', 'ダメな日'] + [''] * 31
        row_inoue[3 + 3] = '×'; row_inoue[3 + 10] = '×'
        row_inoue[3 + 5] = '10'        # 6日（火）：来られる日のうちの希望（10時）
        row_inoue[3 + 8] = '○'         # 9日（金）：終日に○
        row_none = ['山田', '1', '○'] + ['○'] * 31
        out = sheet.parse_rows([head, ['', '', '', '木'], row_aoki, row_inoue, row_none, [], ['書き方の説明']], self.f, 2026, 10, self.s)
        self.assertEqual(out['errors'], [])
        self.assertEqual(out['unmatched'], ['山田'])
        a, i = out['items']
        self.assertEqual((a['beneficiary'], a['desired'], a['mode']), (self.aoki, 3, 'ok'))
        self.assertEqual(a['wishes'], {'2026-10-02': 'all', '2026-10-03': [10, 11], '2026-10-06': 'all'})
        self.assertEqual(a['notes'], ['5日はお休みの日'])
        self.assertEqual((i['beneficiary'], i['desired'], i['mode'], i['ng_dates'], i['wishes']),
                         (self.inoue, 2, 'ng', ['2026-10-04', '2026-10-11'], {'2026-10-06': [10], '2026-10-09': 'all'}))
        bad = sheet.parse_rows([['名前', '1']], self.f, 2026, 10, self.s)
        self.assertIn('「氏名」の見出し', bad['errors'][0])

    def test_upload_xlsx_and_csv_saves_requests(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        head = ['氏名', '希望回数', '書き方'] + list(range(1, 32))
        row = ['青木 子', 4, '○'] + [None] * 31
        row[3 + 1] = '○'
        url = reverse('reservations:request_scan_upload', args=[2026, 10])
        res = self.client.post(url, {'images': [self.xlsx([head, row])]}, follow=True)
        self.assertContains(res, '1 人の利用希望を保存しました：青木 子')
        req = MonthlyRequest.objects.get(beneficiary=self.aoki, year=2026, month=10)
        self.assertEqual((req.source, req.desired_count, req.wishes), ('file', 4, {'2026-10-02': 'all'}))
        self.assertContains(res, 'Excel・CSV から')
        self.assertFalse(RequestScan.objects.exists())                       # 写真ではないので読み取り待ちには入らない
        csv_text = '氏名,希望回数,書き方,' + ','.join(str(d) for d in range(1, 32)) + '\n井上,2,ダメな日,' + ','.join('×' if d in (4, 11) else '' for d in range(1, 32)) + '\n山本,1,○,' + ','.join('○' for _ in range(31)) + '\n'
        res = self.client.post(url, {'images': [SimpleUploadedFile('kibou.csv', csv_text.encode('cp932'), content_type='text/csv')]}, follow=True)
        self.assertContains(res, '1 人の利用希望を保存しました：井上 太郎')
        self.assertContains(res, '台帳に見つからない名前があり、入れていません：山本')
        req = MonthlyRequest.objects.get(beneficiary=self.inoue, year=2026, month=10)
        self.assertEqual((req.wish_mode, req.ng_dates), ('ng', ['2026-10-04', '2026-10-11']))
        # 読めないファイル
        res = self.client.post(url, {'images': [SimpleUploadedFile('kibou.xlsx', b'not a zip', content_type='application/octet-stream')]}, follow=True)
        self.assertContains(res, 'を読めませんでした')

    def test_non_ryoiku_ignores_sheets(self):
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        head = ['氏名', '希望回数', '書き方'] + list(range(1, 32))
        res = self.client.post(reverse('reservations:request_scan_upload', args=[2026, 10]),
                               {'images': [self.xlsx([head, ['青木 子', 1, '○', '○']])]}, follow=True)
        self.assertContains(res, '用紙の写真（JPEG・PNG・HEIC など）を選んでください')
        self.assertFalse(MonthlyRequest.objects.exists())


class RyoikuScreenTests(TestCase):
    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.kid = child(self.f, '青木')

    def test_request_edit_and_list(self):
        url = reverse('reservations:monthly_request_edit', args=[2026, 10, self.kid.pk])
        res = self.client.get(url)
        self.assertContains(res, '終日')
        res = self.client.post(url, {'desired_count': '2', 'note': '午前がよい',
                                     'all_2026-10-03': '1', 'h_2026-10-02': ['10', '11'], 'h_2026-10-05': ['10']})
        self.assertRedirects(res, reverse('reservations:monthly_requests', args=[2026, 10]))
        req = MonthlyRequest.objects.get(beneficiary=self.kid)
        self.assertEqual(req.wishes, {'2026-10-03': 'all', '2026-10-02': [10, 11]})    # 月曜は落ちる
        self.assertEqual((req.desired_count, req.note), (2, '午前がよい'))
        res = self.client.get(reverse('reservations:monthly_requests', args=[2026, 10]))
        self.assertContains(res, '青木 子')
        self.assertContains(res, '2 回')
        # 一覧の「月間予定表を作る」で割り当て
        res = self.client.post(reverse('reservations:monthly_requests', args=[2026, 10]))
        self.assertRedirects(res, reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertEqual(Reservation.objects.filter(beneficiary=self.kid, status='confirmed').count(), 2)
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(res, '青木 子')
        res = self.client.get(reverse('reservations:monthly_schedule_pdf', args=[2026, 10]) + '?fmt=html')
        self.assertContains(res, '月間予定表')
        self.assertContains(res, 'size: A3 landscape')          # 既定は A3 横 1 枚
        self.assertNotContains(res, 'class="st"')               # 担当が無い月は担当の行を出さない
        DayStaff.objects.create(facility=self.f, date=datetime.date(2026, 10, 2), text='終日土田 pm大坂')
        res = self.client.get(reverse('reservations:monthly_schedule_pdf', args=[2026, 10]) + '?fmt=html&paper=a4')
        self.assertContains(res, 'size: A4 landscape')
        self.assertContains(res, '終日土田 pm大坂')              # 日付の下に「その日の担当」
        self.assertContains(res, 'width:33%')                   # 1 枠 3 人の箱を横に並べる
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(res, 'id="ms-fit-toggle"')          # 全体表示の切り替え
        self.assertContains(res, '印刷（A3）')
        res = self.client.get(reverse('reservations:monthly_request_form', args=[2026, 10]) + f'?fmt=html&b={self.kid.pk}')
        self.assertContains(res, '10月予約利用希望')
        self.assertContains(res, '月曜日・木曜日はお休み')

    def test_day_screen_add_with_time(self):
        day = datetime.date(2026, 10, 2)
        url = reverse('reservations:day', args=[2026, 10, 2])
        res = self.client.get(url + '?hour=14')
        self.assertContains(res, '時間枠')
        res = self.client.post(url, {'action': 'add', 'beneficiary': self.kid.pk, 'start_time': '14'})
        self.assertRedirects(res, url)
        r = Reservation.objects.get(beneficiary=self.kid)
        self.assertEqual((r.date, r.start_time), (day, datetime.time(14, 0)))
        res = self.client.get(url)
        self.assertContains(res, '14:00')
        self.assertContains(res, reverse('therapy:child', args=[self.kid.pk]))

    def test_slot_mode_required_for_monthly_pages(self):
        self.s.slot_mode = False
        self.s.save()
        res = self.client.get(reverse('reservations:monthly_requests', args=[2026, 10]))
        self.assertRedirects(res, reverse('reservations:calendar'))

    def test_settings_save_slot_fields(self):
        res = self.client.post(reverse('reservations:settings'), {
            'capacity': '10', 'booking_mode': 'auto', 'booking_from_days': '1', 'booking_until_days': '60',
            'slot_mode': 'on', 'slot_capacity': '2', 'slot_minutes': '45',
            'weekday_first_hour': '9', 'weekday_last_hour': '17', 'holiday_first_hour': '9', 'holiday_last_hour': '16',
            'break_hours': '12, 13', 'closed_weekdays': ['0'],
        })
        self.assertRedirects(res, reverse('reservations:calendar'))
        self.s.refresh_from_db()
        self.assertEqual((self.s.slot_capacity, self.s.weekday_first_hour, self.s.break_hours, self.s.closed_weekdays),
                         (2, 9, [12, 13], [0]))


    def test_schedule_swap_and_move_screen(self):
        other = child(self.f, '井上')
        d1, d2 = datetime.date(2026, 10, 3), datetime.date(2026, 10, 10)
        a, _ = services.create_reservation(self.f, self.kid, d1, start_time=datetime.time(9, 0), notify=False)
        b, _ = services.create_reservation(self.f, other, d2, start_time=datetime.time(10, 0), notify=False)
        page = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(page, '入れ替え・移動')
        self.assertContains(page, f'data-res="{a.pk}"')
        url = reverse('reservations:monthly_schedule_swap', args=[2026, 10])
        res = self.client.post(url, {'action': 'swap', 'a': a.pk, 'b': b.pk}, follow=True)
        self.assertContains(res, '入れ替えました')
        a.refresh_from_db()
        self.assertEqual((a.date, a.start_time.hour), (d2, 10))
        res = self.client.post(url, {'action': 'move', 'a': a.pk, 'date': '2026-10-17', 'hour': '11'}, follow=True)
        self.assertContains(res, '移しました')
        a.refresh_from_db()
        self.assertEqual((a.date, a.start_time.hour), (datetime.date(2026, 10, 17), 11))
        res = self.client.post(url, {'action': 'move', 'a': a.pk, 'date': '2026-10-19', 'hour': '11'}, follow=True)
        self.assertContains(res, '休業日')                              # 月曜はお休み
        self.assertContains(res, 'alert-danger')                        # エラーは赤い枠で出る
        # ほかの事業所の予約は触れない
        f2 = Facility.objects.create(name='ほか', use_reservation=True)
        x, _ = services.create_reservation(f2, child(f2, '他'), d1, notify=False)
        self.assertEqual(self.client.post(url, {'action': 'swap', 'a': x.pk, 'b': a.pk}).status_code, 404)

class CustomerWishPageTests(TestCase):
    def setUp(self):
        self.f, self.s = ryoiku()
        self.kid = child(self.f, '青木')
        self.customer = Customer.objects.create(facility=self.f, name='青木 母')
        self.customer.children.add(self.kid)
        self.url = reverse('reservations_public:customer', args=[self.customer.token])

    def test_wish_form_shown_and_saved(self):
        res = self.client.get(self.url)
        self.assertContains(res, '月の利用希望を送る')
        today = datetime.date.today()
        y, m = monthly.next_month(today.year, today.month)
        first = datetime.date(y, m, 1)
        # 来月の最初の「お休みでない日」
        day = next(d for d in monthly.month_days(y, m) if self.s.slot_hours(d))
        hour = self.s.slot_hours(day)[0]
        res = self.client.post(self.url, {'action': 'wish', 'beneficiary': self.kid.pk, 'wish_year': y, 'wish_month': m,
                                          'desired_count': '3', f'h_{day.isoformat()}': [str(hour)]})
        self.assertEqual(res.status_code, 302)
        req = MonthlyRequest.objects.get(beneficiary=self.kid, year=y, month=m)
        self.assertEqual((req.source, req.customer, req.wishes), ('web', self.customer, {day.isoformat(): [hour]}))
        # 過ぎた月は受けない
        res = self.client.post(self.url, {'action': 'wish', 'beneficiary': self.kid.pk, 'wish_year': first.year - 1,
                                          'wish_month': 1, 'desired_count': '1'}, follow=True)
        self.assertContains(res, '受け付けられません')

    def test_booking_with_time(self):
        today = datetime.date.today()
        day = next(d for d in (today + datetime.timedelta(days=i) for i in range(1, 10)) if self.s.slot_hours(d))
        hour = self.s.slot_hours(day)[0]
        res = self.client.post(self.url, {'action': 'book', 'beneficiary': self.kid.pk, 'date': day.isoformat(),
                                          'start_time': str(hour)}, follow=True)
        self.assertContains(res, 'ご予約を承りました')
        self.assertEqual(Reservation.objects.get(beneficiary=self.kid).start_time, datetime.time(hour, 0))


class WishAskTests(TestCase):
    """保護者に入力ページの URL を配る"""

    def setUp(self):
        from beneficiaries.models import Guardian
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.aoki, self.inoue, self.ueda, self.eguchi = (child(self.f, n) for n in ('青木', '井上', '上田', '江口'))
        self.line_mom = Customer.objects.create(facility=self.f, name='青木 母', line_user_id='U-aoki')
        self.line_mom.children.add(self.aoki)
        self.paper_mom = Customer.objects.create(facility=self.f, name='井上 母')
        self.paper_mom.children.add(self.inoue)
        # 上田・江口は顧客が無い。上田は保護者台帳に LINE つきの母、江口は台帳も無い
        Guardian.objects.create(beneficiary=self.ueda, last_name='上田', first_name='花子', phone='090-1111-2222',
                                is_primary=True)
        self.url = reverse('reservations:monthly_wish_ask', args=[2026, 10])

    def test_targets_and_ask(self):
        rows, no_contact = monthly.wish_targets(self.f, 2026, 10, 'http://t.example/')
        self.assertEqual({r['customer'] for r in rows}, {self.line_mom, self.paper_mom})
        self.assertEqual(set(no_contact), {self.ueda, self.eguchi})
        aoki_row = next(r for r in rows if r['customer'] == self.line_mom)
        self.assertEqual(aoki_row['url'], f'http://t.example/yoyaku/mypage/{self.line_mom.token}/?wish=2026-10#wishCard')
        # 井上は用紙が届いている → 「まだの方だけ」なら青木の母だけ
        monthly.save_request(self.f, self.inoue, 2026, 10, 2, {'2026-10-03': 'all'})
        made = monthly.ask_for_wishes(self.f, 2026, 10, self.s, base='http://t.example/',
                                      deadline=datetime.date(2026, 9, 20))
        self.assertEqual([n.customer for n in made], [self.line_mom])
        n = made[0]
        self.assertEqual((n.kind, n.status, n.to_line_id), ('wish', 'pending', 'U-aoki'))
        self.assertIn('10月のご利用希望の入力をお願いします（青木 子さん）', n.body)
        self.assertIn('9月20日(日)までに', n.body)
        self.assertIn(aoki_row['url'], n.body)
        # もう一度積むと、まだ送っていないお願いは置き換える（重ならない）
        monthly.ask_for_wishes(self.f, 2026, 10, self.s, base='http://t.example/', only_missing=False)
        self.assertEqual(ReservationNotice.objects.filter(kind='wish', customer=self.line_mom).count(), 1)
        paper = ReservationNotice.objects.get(kind='wish', customer=self.paper_mom)
        self.assertEqual(paper.status, 'manual')           # LINE が無い方は「コピーして送る」

    def test_make_contacts_from_guardians(self):
        from beneficiaries.models import Guardian
        # 江口のきょうだい（江口 弟）も同じ母。LINE でつながっている
        Guardian.objects.create(beneficiary=self.eguchi, last_name='江口', first_name='母', line_user_id='U-eguchi',
                                line_linked=True, is_primary=True)
        brother = child(self.f, '江口', '弟')
        Guardian.objects.create(beneficiary=brother, last_name='江口', first_name='母', line_user_id='U-eguchi',
                                line_linked=True, is_primary=True)
        _, no_contact = monthly.wish_targets(self.f, 2026, 10)
        made, joined = monthly.make_contacts(self.f, sorted(no_contact, key=lambda b: b.pk))
        self.assertEqual(len(made), 2)
        self.assertEqual(len(joined), 1)
        eguchi = Customer.objects.get(line_user_id='U-eguchi')
        self.assertEqual(set(eguchi.children.all()), {self.eguchi, brother})
        ueda = Customer.objects.get(children=self.ueda)
        self.assertEqual((ueda.name, ueda.phone, ueda.line_user_id), ('上田 花子', '090-1111-2222', ''))
        self.assertEqual(monthly.wish_targets(self.f, 2026, 10)[1], [])

    @mock.patch('line_integration.sending.push_text', return_value=(True, ''))
    def test_screen_ask_sends_line_only_for_wish(self, push):
        # ほかの送信待ち（ご利用日が決まりました など）は一緒に送らない
        other = services.queue_notice(self.f, ReservationNotice.KIND_ACCEPTED, self.s, customer=self.line_mom,
                                      day=datetime.date(2026, 10, 1), body='別の通知')
        page = self.client.get(reverse('reservations:monthly_requests', args=[2026, 10]))
        self.assertContains(page, '保護者に入力してもらう')
        self.assertContains(page, '上田 子')                # 連絡先の無い方
        self.assertContains(page, f'/yoyaku/mypage/{self.paper_mom.token}/?wish=2026-10')
        res = self.client.post(self.url, {'action': 'ask', 'deadline': '2026-09-20', 'scope': 'missing'}, follow=True)
        self.assertContains(res, 'LINE で 1 件送りました')
        self.assertContains(res, '文面をコピー')
        self.assertEqual(push.call_count, 1)
        self.assertEqual(push.call_args.args[1], 'U-aoki')
        other.refresh_from_db()
        self.assertEqual(other.status, 'pending')
        # 手渡しのお願いを「送った」にする
        paper = ReservationNotice.objects.get(kind='wish', customer=self.paper_mom)
        self.client.post(self.url, {'action': 'handed', 'notice': paper.pk})
        paper.refresh_from_db()
        self.assertEqual(paper.status, 'sent')
        # 連絡先を作る
        self.client.post(self.url, {'action': 'contacts'})
        self.assertTrue(Customer.objects.filter(children=self.eguchi).exists())

    def test_parent_input_reaches_list_and_staff_group(self):
        self.s.notify_group_id = 'G-staff'
        self.s.save()
        sister = child(self.f, '青木', '妹')
        self.line_mom.children.add(sister)
        today = datetime.date.today()
        y, m = monthly.next_month(today.year, today.month)
        page_url = reverse('reservations_public:customer', args=[self.line_mom.token])
        page = self.client.get(page_url + f'?wish={y}-{m:02d}')
        self.assertContains(page, '（まだ）')
        day = next(d for d in monthly.month_days(y, m) if self.s.slot_hours(d))
        res = self.client.post(page_url, {'action': 'wish', 'beneficiary': self.aoki.pk, 'wish_year': y, 'wish_month': m,
                                          'desired_count': '2', f'all_{day.isoformat()}': '1'})
        # 送ったあとは、まだのきょうだいの欄へ
        self.assertEqual(res['Location'], f'{page_url}?wish={y}-{m:02d}&child={sister.pk}#wishCard')
        req = MonthlyRequest.objects.get(beneficiary=self.aoki, year=y, month=m)
        self.assertEqual((req.source, req.desired_count, req.wishes), ('web', 2, {day.isoformat(): 'all'}))
        group = ReservationNotice.objects.get(kind='group', to_line_id='G-staff')
        self.assertIn('青木 子さんの利用希望が届きました', group.body)
        self.assertNotIn('母', group.body)
        # 職員の一覧に「保護者が入力」で出る
        lst = self.client.get(reverse('reservations:monthly_requests', args=[y, m]))
        self.assertContains(lst, '保護者が入力')
        # 予定を組んだあとに直すと、その旨を返す
        monthly.assign_month(self.f, y, m, self.s, notify=False)
        res = self.client.post(page_url, {'action': 'wish', 'beneficiary': self.aoki.pk, 'wish_year': y, 'wish_month': m,
                                          'desired_count': '1', f'all_{day.isoformat()}': '1'}, follow=True)
        self.assertContains(res, 'すでに組んでいる')
        self.assertTrue(ReservationNotice.objects.filter(kind='group', body__contains='直されました').exists())


_SCAN_TMP = __import__('tempfile').mkdtemp()


def _photo(name='form.png'):
    import io
    from PIL import Image
    from django.core.files.uploadedfile import SimpleUploadedFile
    buf = io.BytesIO()
    Image.new('RGB', (60, 80), (255, 255, 255)).save(buf, format='PNG')
    return SimpleUploadedFile(name, buf.getvalue(), content_type='image/png')


def _ai_reply(payload):
    from types import SimpleNamespace
    import json
    return SimpleNamespace(stop_reason='end_turn', content=[
        SimpleNamespace(type='thinking', thinking=''),
        SimpleNamespace(type='text', text=json.dumps(payload, ensure_ascii=False))])


READ = {'name': '青木 子', 'year': 2026, 'month': 10, 'desired_count': 4,
        'days': [{'day': 2, 'all_day': False, 'hours': [10, 11, 9]},
                 {'day': 3, 'all_day': True, 'hours': []},
                 {'day': 5, 'all_day': True, 'hours': []}],
        'note': '午前がよい', 'unreadable': '12日の行はかすれている'}


@override_settings(MEDIA_ROOT=_SCAN_TMP, ANTHROPIC_API_KEY='test-key')
class RequestScanTests(TestCase):
    """紙の利用希望の写真を読み取って反映する"""

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        import shutil
        shutil.rmtree(_SCAN_TMP, ignore_errors=True)

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.kid = child(self.f, '青木')
        self.other = child(self.f, '井上')

    def test_to_wishes_drops_closed_days_and_missing_hours(self):
        from . import scan
        out = scan.normalize(READ, self.f, 2026, 10, self.s)
        self.assertEqual(out['wishes'], {'2026-10-02': [10, 11], '2026-10-03': 'all'})
        self.assertEqual(out['ignored'], ['2日 9時（枠の無い時刻）', '5日（お休みの日）'])
        self.assertEqual((out['desired_count'], out['note'], out['month_mismatch']), (4, '午前がよい', False))
        self.assertEqual(out['slot_count'], 2 + len(self.s.slot_hours(datetime.date(2026, 10, 3))))
        unknown = scan.normalize(dict(READ, desired_count=-1, month=11, days=[]), self.f, 2026, 10, self.s)
        self.assertIsNone(unknown['desired_count'])
        self.assertTrue(unknown['month_mismatch'])
        cal = scan.calendar_text(self.f, 2026, 10, self.s).splitlines()
        self.assertEqual((cal[0], cal[2]), ('1日(木) お休み', '3日(土) 9時、10時、11時、13時、14時、15時、16時、17時'))

    @mock.patch('reservations.scan.anthropic.Anthropic')
    def test_upload_read_review_and_save(self, client_cls):
        client_cls.return_value.messages.create.return_value = _ai_reply(READ)
        lst = reverse('reservations:monthly_requests', args=[2026, 10])
        res = self.client.post(reverse('reservations:request_scan_upload', args=[2026, 10]),
                               {'images': [_photo('a.png'), _photo('b.png')]})
        self.assertRedirects(res, lst + '#scans', fetch_redirect_response=False)
        scans = list(RequestScan.objects.order_by('pk'))
        self.assertEqual(len(scans), 2)
        page = self.client.get(lst)
        self.assertContains(page, '紙の用紙を写真で取り込む')
        self.assertContains(page, 'AIで読み取る')
        # 写真はこの事業所の職員だけが見られる
        self.assertEqual(self.client.get(scans[0].image.url).status_code, 200)

        res = self.client.post(reverse('reservations:request_scan_extract', args=[2026, 10, scans[0].pk]))
        data = res.json()
        self.assertEqual(res.status_code, 200, data)
        self.assertEqual((data['beneficiary_id'], data['desired_count'], data['days']), (self.kid.pk, 4, 2))
        edit = reverse('reservations:monthly_request_edit', args=[2026, 10, self.kid.pk])
        self.assertEqual(data['review_url'], f'{edit}?scan={scans[0].pk}')
        # 画像と、どの月の用紙かを AI に渡している
        kwargs = client_cls.return_value.messages.create.call_args.kwargs
        content = kwargs['messages'][0]['content']
        self.assertEqual(content[0]['type'], 'image')
        self.assertIn('2026年10月', content[1]['text'])
        self.assertIn('5日(月) お休み', content[1]['text'])
        self.assertEqual(kwargs['output_config']['format']['type'], 'json_schema')

        # 確認画面：写真と、読み取った○が入った表（まだ保存しない）
        page = self.client.get(data['review_url'])
        self.assertContains(page, '写真から読み取りました')
        self.assertContains(page, '5日（お休みの日）')
        self.assertContains(page, 'name="all_2026-10-03" value="1" class="allday" checked')
        self.assertContains(page, 'value="4"')
        self.assertFalse(MonthlyRequest.objects.exists())

        # 保存すると利用希望に入り、写真は反映ずみ。次の写真（氏名が読めた）があればそこへ進む
        scans[1].beneficiary, scans[1].status, scans[1].extracted = self.other, RequestScan.STATUS_EXTRACTED, {'wishes': {}}
        scans[1].save()
        res = self.client.post(edit, {'scan': scans[0].pk, 'desired_count': '4', 'note': '午前がよい',
                                      'all_2026-10-03': '1', 'h_2026-10-02': ['10', '11']})
        self.assertRedirects(res, reverse('reservations:monthly_request_edit', args=[2026, 10, self.other.pk])
                             + f'?scan={scans[1].pk}', fetch_redirect_response=False)
        req = MonthlyRequest.objects.get(beneficiary=self.kid, year=2026, month=10)
        self.assertEqual((req.source, req.desired_count, req.wishes),
                         ('photo', 4, {'2026-10-02': [10, 11], '2026-10-03': 'all'}))
        scans[0].refresh_from_db()
        self.assertEqual((scans[0].status, scans[0].request), ('imported', req))
        self.assertContains(self.client.get(lst), '用紙の写真から')

    @mock.patch('reservations.scan.anthropic.Anthropic')
    def test_unknown_name_then_assign(self, client_cls):
        client_cls.return_value.messages.create.return_value = _ai_reply(dict(READ, name=''))
        self.client.post(reverse('reservations:request_scan_upload', args=[2026, 10]), {'images': [_photo()]})
        sc = RequestScan.objects.get()
        data = self.client.post(reverse('reservations:request_scan_extract', args=[2026, 10, sc.pk])).json()
        self.assertEqual((data['beneficiary_id'], data['review_url']), (None, ''))
        res = self.client.post(reverse('reservations:request_scan_action', args=[2026, 10, sc.pk]),
                               {'action': 'assign', 'beneficiary': self.other.pk})
        self.assertRedirects(res, reverse('reservations:monthly_request_edit', args=[2026, 10, self.other.pk])
                             + f'?scan={sc.pk}', fetch_redirect_response=False)
        sc.refresh_from_db()
        self.assertEqual(sc.beneficiary, self.other)
        # 氏名が別の子なら確認画面で知らせる
        sc.extracted = dict(sc.extracted, name='青木 子')
        sc.save()
        page = self.client.get(reverse('reservations:monthly_request_edit', args=[2026, 10, self.other.pk]) + f'?scan={sc.pk}')
        self.assertContains(page, '「青木 子」さんの用紙かもしれません')
        # 消す
        self.client.post(reverse('reservations:request_scan_action', args=[2026, 10, sc.pk]), {'action': 'delete'})
        self.assertFalse(RequestScan.objects.exists())

    def test_single_upload_from_edit_page_and_errors(self):
        edit = reverse('reservations:monthly_request_edit', args=[2026, 10, self.kid.pk])
        self.assertContains(self.client.get(edit), '用紙の写真から読み取る')
        res = self.client.post(reverse('reservations:request_scan_upload', args=[2026, 10]),
                               {'images': [_photo()], 'beneficiary': self.kid.pk, 'next': 'edit'})
        sc = RequestScan.objects.get()
        self.assertRedirects(res, f'{edit}?scan={sc.pk}', fetch_redirect_response=False)
        self.assertContains(self.client.get(f'{edit}?scan={sc.pk}'), 'AI が用紙を読み取っています')
        with mock.patch('reservations.scan.anthropic.Anthropic', side_effect=RuntimeError('boom')):
            res = self.client.post(reverse('reservations:request_scan_extract', args=[2026, 10, sc.pk]))
        self.assertEqual(res.status_code, 500)
        self.assertIn('読み取りに失敗', res.json()['error'])
        with override_settings(ANTHROPIC_API_KEY=''):
            res = self.client.post(reverse('reservations:request_scan_extract', args=[2026, 10, sc.pk]))
        self.assertIn('ANTHROPIC_API_KEY', res.json()['error'])
        # 画像でないファイルは受けない
        from django.core.files.uploadedfile import SimpleUploadedFile
        res = self.client.post(reverse('reservations:request_scan_upload', args=[2026, 10]),
                               {'images': [SimpleUploadedFile('a.txt', b'x', content_type='text/plain')]}, follow=True)
        self.assertContains(res, '写真（JPEG')
        # ほかの事業所の写真は触れない・見られない
        f2 = Facility.objects.create(name='ほか', use_reservation=True)
        other = RequestScan.objects.create(facility=f2, year=2026, month=10, image=_photo('x.png'))
        self.assertEqual(self.client.post(reverse('reservations:request_scan_extract', args=[2026, 10, other.pk])).status_code, 404)
        self.assertEqual(self.client.get(f'{edit}?scan={other.pk}').status_code, 404)
        self.assertEqual(self.client.get(other.image.url).status_code, 404)

class PresetAndSeedTests(TestCase):
    def test_rename_migration(self):
        """「りょういく」→「発達支援ルーム　ゆあーず」の書き換え（ほかの事業所・独自の署名は触らない）"""
        import importlib
        from django.apps import apps as django_apps
        mig = importlib.import_module('reservations.migrations.0009_rename_ryoiku_facility')
        old = Facility.objects.create(name='りょういく')
        services.get_setting(old).__class__.objects.filter(facility=old).update(signature='りょういく')
        custom = Facility.objects.create(name='りょういく2')
        mig.rename(django_apps, None)
        old.refresh_from_db(); custom.refresh_from_db()
        self.assertEqual((old.name, custom.name), ('発達支援ルーム　ゆあーず', 'りょういく2'))
        self.assertEqual(services.get_setting(old).signature, '発達支援ルーム　ゆあーず')
        mig.rename_back(django_apps, None)
        old.refresh_from_db()
        self.assertEqual(old.name, 'りょういく')

    def test_create_facility_preset(self):
        from io import StringIO
        from django.core.management import call_command
        out = StringIO()
        call_command('create_facility', '発達支援ルーム　ゆあーず', '--admin', 'ryoiku', '--password', 'pw12345678',
                     '--preset', 'ryoiku', stdout=out)
        f = Facility.objects.get(name='発達支援ルーム　ゆあーず')
        self.assertTrue(f.use_reservation and f.use_therapy_record)
        self.assertFalse(f.use_billing or f.use_schedule)      # 請求と来所予定（出欠）は使わない
        self.assertEqual((f.layout, f.trial_ai_limit), ('ryoiku', 20))   # メニューは基本機能／お試し／設定、お試しの AI は 20 回
        s = services.get_setting(f)
        self.assertTrue(s.slot_mode)
        self.assertEqual((s.slot_capacity, s.closed_weekdays, s.weekday_first_hour, s.holiday_first_hour), (3, [0, 3], 10, 9))

    def test_seed_demo_for_slot_facility(self):
        from io import StringIO
        from django.core.management import call_command
        f, s = ryoiku()
        StaffAccount.objects.create_user('ryo', password='pw12345678', facility=f, role=StaffAccount.ROLE_ADMIN)
        call_command('seed_demo', '--facility', str(f.pk), stdout=StringIO())
        today = datetime.date.today()
        self.assertTrue(MonthlyRequest.objects.filter(facility=f, year=today.year, month=today.month).exists())
        self.assertTrue(Reservation.objects.filter(facility=f, status='confirmed', start_time__isnull=False).exists())
        from therapy.models import TherapyProfile, TherapyRecord
        self.assertTrue(TherapyRecord.objects.filter(facility=f).exists())
        self.assertTrue(TherapyProfile.objects.filter(beneficiary__facility=f).exists())
        # --reset で入れ直せる
        call_command('seed_demo', '--facility', str(f.pk), '--reset', stdout=StringIO())


class RyoikuConflictTests(TestCase):
    """編集の競合：利用希望の転記・その日の担当・職員のシフト"""

    def setUp(self):
        self.f, self.s = ryoiku()
        StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.kid = child(self.f, '青木')

    def test_monthly_request_conflict_with_parent_or_staff(self):
        from config.concurrency import version_token
        url = reverse('reservations:monthly_request_edit', args=[2026, 10, self.kid.pk])
        page = self.client.get(url)
        self.assertContains(page, 'name="version" value=""')                 # まだ無いときは版なし
        # 開いているあいだに保護者が入力ページから送った
        monthly.save_request(self.f, self.kid, 2026, 10, 2, {'2026-10-02': 'all'}, source=MonthlyRequest.SOURCE_WEB)
        res = self.client.post(url, {'desired_count': '5', 'h_2026-10-03': ['10'], 'version': ''}, follow=True)
        self.assertContains(res, '他の職員か保護者が')
        req = MonthlyRequest.objects.get(beneficiary=self.kid)
        self.assertEqual((req.desired_count, req.source), (2, 'web'))            # 保護者の入力を上書きしない
        page = self.client.get(url)
        self.assertContains(page, f'data-concurrency="monthly_request:{req.pk}"')
        res = self.client.post(url, {'desired_count': '5', 'h_2026-10-03': ['10'], 'version': version_token(req)})
        req.refresh_from_db()
        self.assertEqual(req.desired_count, 5)
        # 消すときも確かめる
        opened = version_token(req)
        monthly.save_request(self.f, self.kid, 2026, 10, 3, {}, source=MonthlyRequest.SOURCE_WEB)
        self.client.post(url, {'action': 'delete', 'version': opened})
        self.assertTrue(MonthlyRequest.objects.filter(beneficiary=self.kid).exists())

    def test_day_staff_saves_only_changed_days(self):
        from .models import DayStaff
        DayStaff.objects.create(facility=self.f, date=datetime.date(2026, 10, 2), text='終日土田')
        DayStaff.objects.create(facility=self.f, date=datetime.date(2026, 10, 3), text='am大坂')
        page = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(page, 'name="staff_was_2026-10-02" value="終日土田"')
        # 画面を開いたあとに、他の職員が 2 日と 3 日を直した
        DayStaff.objects.filter(date=datetime.date(2026, 10, 2)).update(text='終日早田')
        DayStaff.objects.filter(date=datetime.date(2026, 10, 3)).update(text='pm大坂')
        # わたしは 3 日だけ直して、月の担当をまとめて保存（2 日は開いたときの値のまま送られる）
        res = self.client.post(reverse('reservations:monthly_day_staff', args=[2026, 10]), {
            'staff_2026-10-02': '終日土田', 'staff_was_2026-10-02': '終日土田',
            'staff_2026-10-03': '終日土田', 'staff_was_2026-10-03': 'am大坂',
            'staff_2026-10-06': 'pm大坂', 'staff_was_2026-10-06': '',
        }, follow=True)
        texts = dict(DayStaff.objects.filter(facility=self.f).values_list('date__day', 'text'))
        self.assertEqual(texts[2], '終日早田')        # 変えていない日は古い値で戻さない
        self.assertEqual(texts[3], 'pm大坂')          # 他の職員が先に直した日は上書きしない
        self.assertEqual(texts[6], 'pm大坂')          # 直した日は保存
        self.assertContains(res, '10/3 の担当は、この画面を開いたあとに他の職員が直していたため')
        self.assertContains(res, '直した日 1 日')
        # きょうの予定の画面（1 日だけ）も同じ
        board = reverse('reservations:daily_board') + '?d=2026-10-03'
        self.assertContains(self.client.get(board), 'name="was" value="pm大坂"')
        res = self.client.post(board, {'action': 'staff', 'text': '終日土田', 'was': 'am大坂'}, follow=True)
        self.assertContains(res, 'この画面を開いたあとに他の職員が直していた')
        self.client.post(board, {'action': 'staff', 'text': '終日土田', 'was': 'pm大坂'})
        self.assertEqual(DayStaff.objects.get(date=datetime.date(2026, 10, 3)).text, '終日土田')

    def test_staff_shift_edit_conflict(self):
        from config.concurrency import version_token
        from .models import StaffShift
        shift = StaffShift.objects.create(facility=self.f, name='土田', weekdays=[1, 2], part='all')
        opened = version_token(shift)
        self.assertContains(self.client.get(reverse('reservations:staff_shifts')), f'data-concurrency="staff_shift:{shift.pk}"')
        self.client.post(reverse('reservations:staff_shifts'), {'action': 'toggle', 'shift': shift.pk})   # 他の職員が休止にした
        res = self.client.post(reverse('reservations:staff_shifts'), {'action': 'edit', 'shift': shift.pk, 'name': '土田',
                                                                      'weekdays': ['4'], 'part': 'am', 'version': opened}, follow=True)
        self.assertContains(res, '他の職員が')
        shift.refresh_from_db()
        self.assertEqual((shift.weekdays, shift.part, shift.is_active), ([1, 2], 'all', False))


class CannotPairTests(TestCase):
    """同じ時間にできない利用者（共演NG）：同じ日の同じ時間枠には入れない。時間がずれていれば同じ日でも入れる"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.force_login(self.user)
        d = datetime.date(2018, 1, 1)
        self.a = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', last_name_kana='あおき', date_of_birth=d)
        self.b = Beneficiary.objects.create(facility=self.f, last_name='木村', first_name='空', last_name_kana='きむら', date_of_birth=d)
        self.c = Beneficiary.objects.create(facility=self.f, last_name='佐藤', first_name='花', last_name_kana='さとう', date_of_birth=d)
        self.a.cannot_pair.add(self.b)
        self.day = datetime.date(2026, 10, 6)      # 火
        self.day2 = datetime.date(2026, 10, 7)     # 水

    def test_symmetric_and_shown_on_detail(self):
        self.assertEqual(self.b.pair_names(), ['青木 子'])
        res = self.client.get(reverse('beneficiaries:detail', args=[self.a.pk]))
        self.assertContains(res, '同じ時間にできない利用者')
        self.assertContains(res, 'name="cannot_pair" value="%d" checked' % self.b.pk)
        # 編集で付け替えられる（相手の側にも付く）
        res = self.client.post(reverse('beneficiaries:update', args=[self.a.pk]), {
            'last_name': '青木', 'first_name': '子', 'date_of_birth': '2018-01-01', 'status': 'active', 'gender': 'male',
            'cannot_pair': [self.c.pk]})
        self.assertEqual(res.status_code, 302)
        self.assertEqual(set(self.a.cannot_pair.values_list('pk', flat=True)), {self.c.pk})
        self.assertEqual(list(self.b.cannot_pair.all()), [])
        self.assertEqual(list(self.c.cannot_pair.all()), [self.a])

    def test_same_day_other_time_is_ok(self):
        services.create_reservation(self.f, self.b, self.day, start_time=10)
        ra, _ = services.create_reservation(self.f, self.a, self.day, start_time=11)     # 時間がずれていればよい
        self.assertEqual(ra.status, Reservation.STATUS_CONFIRMED)
        self.assertNotContains(self.client.get(reverse('reservations:day', args=[2026, 10, 6])), '組み合わせNG')
        with self.assertRaises(services.ReservationError) as cm:
            services.move_reservation(ra, self.day, start_time=10)          # 同じ時間へは動かせない
        self.assertIn('10:00', str(cm.exception))
        self.assertIn('時間をずらせば入れられます', str(cm.exception))

    def test_create_move_link_swap_blocked(self):
        services.create_reservation(self.f, self.b, self.day, start_time=10)
        with self.assertRaises(services.ReservationError) as cm:
            services.create_reservation(self.f, self.a, self.day, start_time=10)
        self.assertIn('木村 空 さんの予約があるため', str(cm.exception))
        # 保護者からの申し込み（LINE・入力ページ）は止めずに受け、職員の画面で赤く示す
        ra_line, _ = services.create_reservation(self.f, self.a, self.day, start_time=10, source=Reservation.SOURCE_LINE)
        self.assertEqual(ra_line.status, Reservation.STATUS_CONFIRMED)
        res = self.client.get(reverse('reservations:day', args=[2026, 10, 6]))
        self.assertContains(res, '組み合わせNG：木村 空')
        self.assertContains(res, '組み合わせNG：青木 子')
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(res, 'ms-box pair')
        self.assertContains(res, '組み合わせNG：木村 空 さんと同じ時間')
        res = self.client.get(reverse('reservations:daily_board') + '?d=2026-10-06')
        self.assertContains(res, 'pair-ng')
        services.cancel_reservation(ra_line, notify=False)
        self.assertNotContains(self.client.get(reverse('reservations:day', args=[2026, 10, 6])), '組み合わせNG')
        services.create_reservation(self.f, self.c, self.day, start_time=10)             # 関係のない人は入れる
        ra, _ = services.create_reservation(self.f, self.a, self.day2, start_time=10)
        with self.assertRaises(services.ReservationError):
            services.move_reservation(ra, self.day, start_time=10)
        rc = Reservation.objects.get(beneficiary=self.c)
        with self.assertRaises(services.ReservationError):
            monthly.swap_reservations(ra, rc)      # 青木が木村と同じ日・同じ時間へ移ることになる
        guest, _ = services.create_reservation(self.f, None, self.day, guest_name='新しい子', start_time=10)
        with self.assertRaises(services.ReservationError):
            services.link_reservation(guest, self.a)
        # 画面からも止まる
        res = self.client.post(reverse('reservations:day', args=[2026, 10, 6]),
                               {'action': 'add', 'beneficiary': self.a.pk, 'start_time': '10'}, follow=True)
        self.assertContains(res, '同じ時間にできない')
        self.assertEqual(Reservation.objects.filter(beneficiary=self.a, date=self.day, status__in=Reservation.ACTIVE_STATUSES).count(), 0)

    def test_assign_month_keeps_pairs_apart(self):
        for kid in (self.a, self.b):
            monthly.save_request(self.f, kid, 2026, 10, 2, {'2026-10-06': [10], '2026-10-07': [10], '2026-10-09': [10], '2026-10-10': [10]})
        result = monthly.assign_month(self.f, 2026, 10, self.s, notify=False)
        self.assertEqual((len(result.made), result.short), (4, []))      # 4 日を 2 人で分け合う
        slots_a = set(Reservation.objects.filter(beneficiary=self.a).values_list('date', 'start_time'))
        slots_b = set(Reservation.objects.filter(beneficiary=self.b).values_list('date', 'start_time'))
        self.assertEqual(slots_a & slots_b, set())
        rows = {r['beneficiary'].pk: r for r in monthly.request_rows(self.f, 2026, 10, self.s)}
        self.assertEqual(rows[self.a.pk]['pair_names'], ['木村 空'])
        self.assertContains(self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10])), '同じ時間にできない：木村 空')

    def test_assign_month_same_day_other_hour(self):
        """1日しか来られない2人でも、時間をずらして同じ日に入れる"""
        for kid in (self.a, self.b):
            monthly.save_request(self.f, kid, 2026, 10, 1, {'2026-10-06': [10, 11]})
        result = monthly.assign_month(self.f, 2026, 10, self.s, notify=False)
        self.assertEqual((len(result.made), result.short), (2, []))
        hours = sorted(r.start_time.hour for r in Reservation.objects.filter(date=self.day))
        self.assertEqual(hours, [10, 11])


class SiblingSeatTests(TestCase):
    """きょうだい：同じ時間枠の日は「きょうだいで1枠にまとめる」を選べる（まとめない日はそれぞれ1枠）"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.force_login(self.user)
        d = datetime.date(2018, 1, 1)
        mk = lambda last, first, kana: Beneficiary.objects.create(facility=self.f, last_name=last, first_name=first,
                                                                   last_name_kana=kana, date_of_birth=d)
        self.ani, self.imo = mk('山田', '兄', 'やまだ'), mk('山田', '妹', 'やまだ')
        self.ani.siblings.add(self.imo)
        self.o1, self.o2, self.o3 = mk('青木', '一', 'あおき'), mk('井上', '二', 'いのうえ'), mk('上田', '三', 'うえだ')
        self.day = datetime.date(2026, 10, 6)      # 火

    def test_siblings_take_separate_seats_unless_shared(self):
        ra, _ = services.create_reservation(self.f, self.ani, self.day, start_time=10)
        ri, _ = services.create_reservation(self.f, self.imo, self.day, start_time=10)
        services.create_reservation(self.f, self.o1, self.day, start_time=10)
        r2, _ = services.create_reservation(self.f, self.o2, self.day, start_time=10)
        self.assertEqual(r2.status, Reservation.STATUS_WAITLIST)        # まとめない日は 3 枠ぶん
        # まとめると席が空き、キャンセル待ちが繰り上がる
        services.set_share_seat(ri, True)
        r2.refresh_from_db()
        self.assertEqual(r2.status, Reservation.STATUS_CONFIRMED)
        st = services.slot_state(self.f, self.day, 10)
        self.assertEqual((st['confirmed'], st['seats'], st['full']), (4, 3, True))
        # いま分けると 4 枠になってしまうので分けられない
        with self.assertRaises(services.ReservationError):
            services.set_share_seat(ri, False)
        # きょうだいのいない枠ではまとめられない
        r3, _ = services.create_reservation(self.f, self.o3, self.day, start_time=11)
        with self.assertRaises(services.ReservationError):
            services.set_share_seat(r3, True)
        # 新しく入れるときに「まとめる」なら満員の枠でも入る（きょうだいがいるので席は増えない）
        services.cancel_reservation(ra, notify=False)

    def test_create_with_share_seat_and_customer_siblings(self):
        from reservations.models import Customer
        c = Customer.objects.create(facility=self.f, name='青木 母')
        c.children.add(self.o1, self.o2)                                  # 連絡先が同じ子もきょうだい
        self.assertIn(self.o2.pk, services.sibling_map(self.f)[self.o1.pk])
        services.create_reservation(self.f, self.o1, self.day, start_time=10)
        services.create_reservation(self.f, self.o3, self.day, start_time=10)
        services.create_reservation(self.f, self.ani, self.day, start_time=10)
        r2, _ = services.create_reservation(self.f, self.o2, self.day, start_time=10, share_seat=True)
        self.assertEqual(r2.status, Reservation.STATUS_CONFIRMED)

    def test_day_page_buttons_and_schedule_box(self):
        ra, _ = services.create_reservation(self.f, self.ani, self.day, start_time=10)
        ri, _ = services.create_reservation(self.f, self.imo, self.day, start_time=10)
        url = reverse('reservations:day', args=[2026, 10, 6])
        self.assertContains(self.client.get(url), 'きょうだいで1枠にまとめる')
        res = self.client.post(url, {'action': 'share', 'reservation': ri.pk, 'value': '1'}, follow=True)
        self.assertContains(res, 'きょうだいと1枠にまとめました')
        self.assertContains(res, '分ける')
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        self.assertContains(res, 'ms-box grp')
        self.assertContains(res, f'data-sibs="{self.imo.pk}"')
        res = self.client.post(url, {'action': 'share', 'reservation': ra.pk, 'value': '0'}, follow=True)
        self.assertContains(res, '別の枠に分けました')
        ri.refresh_from_db()
        self.assertFalse(ri.share_seat)

    def test_join_by_drop_on_sibling_and_swap_resets(self):
        ra, _ = services.create_reservation(self.f, self.ani, self.day, start_time=10)
        services.create_reservation(self.f, self.o1, self.day, start_time=10)
        services.create_reservation(self.f, self.o2, self.day, start_time=10)          # 10時は満員
        ri, _ = services.create_reservation(self.f, self.imo, self.day, start_time=11)
        swap = reverse('reservations:monthly_schedule_swap', args=[2026, 10])
        res = self.client.post(swap, {'action': 'move', 'a': ri.pk, 'date': '2026-10-06', 'hour': '10'}, follow=True)
        self.assertContains(res, '満員')
        res = self.client.post(swap, {'action': 'join', 'a': ri.pk, 'b': ra.pk}, follow=True)
        self.assertContains(res, 'きょうだいで1枠にまとめました')
        ri.refresh_from_db()
        self.assertEqual((ri.start_time.hour, ri.share_seat, ri.status), (10, True, Reservation.STATUS_CONFIRMED))
        # 兄弟でない相手には「まとめる」はできない
        r1 = Reservation.objects.get(beneficiary=self.o1)
        res = self.client.post(swap, {'action': 'join', 'a': r1.pk, 'b': ra.pk}, follow=True)
        self.assertContains(res, 'きょうだいがいません')
        # 入れ替えで枠が変わると、まとめは外れる
        r3, _ = services.create_reservation(self.f, self.o3, self.day, start_time=11)
        monthly.swap_reservations(ri, r3)
        ri.refresh_from_db()
        self.assertFalse(ri.share_seat)

    def test_sibling_field_on_edit(self):
        res = self.client.post(reverse('beneficiaries:update', args=[self.o1.pk]), {
            'last_name': '青木', 'first_name': '一', 'date_of_birth': '2018-01-01', 'status': 'active', 'gender': 'male',
            'siblings': [self.o2.pk]})
        self.assertEqual(res.status_code, 302)
        self.assertEqual(list(self.o2.siblings.all()), [self.o1])
        self.assertContains(self.client.get(reverse('beneficiaries:detail', args=[self.o1.pk])), 'きょうだい')


class ScheduleListTests(TestCase):
    """月間予定表の右の一覧：希望 0 回で予約の無い子は出さない。合計を出す。ドラッグで動かせる"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.force_login(self.user)
        d = datetime.date(2018, 1, 1)
        self.a = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=d)
        self.zero = Beneficiary.objects.create(facility=self.f, last_name='ゼロ', first_name='子', date_of_birth=d)
        self.dup = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2000, 1, 1))

    def test_zero_and_duplicates_hidden_and_totals(self):
        monthly.save_request(self.f, self.a, 2026, 10, 3, {'2026-10-06': [10], '2026-10-07': [10]})
        monthly.save_request(self.f, self.zero, 2026, 10, 0, {})
        monthly.assign_month(self.f, 2026, 10, self.s, notify=False)
        res = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10]))
        side = res.content.decode().split('id="ms-side"')[1]
        self.assertNotIn('ゼロ 子', side)
        self.assertEqual(side.count('青木 子'), 1)                    # 同じ名前の二重登録（予約なし）は出さない
        self.assertContains(res, '合計 1 名')
        self.assertContains(res, '予約 合計 <strong>2</strong> 件')
        self.assertContains(res, '計2')                               # 週の合計
        self.assertContains(res, 'draggable="true"')
        self.assertContains(res, 'ms-drop-menu')
        pdf = self.client.get(reverse('reservations:monthly_schedule_pdf', args=[2026, 10]) + '?fmt=html')
        self.assertNotContains(pdf, 'ゼロ 子')


class ScheduleSideTableTests(TestCase):
    """右の一覧の数字：確定＝予約の回数（欠席・キャンセルは除く）、あと n＝希望に足りない回数、残＝契約（支給量）−確定"""

    def setUp(self):
        from beneficiaries.models import RecipientCertificate
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.force_login(self.user)
        self.a = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2018, 1, 1))
        RecipientCertificate.objects.create(beneficiary=self.a, granted_days=10, valid_from=datetime.date(2026, 4, 1),
                                            valid_until=datetime.date(2027, 3, 31))

    def test_numbers(self):
        monthly.save_request(self.f, self.a, 2026, 10, 4, {'2026-10-06': [10], '2026-10-07': [10], '2026-10-09': [10]})
        monthly.assign_month(self.f, 2026, 10, self.s, notify=False)        # 3 回しか入らない
        res = Reservation.objects.filter(beneficiary=self.a).order_by('date').first()
        monthly.set_attendance(res, Reservation.ATT_ABSENT)
        row = next(r for r in monthly.request_rows(self.f, 2026, 10, self.s) if r['beneficiary'].pk == self.a.pk)
        self.assertEqual((row['desired'], row['confirmed'], row['used'], row['absent'], row['short'], row['contract_left']),
                         (4, 3, 2, 1, 1, 8))
        totals = monthly.row_totals(monthly.schedule_rows([row]))
        self.assertEqual((totals['used'], totals['short'], totals['contract_left'], totals['granted']), (2, 1, 8, 10))
        page = self.client.get(reverse('reservations:monthly_schedule', args=[2026, 10])).content.decode()
        side = page.split('id="ms-side"')[1]
        self.assertIn('あと1', side)
        self.assertIn('欠席・キャンセル 1 回は数えない', side)
        self.assertIn('契約（支給量）10 − 確定 2', side)
        self.assertIn('ms-fit', page)                      # 全体表示は画面いっぱい（左メニュー・上の帯を消す）
        self.assertIn('requestFullscreen', page)


class AddTodayTests(TestCase):
    """利用者の画面の「今日の予定に追加」：枠を選んできょうの予約を作り、そのまま療育記録へ"""

    def setUp(self):
        self.f, self.s = ryoiku()
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.kids = [child(self.f, n) for n in ('青木', '井上', '上田', '江口', '大野')]
        self.day = datetime.date(2026, 10, 2)    # 金
        patcher = mock.patch('django.utils.timezone.localdate', return_value=self.day)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_detail_shows_slots_and_add_goes_to_therapy(self):
        a, b = self.kids[0], self.kids[1]
        a.cannot_pair.add(b)
        services.create_reservation(self.f, b, self.day, start_time='10')
        detail = reverse('beneficiaries:detail', args=[a.pk])
        res = self.client.get(detail)
        self.assertContains(res, '今日の予定に追加')
        self.assertContains(res, '共演NG：井上 子')                  # 10時は同じ時間にできない子がいる
        self.assertContains(res, 'value="11:00"')
        url = reverse('reservations:add_today', args=[a.pk])
        res = self.client.post(url, {'start_time': '10:00'})
        self.assertRedirects(res, detail)
        self.assertFalse(Reservation.objects.filter(beneficiary=a).exists())
        res = self.client.post(url, {'start_time': '11:00'})
        self.assertRedirects(res, reverse('therapy:child', args=[a.pk]) + '?date=2026-10-02&time=11:00#add', fetch_redirect_response=False)
        r = Reservation.objects.get(beneficiary=a)
        self.assertEqual((r.date, r.start_time, r.status, r.note), (self.day, datetime.time(11), Reservation.STATUS_CONFIRMED, '当日追加'))
        self.assertFalse(ReservationNotice.objects.filter(reservation=r).exists())     # 保護者への通知は出さない
        res = self.client.get(detail)
        self.assertContains(res, 'きょう 11時の予定あり')
        self.assertContains(res, '?date=2026-10-02&time=11:00#add')
        # もう一度押しても二重にしない
        self.client.post(url, {'start_time': '13:00'})
        self.assertEqual(Reservation.objects.filter(beneficiary=a).count(), 1)

    def test_full_slot_goes_to_waitlist_and_closed_day(self):
        for k in self.kids[1:4]:
            services.create_reservation(self.f, k, self.day, start_time='10')
        self.client.post(reverse('reservations:add_today', args=[self.kids[0].pk]), {'start_time': '10:00'})
        self.assertEqual(Reservation.objects.get(beneficiary=self.kids[0]).status, Reservation.STATUS_WAITLIST)
        with mock.patch('django.utils.timezone.localdate', return_value=datetime.date(2026, 10, 5)):   # 月曜はお休み
            res = self.client.get(reverse('beneficiaries:detail', args=[self.kids[4].pk]))
        self.assertContains(res, 'きょうは休業日')
