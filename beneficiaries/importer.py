"""
利用者情報の Excel・CSV 取り込み（ゆあーず）。

- COLUMNS が列の定義（1か所）。見出しの文字（label）で列を探すので、雛形の並びを変えても読める。
  依頼者から項目の一覧が届いたら、ここを差し替える
- 同じ「姓・名・生年月日」の利用者がいれば更新（空の欄は上書きしない）、いなければ新規登録
- 保護者は「保護者 姓・名」が同じなら更新、受給者証は受給者証番号が同じなら更新
- まず「確かめる」（登録しない）で行ごとの結果を見せ、「登録する」で保存する
- もう1つの形「保護者一覧」（前のシステムからのデータ移行用。1行が保護者1人。児童は名前で台帳と照合）にも対応する。
  1行目に「保護者（名前）」があればこの形とみなす（GUARDIAN_COLUMNS）
- もう1つの形「児童一覧」（前のシステムの CSV。1行が契約1件で、同じお子さまが複数行に出る）にも対応する。
  見出しあり（17列：保護者（名前）…児童（名前）…保訪退所日。1行目に「児童（名前）」があればこの形）と、見出しなし（15列）の両方を読む。
  列の意味は CHILDREN_COLUMNS。同じ姓・名・生年月日の行は1人にまとめ、利用中があれば在籍中、なければ退所にする
- すでに台帳にいる利用者は、登録のときに「取り込んだ値で上書きする」か「台帳の値を残す」かを選ぶ（MODE_*）。
  どちらでも、台帳の欄が空なら取り込んだ値を入れる
"""
import datetime
import io
import re

from django.db import transaction

from .models import Beneficiary, Guardian, RecipientCertificate, to_hiragana

# (キー, 見出し, 必須, 書き方の説明, 雛形の例)
COLUMNS = [
    ('last_name', '姓', True, '', '山田'),
    ('first_name', '名', True, '', '太郎'),
    ('last_name_kana', 'せい（ふりがな）', False, 'ひらがな', 'やまだ'),
    ('first_name_kana', 'めい（ふりがな）', False, 'ひらがな', 'たろう'),
    ('date_of_birth', '生年月日', True, '2019-04-01 か 2019/4/1 か 2019年4月1日', '2019-04-01'),
    ('gender', '性別', False, '男・女・その他', '男'),
    ('disability_class', '障害区分', False, '1級・2級（空でもよい）', ''),
    ('disability_type', '障害種別', False, '自由に書く（例：自閉スペクトラム症）', '自閉スペクトラム症'),
    ('is_severe', '重症心身障害児', False, '○ か 空', ''),
    ('postal_code', '郵便番号', False, '6000000（ハイフンは無くてもよい）', '600-8216'),
    ('address', '住所', False, '', '京都市下京区○○町1-1'),
    ('mobile_phone', '携帯電話番号', False, '', '090-0000-0000'),
    ('home_phone', '自宅電話番号', False, '', ''),
    ('school_name', '通学学校名', False, '', '○○小学校'),
    ('grade', '学年', False, '未就学・小1〜小6・中1〜中3・高1〜高3・その他', '小1'),
    ('admission_date', '入所日', False, '日付', '2026-04-01'),
    ('status', '在籍状況', False, '在籍中・退所・卒業（空なら変えない。新しい人は在籍中）', '在籍中'),
    ('discharge_date', '退所日', False, '日付（退所のとき）', ''),
    ('weekdays', '利用予定曜日', False, '月・水・金 のように「・」か「,」で区切る', '月・水'),
    ('notes', '備考', False, '', ''),
    ('guardian_last_name', '保護者 姓', False, '', '山田'),
    ('guardian_first_name', '保護者 名', False, '', '花子'),
    ('guardian_kana', '保護者 ふりがな', False, '', 'やまだ はなこ'),
    ('guardian_relation', '保護者 続柄', False, '父・母・その他', '母'),
    ('guardian_phone', '保護者 電話番号', False, '', '090-0000-0001'),
    ('guardian_email', '保護者 メール', False, '', ''),
    ('certificate_number', '受給者証番号', False, '', '2600001234'),
    ('granted_days', '支給量（日/月）', False, '数字', '10'),
    ('monthly_cap', '負担上限月額（円）', False, '数字', '4600'),
    ('valid_from', '受給者証 有効期間（開始）', False, '日付', '2026-04-01'),
    ('valid_until', '受給者証 有効期間（終了）', False, '日付', '2027-03-31'),
    ('municipality', '受給者証 市区町村', False, '', '京都市'),
    ('support_office', '相談支援事業所', False, '', ''),
]
HEADERS = [c[1] for c in COLUMNS]
LABEL_TO_KEY = {c[1]: c[0] for c in COLUMNS}
LABEL_TO_KEY.update({c[0]: c[0] for c in COLUMNS})     # 英語のキーで書いた見出しも受け付ける
REQUIRED = [c[0] for c in COLUMNS if c[2]]
MAX_ROWS = 500

GENDER = {'男': 'male', '男性': 'male', '女': 'female', '女性': 'female', 'その他': 'other',
          'male': 'male', 'female': 'female', 'other': 'other'}
RELATION = {'父': 'father', '母': 'mother', 'その他': 'other', 'father': 'father', 'mother': 'mother', 'other': 'other'}
STATUS = {'在籍中': 'active', '在籍': 'active', '利用中': 'active', '退所': 'inactive', '退所済': 'inactive', '退所済み': 'inactive',
          '卒業': 'graduated', '卒業済': 'graduated', '卒業済み': 'graduated',
          'active': 'active', 'inactive': 'inactive', 'graduated': 'graduated'}
DISABILITY = {'1級': '1', '2級': '2', '1': '1', '2': '2', '１級': '1', '２級': '2'}
GRADE = {label: key for key, label in Beneficiary.GRADE_CHOICES if key}
WEEKDAY_FIELDS = {'月': 'weekday_mon', '火': 'weekday_tue', '水': 'weekday_wed', '木': 'weekday_thu', '金': 'weekday_fri', '土': 'weekday_sat'}
TRUE_WORDS = {'○', '◯', '〇', 'はい', 'あり', '有', 'yes', 'true', '1', 'y'}


# 保護者一覧の形：(見出し, 同じ見出しの何番目か, キー)。前のシステムの書き出しをそのまま読む
GUARDIAN_COLUMNS = [
    ('保護者（名前）', 0, 'g_name'), ('保護者（カナ）', 0, 'g_kana'), ('続柄', 0, 'relation'), ('児童', 0, 'children'),
    ('郵便番号', 0, 'postal'), ('都道府県', 0, 'pref'), ('市区町村', 0, 'city'), ('番地', 0, 'street'), ('ビル・マンション名', 0, 'building'),
    ('備考', 0, 'note'),
    ('連絡先', 0, 'c1_label'), ('電話番号1', 0, 'c1_phone1'), ('電話番号2', 0, 'c1_phone2'), ('その他（メールアドレス等）', 0, 'c1_other'),
    ('連絡先', 1, 'c2_label'), ('電話番号1', 1, 'c2_phone1'), ('電話番号2', 1, 'c2_phone2'), ('その他（メールアドレス等）', 1, 'c2_other'),
    ('支払い方法', 0, 'pay_method'), ('請求先区分', 0, 'bill_type'), ('宛名', 0, 'bill_name'),
    ('郵便番号', 1, 'bill_postal'), ('都道府県', 1, 'bill_pref'), ('市区町村', 1, 'bill_city'), ('番地', 1, 'bill_street'),
    ('ビル・マンション名', 1, 'bill_building'), ('電話番号1', 2, 'bill_phone1'), ('電話番号2', 2, 'bill_phone2'), ('備考', 1, 'bill_note'),
]
GUARDIAN_HEADERS = [c[0] for c in GUARDIAN_COLUMNS]
GUARDIAN_MARK = '保護者（名前）'
FORMAT_BENEFICIARY = 'beneficiary'
FORMAT_GUARDIAN = 'guardian'
FORMAT_CHILDREN = 'children'
MODE_OVERWRITE = 'overwrite'     # 台帳と違う値は、取り込んだ値で上書きする
MODE_KEEP = 'keep'               # 台帳の値を残す（台帳の欄が空のときだけ取り込んだ値を入れる）
MODES = (MODE_OVERWRITE, MODE_KEEP)
# 前のシステムの「児童一覧」CSV（1行が契約1件）。見出しの文字と列の並び（見出しなしの CSV はこの並びで読む）
CHILDREN_COLUMNS = ['保護者（名前）', '保護者（カナ）', '管理番号', '児童（名前）', '児童（カナ）', '性別', '生年月日', '備考',
                    '児発状態', '児発利用契約日', '児発退所日', '放デイ状態', '放デイ利用契約日', '放デイ退所日',
                    '保訪状態', '保訪利用契約日', '保訪退所日']
CHILDREN_WIDTH = len(CHILDREN_COLUMNS)
CHILDREN_MARK = '児童（名前）'
CHILDREN_SERVICES = ((8, '児童発達支援'), (11, '放課後等デイ'), (14, '保育所等訪問支援'))   # (状態の列, 名前)。契約日・退所日はその右の2列
CHILDREN_STATUS_WORDS = {'利用なし', '退所', '利用中'}
CHILDREN_ACTIVE_WORDS = {'利用中', '契約中', '在籍', '在籍中', '利用'}
CHILDREN_NONE_WORDS = {'', '利用なし', 'なし', '-', '－', '—'}
PLACEHOLDER_DOB = datetime.date(2000, 1, 1)     # 児童の生年月日が無いときの仮の値（あとで直してもらう）
RELATION_WORDS = {'父': 'father', '母': 'mother', '父親': 'father', '母親': 'mother', 'お父さん': 'father', 'お母さん': 'mother'}


class RowError(Exception):
    pass


def template_xlsx():
    """雛形（1枚目：見出しと例、2枚目：書き方）"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = '利用者'
    ws.append(HEADERS)
    ws.append([c[4] for c in COLUMNS])
    for i, c in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='FFF3E0' if c[2] else 'EEEEEE')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(28, len(c[1]) * 2 + 4))
    ws.freeze_panes = 'A2'
    ws2 = wb.create_sheet('書き方')
    ws2.append(['見出し', '必須', '書き方'])
    for c in COLUMNS:
        ws2.append([c[1], '必須' if c[2] else '', c[3]])
    ws2.append([])
    ws2.append(['1枚目の2行目は例です。消して、2行目から利用者を1人1行で入れてください。'])
    ws2.append(['同じ「姓・名・生年月日」の利用者がすでにいれば、登録のときに「取り込んだ値で上書きする」か「台帳の値を残す」かを選べます（台帳の欄が空なら、どちらでも取り込んだ値を入れます。取り込む側の空の欄は何も変えません）。'])
    ws2.append(['見出しの並びは変えても構いません。使わない列は消しても構いません（姓・名・生年月日は必要）。'])
    for col, w in (('A', 30), ('B', 8), ('C', 60)):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rows_from_file(uploaded):
    """xlsx / csv / タブ区切り → 見出し→値 の dict のリスト（見出しの行は自動で探す。表題行があっても読める。空行は飛ばす）"""
    from config.tabular import find_header, read_rows
    rows = read_rows(uploaded)
    if not rows:
        return []
    head = find_header(rows, [CHILDREN_MARK, '生年月日'])
    if head is not None:
        return _children_rows(_children_by_header(rows[head], rows[head + 1:head + 1 + MAX_ROWS]), head + 1, strict=False)
    head = find_header(rows, [GUARDIAN_MARK])
    if head is not None:
        return _guardian_rows([(h or '').strip() for h in rows[head]], rows[head + 1:head + 1 + MAX_ROWS])
    start = _children_start(rows)
    if start is not None:
        return _children_rows(rows[start:start + MAX_ROWS], start)
    head = find_header(rows, ['姓', '名', '生年月日'])
    if head is None:
        head = find_header(rows, ['last_name', 'first_name', 'date_of_birth'])
    if head is None:
        raise ValueError('1行目の見出しに「姓」「名」「生年月日」が見つかりません。雛形の見出しを使ってください'
                         '（前のシステムの「保護者一覧」なら「保護者（名前）」の見出しがあれば読めます）。')
    header = [(h or '').strip() for h in rows[head]]
    keys = [LABEL_TO_KEY.get(h) for h in header]
    out = []
    for r in rows[head + 1:head + 1 + MAX_ROWS]:
        d = {}
        for i, k in enumerate(keys):
            if k:
                d[k] = (r[i] if i < len(r) else '').strip()
        if any(d.values()):
            out.append(d)
    return out


def _guardian_rows(header, body):
    """保護者一覧（同じ見出しが何度も出る）を、何番目かで見分けてキーに割り当てる"""
    seen = {}
    keys = []
    for h in header:
        n = seen.get(h, 0)
        seen[h] = n + 1
        keys.append(next((k for label, idx, k in GUARDIAN_COLUMNS if label == h and idx == n), None))
    out = []
    for r in body:
        d = {'_format': FORMAT_GUARDIAN}
        for i, k in enumerate(keys):
            if k:
                d[k] = (r[i] if i < len(r) else '').strip()
        if any(v for k, v in d.items() if k != '_format'):
            out.append(d)
    return out


def _children_by_header(header, body):
    """見出しのある児童一覧 → CHILDREN_COLUMNS の並びにそろえた行（見出しの並びが変わっても読める）"""
    pos = {}
    for i, h in enumerate(header):
        h = (h or '').strip()
        if h in CHILDREN_COLUMNS and h not in pos:
            pos[h] = i
    return [[(r[pos[label]] if label in pos and pos[label] < len(r) else '') for label in CHILDREN_COLUMNS] for r in body]


def _is_children_row(r):
    """見出しなしの児童一覧の1行か：15列以上、性別が男/女、利用状況の語、管理番号が数字、生年月日が日付"""
    if len(r) < 15:
        return False
    c = [(x or '').strip() for x in r]
    try:
        return (c[5] in GENDER and c[8] in CHILDREN_STATUS_WORDS and c[11] in CHILDREN_STATUS_WORDS
                and re.fullmatch(r'\d{3,}', c[2]) is not None and parse_date(c[6]) is not None)
    except RowError:
        return False


def _children_start(rows):
    """先頭 3 行のうち、児童一覧の行として読める最初の行番号（見出し行が足されていても飛ばす）。無ければ None"""
    for i, r in enumerate(rows[:3]):
        if _is_children_row(r):
            return i
    return None


def _hira(s):
    """カタカナ → ひらがな（台帳のふりがなはひらがな）"""
    return to_hiragana(s)


def _fmt_date(v):
    try:
        d = parse_date(v)
    except RowError:
        return v
    return f'{d.year}/{d.month}/{d.day}' if d else ''


def _children_rows(body, offset, strict=True):
    """児童一覧（1行が契約1件）→ 同じお子さまを1人にまとめて、雛形と同じキーの dict にする。
    strict は見出しなしの CSV（並びを決め打ちで読むので、行の形を確かめる）"""
    groups = {}
    order = []
    for n, r in enumerate(body):
        c = ([(x or '').strip() for x in r] + [''] * CHILDREN_WIDTH)[:CHILDREN_WIDTH]
        if not any(c):
            continue
        if strict and not _is_children_row(c):
            key = ('?', n)
            groups[key] = {'_format': FORMAT_CHILDREN, '_line': offset + n + 1, '_error': '児童一覧の行として読めません（列の並びを確かめてください）',
                           'last_name': c[3], 'first_name': ''}
            order.append(key)
            continue
        last, first = split_name(c[3])
        key = (last, first, c[6])
        regs = []
        for col, name in CHILDREN_SERVICES:
            st = c[col]
            start, end = c[col + 1], c[col + 2]
            if st not in CHILDREN_NONE_WORDS or start or end:
                regs.append({'service': name, 'status': st or '（状態なし）', 'start': start, 'end': end, 'number': c[2],
                             'active': st in CHILDREN_ACTIVE_WORDS or (not st and start and not end)})
        if not regs:
            regs.append({'service': '', 'status': '利用なし', 'start': '', 'end': '', 'number': c[2], 'active': False})
        g = groups.get(key)
        if g is None:
            g = {'_format': FORMAT_CHILDREN, '_line': offset + n + 1, '_regs': [], '_notes': [],
                 'last_name': last, 'first_name': first, 'date_of_birth': c[6], 'gender': c[5]}
            groups[key] = g
            order.append(key)
        kl, kf = split_name(c[4])
        if kl or kf:
            g['last_name_kana'], g['first_name_kana'] = _hira(kl), _hira(kf)
        if c[0]:
            gl, gf = split_name(c[0])
            g['guardian_last_name'], g['guardian_first_name'], g['guardian_kana'] = gl, gf, _hira(c[1])
        if c[7] and c[7] not in g['_notes']:
            g['_notes'].append(c[7])
        g['_regs'].extend(regs)
    by_first = {}
    for key in order:
        if key[0] != '?':
            by_first.setdefault((key[1], key[2]), []).append(key)
    out = []
    for key in order:
        g = groups[key]
        if '_error' in g:
            out.append(g)
            continue
        regs = g.pop('_regs')
        notes = g.pop('_notes')
        others = [k for k in by_first.get((key[1], key[2]), []) if k != key]
        actives = [x for x in regs if x['active']]
        starts = sorted((parse_date_or_none(x['start']), x['start']) for x in regs if parse_date_or_none(x['start']))
        ends = sorted((parse_date_or_none(x['end']), x['end']) for x in regs if parse_date_or_none(x['end']))
        g['status'] = '在籍中' if actives else '退所'
        g['admission_date'] = starts[0][1] if starts else ''
        g['discharge_date'] = '' if actives else (ends[-1][1] if ends else '')
        g['notes_append'] = '\n'.join(
            [f"前のシステム：管理番号 {x['number']}" + (f" {x['service']}" if x['service'] else '') + f" {x['status']}"
             + (f" {_fmt_date(x['start'])}〜{_fmt_date(x['end'])}" if x['start'] or x['end'] else '') for x in regs]
            + [f'備考：{x}' for x in notes])
        g['_detail'] = '・'.join(
            (f"{x['service']} " if x['service'] else '') + x['status']
            + (f"（{_fmt_date(x['start'])}〜{_fmt_date(x['end'])}）" if x['start'] or x['end'] else '') for x in regs)
        if others:
            g['_detail'] += '。※ 姓だけ違う同じ名・生年月日の行（' + '・'.join(f'{k[0]} {k[1]}' for k in others) + \
                            '）もあります。姓が変わったのなら、登録後にどちらかを消してください'
        out.append(g)
    return out


def format_of(rows):
    if not rows:
        return FORMAT_BENEFICIARY
    return rows[0].get('_format') if rows[0].get('_format') in (FORMAT_GUARDIAN, FORMAT_CHILDREN) else FORMAT_BENEFICIARY


def parse_date(v):
    v = (v or '').strip()
    if not v:
        return None
    v = v.translate(str.maketrans('０１２３４５６７８９／．－', '0123456789/.-'))
    m = re.match(r'^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?', v)
    if not m:
        raise RowError(f'日付を読み取れません：{v}（例：2019-04-01）')
    try:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        raise RowError(f'日付が正しくありません：{v}')


def parse_date_or_none(v):
    """日付として読めなければ None（並べ替え用）"""
    try:
        return parse_date(v)
    except RowError:
        return None


def _int(v, label):
    v = (v or '').strip().translate(str.maketrans('０１２３４５６７８９', '0123456789')).replace(',', '').replace('円', '').replace('日', '')
    if not v:
        return None
    if not v.isdigit():
        raise RowError(f'{label}は数字で書いてください：{v}')
    return int(v)


def _choice(v, table, label):
    v = (v or '').strip()
    if not v:
        return ''
    if v not in table:
        raise RowError(f'{label}の書き方が違います：{v}（{"・".join(k for k in table if not k.isascii())}）')
    return table[v]


def _weekdays(v):
    v = (v or '').strip()
    if not v:
        return None
    on = set()
    for ch in re.split(r'[・,、/／\s]+', v.replace('曜日', '').replace('曜', '')):
        if not ch:
            continue
        if ch not in WEEKDAY_FIELDS:
            raise RowError(f'利用予定曜日の書き方が違います：{v}（月・火・水・木・金・土 を「・」で区切る）')
        on.add(WEEKDAY_FIELDS[ch])
    return on


def parse_row(r):
    """1行を検査して、登録に使う値にする。問題があれば RowError"""
    for k in REQUIRED:
        if not r.get(k):
            raise RowError(f'{dict((c[0], c[1]) for c in COLUMNS)[k]}が空です')
    out = {
        'last_name': r['last_name'][:50], 'first_name': r['first_name'][:50],
        'last_name_kana': r.get('last_name_kana', '')[:50], 'first_name_kana': r.get('first_name_kana', '')[:50],
        'date_of_birth': parse_date(r['date_of_birth']),
        'gender': _choice(r.get('gender'), GENDER, '性別'),
        'disability_class': _choice(r.get('disability_class'), DISABILITY, '障害区分'),
        'disability_type': r.get('disability_type', '')[:100],
        'is_severe': (r.get('is_severe', '').strip().lower() in TRUE_WORDS) if r.get('is_severe', '').strip() else None,
        'postal_code': r.get('postal_code', '').replace('-', '').replace('ー', '')[:8],
        'address': r.get('address', '')[:200],
        'mobile_phone': r.get('mobile_phone', '')[:20], 'home_phone': r.get('home_phone', '')[:20],
        'school_name': r.get('school_name', '')[:100],
        'grade': _choice(r.get('grade'), GRADE, '学年'),
        'admission_date': parse_date(r.get('admission_date')),
        'status': _choice(r.get('status'), STATUS, '在籍状況'),
        'discharge_date': parse_date(r.get('discharge_date')),
        'weekdays': _weekdays(r.get('weekdays')),
        'notes': r.get('notes', ''),
        'notes_append': r.get('notes_append', ''),
        'guardian': None, 'certificate': None,
    }
    gl, gf = r.get('guardian_last_name', '').strip(), r.get('guardian_first_name', '').strip()
    if gl or gf:
        out['guardian'] = {
            'last_name': (gl or out['last_name'])[:50], 'first_name': (gf or '保護者')[:50],
            'relation': _choice(r.get('guardian_relation'), RELATION, '保護者 続柄') or 'other',
            'kana': r.get('guardian_kana', '')[:100],
            'phone': r.get('guardian_phone', '')[:20], 'email': r.get('guardian_email', '')[:254],
        }
    cert_no = r.get('certificate_number', '').strip()
    valid_from, valid_until = parse_date(r.get('valid_from')), parse_date(r.get('valid_until'))
    if cert_no or valid_until or valid_from:
        if not valid_until:
            raise RowError('受給者証を入れるときは「受給者証 有効期間（終了）」が必要です')
        out['certificate'] = {
            'certificate_number': cert_no[:20], 'granted_days': _int(r.get('granted_days'), '支給量') or 0,
            'monthly_cap': _int(r.get('monthly_cap'), '負担上限月額') or 0,
            'valid_from': valid_from or datetime.date.today(), 'valid_until': valid_until,
            'municipality': r.get('municipality', '')[:50], 'support_office': r.get('support_office', '')[:100],
        }
    return out


def split_name(name):
    """「山田 太郎」「山田　太郎」→ (山田, 太郎)。区切りが無ければ (全体, '')"""
    parts = re.split(r'[\s\u3000]+', (name or '').strip(), maxsplit=1)
    return (parts[0], parts[1] if len(parts) > 1 else '')


def _norm(name):
    return re.sub(r'[\s\u3000]+', '', name or '')


def _join_address(r, pre=''):
    return ''.join((r.get(pre + k, '') for k in ('pref', 'city', 'street', 'building')))[:200]


def parse_guardian_row(r):
    """保護者一覧の1行 → 保護者の値・児童の名前・児童に写す住所。問題があれば RowError"""
    name = r.get('g_name', '').strip()
    if not name:
        raise RowError('保護者（名前）が空です')
    children = [c.strip() for c in re.split(r'[、,，／/;；\n]+', r.get('children', '')) if c.strip()]
    if not children:
        raise RowError('児童が空です（誰の保護者か分からないので登録できません）')
    last, first = split_name(name)
    rel = r.get('relation', '').strip()
    relation = RELATION_WORDS.get(rel) or RELATION.get(rel) or 'other'
    phones = [p for p in (r.get('c1_phone1'), r.get('c1_phone2'), r.get('c2_phone1'), r.get('c2_phone2')) if p]
    others = [o for o in (r.get('c1_other'), r.get('c2_other')) if o]
    email = next((o for o in others if '@' in o), '')
    extra = {}
    for n, pre in ((1, 'c1_'), (2, 'c2_')):
        label = r.get(pre + 'label', '')
        block = {'名称': label, '電話番号1': r.get(pre + 'phone1', ''), '電話番号2': r.get(pre + 'phone2', ''), 'その他': r.get(pre + 'other', '')}
        for k, v in block.items():
            if v and not (k == 'その他' and v == email):
                extra[f'連絡先{n} {k}'] = v
    if rel and relation == 'other' and rel not in ('その他',):
        extra['続柄（原文）'] = rel
    bill_addr = _join_address(r, 'bill_')
    for k, v in (('支払い方法', r.get('pay_method', '')), ('請求先区分', r.get('bill_type', '')), ('請求先 宛名', r.get('bill_name', '')),
                 ('請求先 郵便番号', r.get('bill_postal', '')), ('請求先 住所', bill_addr),
                 ('請求先 電話番号1', r.get('bill_phone1', '')), ('請求先 電話番号2', r.get('bill_phone2', '')), ('請求先 備考', r.get('bill_note', ''))):
        if v:
            extra[k] = v
    return {
        'children': children,
        'guardian': {'last_name': last[:50], 'first_name': first[:50], 'kana': r.get('g_kana', '')[:100], 'relation': relation,
                     'phone': (phones[0] if phones else '')[:20], 'phone2': (phones[1] if len(phones) > 1 else '')[:20],
                     'email': email[:254], 'memo': r.get('note', '')[:200], 'extra': extra},
        'postal_code': r.get('postal', '').replace('-', '').replace('ー', '')[:8], 'address': _join_address(r),
    }


def find_child_by_name(facility, name):
    key = _norm(name)
    for b in Beneficiary.objects.filter(facility=facility).order_by('pk'):
        if _norm(b.last_name + b.first_name) == key:
            return b
    return None


def plan_guardians(facility, rows, create_children=False):
    """保護者一覧の行ごとの見込み。児童は名前で台帳と照合し、無ければエラー（create_children なら仮の生年月日で作る）"""
    out = []
    for i, r in enumerate(rows, 2):
        name = r.get('g_name', '').strip() or '（名前なし）'
        try:
            data = parse_guardian_row(r)
        except RowError as e:
            out.append({'line': i, 'name': name, 'action': 'error', 'detail': str(e), 'data': None})
            continue
        found, missing = [], []
        for child in data['children']:
            b = find_child_by_name(facility, child)
            (found if b else missing).append((child, b.pk if b else None))
        if missing and not create_children:
            out.append({'line': i, 'name': name, 'action': 'error', 'data': None,
                        'detail': '台帳にいない児童：' + '・'.join(c for c, _ in missing) + '（先に児童を登録するか、「台帳にいない児童は仮の生年月日で作る」を付けてください）'})
            continue
        parts = []
        if found:
            parts.append('保護者を ' + '・'.join(c for c, _ in found) + ' さんに付けます（同じ名前の保護者がいれば書き換え）')
        if missing:
            parts.append('児童 ' + '・'.join(c for c, _ in missing) + ' さんを仮の生年月日 2000-01-01 で新しく作ります（あとで直してください）')
        data['found'], data['missing'] = found, missing
        out.append({'line': i, 'name': name, 'action': 'update' if found and not missing else 'create',
                    'detail': '。'.join(parts), 'data': data})
    return out


@transaction.atomic
def apply_guardians(facility, planned, mode=MODE_OVERWRITE):
    """plan_guardians() の結果を保存する。mode は台帳にすでに値がある欄の扱い。戻り値 (新しく作った児童, 付けた保護者の数, エラー数)"""
    children_made = guardians = errors = 0
    for item in planned:
        data = item['data']
        if item['action'] == 'error' or data is None:
            errors += 1
            continue
        targets = [Beneficiary.objects.filter(pk=pk, facility=facility).first() for _, pk in data['found']]
        for child, _ in data['missing']:
            last, first = split_name(child)
            b = find_child_by_name(facility, child)     # 同じ取り込みの前の行で作った場合
            if b is None:
                b = Beneficiary.objects.create(facility=facility, last_name=last[:50], first_name=first[:50], date_of_birth=PLACEHOLDER_DOB,
                                               has_prior_records=True, notes='生年月日は取り込み時の仮の値（2000-01-01）です。正しい日付に直してください。')
                children_made += 1
            targets.append(b)
        g = data['guardian']
        for b in targets:
            if b is None:
                continue
            changed = False
            if data['postal_code'] and not b.postal_code:
                b.postal_code, changed = data['postal_code'], True
            if data['address'] and not b.address:
                b.address, changed = data['address'], True
            if changed:
                b.save(update_fields=['postal_code', 'address', 'updated_at'])
            existing = b.guardians.filter(last_name=g['last_name'], first_name=g['first_name']).first()
            fields = {k: v for k, v in g.items() if k in ('kana', 'relation', 'phone', 'phone2', 'email', 'memo') and v}
            if existing is None:
                Guardian.objects.create(beneficiary=b, last_name=g['last_name'], first_name=g['first_name'], extra=g['extra'],
                                        is_primary=not b.guardians.filter(is_primary=True).exists(), **fields)
            else:
                for k, v in fields.items():
                    if k == 'relation':
                        if v != Guardian.RELATION_OTHER and (mode == MODE_OVERWRITE or existing.relation == Guardian.RELATION_OTHER):
                            existing.relation = v
                    elif _take(getattr(existing, k), v, mode):
                        setattr(existing, k, v)
                existing.extra = ({**(existing.extra or {}), **g['extra']} if mode == MODE_OVERWRITE
                                  else {**g['extra'], **{k: v for k, v in (existing.extra or {}).items() if v not in ('', None)}})
                existing.save()
            guardians += 1
    return children_made, guardians, errors


def find_existing(facility, data):
    return Beneficiary.objects.filter(facility=facility, last_name=data['last_name'], first_name=data['first_name'],
                                      date_of_birth=data['date_of_birth']).order_by('pk').first()


FIELD_LABELS = {c[0]: c[1] for c in COLUMNS}
FIELD_LABELS.update({'last_name_kana': 'ふりがな（姓）', 'first_name_kana': 'ふりがな（名）', 'weekdays': '利用予定曜日', 'is_severe': '重症心身障害児'})


def _empty(v):
    return v in ('', None, False)


def _show(b, k, v):
    """見込みの画面に出す値（選択肢は表示名に、日付は 2026/4/1 に）"""
    if v in ('', None):
        return '（空）'
    if isinstance(v, datetime.date):
        return f'{v.year}/{v.month}/{v.day}'
    if isinstance(v, bool):
        return '○' if v else '（空）'
    if isinstance(v, set):
        return '・'.join(d for d, f in WEEKDAY_FIELDS.items() if f in v) or '（空）'
    field = Beneficiary._meta.get_field(k) if k in SIMPLE_FIELDS else None
    if field is not None and field.choices:
        return dict(field.choices).get(v, v)
    text = str(v)
    return text if len(text) <= 30 else text[:30] + '…'


def compare(b, data):
    """台帳の利用者 b と取り込む値 data を比べる → (空欄に入れる欄, 台帳と違う欄 [(見出し, 台帳, 取り込み)])"""
    fills, diffs = [], []
    current = {k: getattr(b, k) for k in SIMPLE_FIELDS}
    current['weekdays'] = {f for f in WEEKDAY_FIELDS.values() if getattr(b, f)}
    current['is_severe'] = b.is_severe
    for k in SIMPLE_FIELDS + ('weekdays', 'is_severe'):
        new = data.get(k)
        if k == 'notes' or new in ('', None):
            continue
        old = current[k]
        if _empty(old) or (k == 'weekdays' and not old):
            fills.append(FIELD_LABELS.get(k, k))
        elif old != new:
            diffs.append((FIELD_LABELS.get(k, k), _show(b, k, old), _show(b, k, new)))
    new_notes = (data.get('notes') or '').strip()
    if new_notes:
        if not (b.notes or '').strip():
            fills.append('備考')
        elif new_notes != (b.notes or '').strip():
            diffs.append(('備考', _show(b, 'notes', b.notes), _show(b, 'notes', new_notes)))
    return fills, diffs


def find_same_name(facility, data):
    """
    姓・名は同じで生年月日だけ違う台帳の子（仮の生年月日 2000-01-01 の子を先に）。生年月日を書き換えて同じ子として扱う候補。
    無ければ None
    """
    key = _norm(data['last_name'] + data['first_name'])
    found = [b for b in Beneficiary.objects.filter(facility=facility, last_name__in=[data['last_name'], data['last_name'].strip()])
             .exclude(date_of_birth=data['date_of_birth']).order_by('pk') if _norm(b.last_name + b.first_name) == key]
    found.sort(key=lambda b: (b.date_of_birth != PLACEHOLDER_DOB, b.pk))
    return found[0] if found else None


def plan(facility, rows):
    """行ごとの見込み：[{'line', 'name', 'action': 'create'|'update'|'error', 'detail', 'data', 'fills', 'diffs'}]"""
    out = []
    for i, r in enumerate(rows, 2):
        name = f'{r.get("last_name", "")} {r.get("first_name", "")}'.strip() or '（名前なし）'
        line = r.get('_line', i)
        if r.get('_error'):
            out.append({'line': line, 'name': name, 'action': 'error', 'detail': r['_error'], 'data': None})
            continue
        try:
            data = parse_row(r)
        except RowError as e:
            out.append({'line': line, 'name': name, 'action': 'error', 'detail': str(e), 'data': None})
            continue
        existing = find_existing(facility, data)
        extras = []
        if data['guardian']:
            extras.append('保護者')
        if data['certificate']:
            extras.append('受給者証')
        fills, diffs = compare(existing, data) if existing else ([], [])
        if existing:
            detail = 'すでに台帳にいます' + ('（台帳と違う値があります。下の「台帳と違う値」の扱いに従います）' if diffs else '')
        else:
            detail = '新しく登録します'
        detail += '（' + '・'.join(extras) + ' も）' if extras else ''
        if r.get('_detail'):
            detail += '。' + r['_detail']
        dob_match = None
        if not existing:
            other = find_same_name(facility, data)
            if other:
                dob_match = {'pk': other.pk, 'old': other.date_of_birth, 'placeholder': other.date_of_birth == PLACEHOLDER_DOB}
                detail += (f'。※ 台帳に同じ名前の子がいます（生年月日 {other.date_of_birth:%Y/%m/%d}'
                           + ('・取り込みのときの仮の日付' if dob_match['placeholder'] else '') + '）。同じ子なら下の印で生年月日を書き換えます')
        if not existing:
            same = Beneficiary.objects.filter(facility=facility, first_name=data['first_name'], date_of_birth=data['date_of_birth']) \
                .exclude(last_name=data['last_name']).first()
            if same:
                detail += f'。※ 姓だけ違う同じ名・生年月日の利用者（{same.full_name}）がいます。姓が変わったのなら、登録後にどちらかを消してください'
        out.append({'line': line, 'name': name, 'action': 'update' if existing else 'create', 'detail': detail, 'data': data,
                    'existing_pk': existing.pk if existing else None, 'fills': fills, 'diffs': diffs, 'dob_match': dob_match})
    return out


SIMPLE_FIELDS = ('last_name_kana', 'first_name_kana', 'gender', 'disability_class', 'disability_type', 'postal_code', 'address',
                 'mobile_phone', 'home_phone', 'school_name', 'grade', 'admission_date', 'status', 'discharge_date', 'notes')


def _take(old, new, mode):
    """取り込んだ値 new を使うか：空なら使わない。台帳が空なら使う。台帳に値があれば上書きのときだけ"""
    if new in ('', None):
        return False
    return mode == MODE_OVERWRITE or _empty(old)


@transaction.atomic
def apply(facility, planned, mode=MODE_OVERWRITE, dob_lines=()):
    """
    plan() の結果を保存する。mode は台帳にすでに値がある欄の扱い（MODE_OVERWRITE／MODE_KEEP）。
    dob_lines は「同じ子として生年月日を書き換える」に印を付けた行の番号。戻り値 (新規, 更新, エラー数)
    """
    created = updated = errors = 0
    dob_lines = {str(x) for x in dob_lines}
    for item in planned:
        data = item['data']
        if item['action'] == 'error' or data is None:
            errors += 1
            continue
        if item.get('dob_match') and str(item['line']) in dob_lines:
            b = Beneficiary.objects.filter(pk=item['dob_match']['pk'], facility=facility).first()
            if b is not None:
                b.date_of_birth = data['date_of_birth']      # 生年月日は印を付けたとおりに書き換える（残す／上書きの選択に関わらず）
                if b.notes and '生年月日は取り込み時の仮の値' in b.notes:
                    b.notes = '\n'.join(ln for ln in b.notes.splitlines() if '生年月日は取り込み時の仮の値' not in ln)
                b.save(update_fields=['date_of_birth', 'notes', 'updated_at'])
                item = dict(item, existing_pk=b.pk)
        b = Beneficiary.objects.filter(pk=item.get('existing_pk'), facility=facility).first() if item.get('existing_pk') else None
        if b is None:
            b = Beneficiary(facility=facility, last_name=data['last_name'], first_name=data['first_name'],
                            date_of_birth=data['date_of_birth'], has_prior_records=True, gender='', status='')
            is_new = True
        else:
            is_new = False
        m = MODE_OVERWRITE if is_new else mode
        status_taken = _take(b.status, data.get('status'), m)
        for k in SIMPLE_FIELDS:
            v = data.get(k)
            if _take(getattr(b, k), v, m):
                setattr(b, k, v)
        if not b.gender:
            b.gender = Beneficiary.GENDER_MALE
        if not b.status:
            b.status = Beneficiary.STATUS_ACTIVE
        if status_taken and data.get('status') == Beneficiary.STATUS_ACTIVE and not data.get('discharge_date'):
            b.discharge_date = None      # 在籍中に戻すときは退所日を消す
        add = (data.get('notes_append') or '').strip()
        if add:
            lines = [ln for ln in add.split('\n') if ln and ln not in (b.notes or '')]
            if lines:
                b.notes = ((b.notes or '').rstrip() + '\n' if b.notes else '') + '\n'.join(lines)
        if data.get('is_severe') is not None and (m == MODE_OVERWRITE or not b.is_severe):
            b.is_severe = data['is_severe']
        has_days = any(getattr(b, f) for f in WEEKDAY_FIELDS.values())
        if data.get('weekdays') is not None and (m == MODE_OVERWRITE or not has_days):
            for f in WEEKDAY_FIELDS.values():
                setattr(b, f, f in data['weekdays'])
        b.save()
        g = data.get('guardian')
        if g:
            existing = b.guardians.filter(last_name=g['last_name'], first_name=g['first_name']).first()
            if existing is None:
                Guardian.objects.create(beneficiary=b, last_name=g['last_name'], first_name=g['first_name'],
                                        relation=g['relation'], phone=g['phone'], email=g['email'], kana=g.get('kana', ''),
                                        is_primary=not b.guardians.filter(is_primary=True).exists())
            else:
                for k in ('phone', 'email', 'kana'):
                    if _take(getattr(existing, k), g.get(k), m):
                        setattr(existing, k, g[k])
                # 続柄：「その他」は空とみなす（続柄の無い行で、台帳の父・母を変えない）
                if g['relation'] != Guardian.RELATION_OTHER and (m == MODE_OVERWRITE or existing.relation == Guardian.RELATION_OTHER):
                    existing.relation = g['relation']
                existing.save()
        c = data.get('certificate')
        if c:
            qs = b.recipient_certificates.filter(certificate_number=c['certificate_number']) if c['certificate_number'] else \
                b.recipient_certificates.filter(valid_until=c['valid_until'])
            cert = qs.order_by('-valid_until').first()
            if cert is None:
                RecipientCertificate.objects.create(beneficiary=b, **c)
            else:
                for k, v in c.items():
                    if v in ('', None, 0) and k not in ('valid_from', 'valid_until'):
                        continue
                    if m == MODE_OVERWRITE or _empty(getattr(cert, k)) or getattr(cert, k) == 0:
                        setattr(cert, k, v)
                cert.save()
        if is_new:
            created += 1
        else:
            updated += 1
    return created, updated, errors
