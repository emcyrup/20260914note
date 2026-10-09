import datetime
import json
from unittest import mock

from django.test import TestCase, override_settings

# Create your tests here.


from django.urls import reverse

from accounts.models import StaffAccount
from facilities.models import Facility
from .models import Beneficiary, BeneficiaryOffice


class BeneficiaryOfficeTests(TestCase):
    """利用事業所（上限管理事業所のフラグ）と重身フラグ"""

    def setUp(self):
        self.facility = Facility.objects.create(name='はぴねす', office_number='2650000001')
        self.user = StaffAccount.objects.create_user(username='staff', password='pw12345678', facility=self.facility)
        self.client.force_login(self.user)
        self.ben = Beneficiary.objects.create(
            facility=self.facility, last_name='中村', first_name='みお', date_of_birth='2015-09-14',
        )

    def test_add_office_and_manager_flag_is_unique(self):
        url = reverse('beneficiaries:office_create', args=[self.ben.pk])
        self.client.post(url, {'name': 'ひまわり', 'office_number': '2650000101', 'is_manager': 'on', 'order': 1})
        self.client.post(url, {'name': 'はぴねす', 'is_this_office': 'on', 'is_manager': 'on', 'order': 0})
        offices = list(self.ben.offices.all())
        self.assertEqual(len(offices), 2)
        managers = [o for o in offices if o.is_manager]
        self.assertEqual(len(managers), 1)
        self.assertTrue(managers[0].is_this_office)
        # 当施設の行は施設設定の番号を使う
        self.assertEqual(managers[0].office_number, '2650000001')
        self.assertTrue(self.ben.is_copayment_manager_here)
        self.assertEqual(self.ben.other_offices.count(), 1)

    def test_detail_shows_offices(self):
        BeneficiaryOffice.objects.create(beneficiary=self.ben, name='そら', office_number='2650000102', is_manager=True)
        res = self.client.get(reverse('beneficiaries:detail', args=[self.ben.pk]))
        self.assertContains(res, 'そら')
        self.assertContains(res, '上限管理事業所')
        self.assertFalse(self.ben.is_copayment_manager_here)

    def test_update_and_delete_office(self):
        o = BeneficiaryOffice.objects.create(beneficiary=self.ben, name='そら', office_number='2650000102')
        self.client.post(reverse('beneficiaries:office_update', args=[self.ben.pk, o.pk]),
                         {'name': 'そら（本店）', 'office_number': '2650000102', 'order': 0, 'fax': '075-000-0200'})
        o.refresh_from_db()
        self.assertEqual(o.name, 'そら（本店）')
        self.assertEqual(o.fax, '075-000-0200')
        self.client.post(reverse('beneficiaries:office_delete', args=[self.ben.pk, o.pk]))
        self.assertFalse(BeneficiaryOffice.objects.filter(pk=o.pk).exists())

    def test_other_facility_cannot_touch_offices(self):
        other = Facility.objects.create(name='別の事業所')
        other_ben = Beneficiary.objects.create(facility=other, last_name='他', first_name='子', date_of_birth='2015-01-01')
        res = self.client.post(reverse('beneficiaries:office_create', args=[other_ben.pk]), {'name': 'x', 'order': 0})
        self.assertEqual(res.status_code, 404)

    def test_is_severe_saved_from_edit(self):
        res = self.client.post(reverse('beneficiaries:update', args=[self.ben.pk]), {
            'last_name': '中村', 'first_name': 'みお', 'date_of_birth': '2015-09-14', 'gender': 'female',
            'status': 'active', 'is_severe': 'on',
        })
        self.assertEqual(res.status_code, 302)
        self.ben.refresh_from_db()
        self.assertTrue(self.ben.is_severe)
        self.facility.base_unit_count = 604
        self.facility.base_unit_count_severe = 1756
        self.assertEqual(self.facility.base_units_for(self.ben), 1756)
        self.facility.base_unit_count_severe = None
        self.assertEqual(self.facility.base_units_for(self.ben), 604)


class AssessmentTests(TestCase):
    """利用者台帳のアセスメント・資料"""

    def setUp(self):
        import shutil
        import tempfile
        from django.test import override_settings
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        self.override = override_settings(MEDIA_ROOT=self.media)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず')
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_STAFF, display_name='永山')
        self.client.login(username='ryo', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子',
                                              date_of_birth=datetime.date(2019, 4, 1))
        self.detail = reverse('beneficiaries:detail', args=[self.kid.pk])

    def pdf(self, name='kensa.pdf', size=100):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(name, b'%PDF-1.4 ' + b'0' * size, content_type='application/pdf')

    def test_add_text_and_file_then_open(self):
        from .models import BeneficiaryAssessment
        url = reverse('beneficiaries:assessment_create', args=[self.kid.pk])
        res = self.client.post(url, {'date': '2026-09-01', 'kind': 'test', 'title': '新版K式',
                                     'content': '言語 3歳相当。\n手先が器用。', 'file': self.pdf()})
        self.assertRedirects(res, self.detail + '#assessments', fetch_redirect_response=False)
        a = BeneficiaryAssessment.objects.get()
        self.assertEqual((a.kind, a.title, a.file_name, a.created_by), ('test', '新版K式', 'kensa.pdf', self.user))
        self.assertTrue(a.file.name.startswith('beneficiary_assessments/'))
        self.assertNotIn('kensa', a.file.name)
        res = self.client.get(self.detail)
        self.assertContains(res, '発達検査・評価')
        self.assertContains(res, '言語 3歳相当。')
        self.assertContains(res, a.file.url)
        # 書類は同じ事業所の職員だけが開ける
        self.assertEqual(self.client.get(a.file.url).status_code, 200)
        other = Facility.objects.create(name='ほか')
        StaffAccount.objects.create_user('oth', password='pw12345678', facility=other, role=StaffAccount.ROLE_STAFF)
        self.client.login(username='oth', password='pw12345678')
        self.assertEqual(self.client.get(a.file.url).status_code, 404)
        self.assertEqual(self.client.post(url, {'date': '2026-09-01', 'content': 'x'}).status_code, 404)

    def test_needs_content_or_file_and_checks_file(self):
        from .models import BeneficiaryAssessment
        url = reverse('beneficiaries:assessment_create', args=[self.kid.pk])
        self.client.post(url, {'date': '2026-09-01', 'title': '空'})
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.post(url, {'date': '2026-09-01', 'content': 'x', 'file': SimpleUploadedFile('a.exe', b'MZ')})
        self.client.post(url, {'date': '', 'content': 'x'})
        self.assertFalse(BeneficiaryAssessment.objects.exists())
        # 件名が空なら種類の名前
        self.client.post(url, {'date': '2026-09-01', 'kind': 'interview', 'content': '母より聞き取り'})
        self.assertEqual(BeneficiaryAssessment.objects.get().title, '面談・聞き取り')

    def test_edit_replace_remove_and_delete_file(self):
        import os
        from .models import BeneficiaryAssessment
        a = BeneficiaryAssessment.objects.create(beneficiary=self.kid, date=datetime.date(2026, 9, 1), title='初回',
                                                 content='x', file=self.pdf(), file_name='kensa.pdf')
        first = os.path.join(self.media, a.file.name)
        self.assertTrue(os.path.exists(first))
        edit = reverse('beneficiaries:assessment_update', args=[self.kid.pk, a.pk])
        self.client.post(edit, {'date': '2026-09-02', 'kind': 'assessment', 'title': '初回（直し）', 'content': 'y',
                                'file': self.pdf('new.pdf')})
        a.refresh_from_db()
        self.assertEqual((a.title, a.file_name), ('初回（直し）', 'new.pdf'))
        self.assertFalse(os.path.exists(first))            # 入れ替えた古い書類は消す
        second = os.path.join(self.media, a.file.name)
        self.client.post(edit, {'date': '2026-09-02', 'title': '初回', 'content': 'y', 'remove_file': '1'})
        a.refresh_from_db()
        self.assertFalse(a.file)
        self.assertFalse(os.path.exists(second))
        self.client.post(reverse('beneficiaries:assessment_delete', args=[self.kid.pk, a.pk]))
        self.assertFalse(BeneficiaryAssessment.objects.exists())


class BeneficiaryDocumentTests(TestCase):
    """利用者の基本情報の書類・画像（写真・PDF・Excel・CSV。ゆあーずだけ）"""

    def setUp(self):
        import datetime
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU)
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_STAFF)
        self.client.login(username='ryo', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2019, 4, 1))
        self.url = reverse('beneficiaries:document_upload', args=[self.kid.pk])

    @staticmethod
    def up(name, data=b'x', ctype='application/octet-stream'):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(name, data, content_type=ctype)

    def test_upload_list_delete(self):
        from .models import BeneficiaryDocument
        res = self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk]))
        self.assertContains(res, 'ここにファイルを置く')
        self.assertContains(res, 'js/dropzone.js')
        png = b'\x89PNG\r\n\x1a\n' + b'0' * 20
        res = self.client.post(self.url, {'files': [self.up('受給者証.png', png, 'image/png'), self.up('診断書.PDF'),
                                                    self.up('名簿.xlsx'), self.up('data.csv'), self.up('virus.exe')], 'title': '無視される'}, follow=True)
        self.assertContains(res, '書類を 4 件取り込みました')
        self.assertContains(res, 'virus.exe（この種類は入れられません）')
        docs = list(BeneficiaryDocument.objects.filter(beneficiary=self.kid).order_by('pk'))
        self.assertEqual([d.kind for d in docs], ['image', 'pdf', 'sheet', 'sheet'])
        self.assertEqual([d.title for d in docs], ['', '', '', ''])          # 複数のときは件名を付けない
        self.assertContains(res, '書類・画像（4）')
        self.assertContains(res, '受給者証.png')
        self.assertContains(res, 'bi-file-earmark-spreadsheet')
        # 1つだけなら件名が付く。ファイルはこの事業所の職員だけが開ける
        self.client.post(self.url, {'files': [self.up('契約書.jpg', png, 'image/jpeg')], 'title': '契約書 2026'})
        d = BeneficiaryDocument.objects.get(file_name='契約書.jpg')
        self.assertEqual(d.label, '契約書 2026')
        # 名前を変える（元のファイル名は残る）
        res = self.client.post(reverse('beneficiaries:document_rename', args=[self.kid.pk, d.pk]), {'title': '契約書（2026年度）'}, follow=True)
        self.assertContains(res, '書類の名前を「契約書（2026年度）」にしました')
        d.refresh_from_db()
        self.assertEqual((d.label, d.file_name), ('契約書（2026年度）', '契約書.jpg'))
        self.assertContains(res, 'doc-rename-btn')
        self.assertTrue(d.file.name.startswith('beneficiary_documents/'))
        self.assertEqual(self.client.get(d.file.url).status_code, 200)
        # 削除
        res = self.client.post(reverse('beneficiaries:document_delete', args=[self.kid.pk, d.pk]), follow=True)
        self.assertContains(res, '「契約書（2026年度）」を削除しました')
        self.assertFalse(BeneficiaryDocument.objects.filter(pk=d.pk).exists())
        # 空・大きすぎ
        res = self.client.post(self.url, {}, follow=True)
        self.assertContains(res, 'ファイルを選ぶか')
        from .models import DOCUMENT_MAX_BYTES
        big = self.up('big.pdf', b'0' * (DOCUMENT_MAX_BYTES + 1))
        res = self.client.post(self.url, {'files': [big]}, follow=True)
        self.assertContains(res, 'MB を超えています')

    def test_other_facility_and_non_ryoiku(self):
        import datetime
        from facilities.models import Facility
        from accounts.models import StaffAccount
        g = Facility.objects.create(name='ほか', layout=Facility.LAYOUT_RYOIKU)
        other = Beneficiary.objects.create(facility=g, last_name='井上', first_name='子', date_of_birth=datetime.date(2019, 4, 1))
        self.assertEqual(self.client.post(reverse('beneficiaries:document_upload', args=[other.pk]), {'files': [self.up('a.pdf')]}).status_code, 404)
        # 標準の型の事業所には出ない・受け付けない
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        res = self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk]))
        self.assertNotContains(res, 'ここにファイルを置く')
        self.assertEqual(self.client.post(self.url, {'files': [self.up('a.pdf')]}).status_code, 404)


class RecentTherapyOnDetailTests(TestCase):
    """利用者の詳細に療育記録の直近5日分（ゆあーず）"""

    def setUp(self):
        import datetime
        from accounts.models import StaffAccount
        from facilities.models import Facility
        from therapy.models import TherapyRecord
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU,
                                         use_reservation=True, use_therapy_record=True)
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_STAFF, display_name='大坂')
        self.client.login(username='ryo', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2019, 4, 1))
        for i in range(7):
            TherapyRecord.objects.create(facility=self.f, beneficiary=self.kid, date=datetime.date(2026, 10, 1) + datetime.timedelta(days=i * 2),
                                         time=datetime.time(10, 0), staff=self.user, activities=['ウレタン棒'], body=f'記録{i}')
        TherapyRecord.objects.create(facility=self.f, beneficiary=self.kid, date=datetime.date(2026, 10, 13), body='同じ日の2件目')

    def test_five_days_shown(self):
        res = self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk]))
        self.assertContains(res, '療育記録（直近5日分）')
        self.assertContains(res, '全部で 8 件')
        for body in ('記録6', '記録5', '記録4', '記録3', '記録2', '同じ日の2件目'):
            self.assertContains(res, body)
        self.assertNotContains(res, '記録1')
        self.assertNotContains(res, '記録0')
        self.assertContains(res, '①ウレタン棒')
        self.assertContains(res, '担当 大坂')
        self.assertContains(res, f'#rec{self.kid.therapy_records.first().pk}')

    def test_hidden_for_other_layouts(self):
        from facilities.models import Facility
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        self.assertNotContains(self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk])), '療育記録（直近5日分）')


class BeneficiaryImportTests(TestCase):
    """利用者情報の Excel・CSV 取り込み（ゆあーず）"""

    def setUp(self):
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU)
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.url = reverse('beneficiaries:import')

    @staticmethod
    def xlsx(rows):
        import io
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        for r in rows:
            ws.append(r)
        buf = io.BytesIO(); wb.save(buf)
        return SimpleUploadedFile('名簿.xlsx', buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @staticmethod
    def csv_file(text, enc='utf-8-sig'):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile('名簿.csv', text.encode(enc), content_type='text/csv')

    def test_template(self):
        from openpyxl import load_workbook
        import io
        res = self.client.get(reverse('beneficiaries:import_template'))
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(io.BytesIO(res.content))
        self.assertEqual(wb.sheetnames, ['利用者', '書き方'])
        header = [c.value for c in wb['利用者'][1]]
        self.assertEqual(header[:5], ['姓', '名', 'せい（ふりがな）', 'めい（ふりがな）', '生年月日'])
        self.assertIn('受給者証番号', header)

    def test_parse_row_and_errors(self):
        import datetime
        from . import importer
        data = importer.parse_row({'last_name': '山田', 'first_name': '太郎', 'date_of_birth': '2019年4月1日', 'gender': '女', 'grade': '小1',
                                   'weekdays': '月・水,金', 'is_severe': '○', 'postal_code': '600-8216', 'disability_class': '2級',
                                   'guardian_last_name': '山田', 'guardian_first_name': '花子', 'guardian_relation': '母',
                                   'certificate_number': '2600001234', 'granted_days': '１０', 'monthly_cap': '4,600円', 'valid_until': '2027/3/31'})
        self.assertEqual((data['date_of_birth'], data['gender'], data['grade'], data['is_severe'], data['postal_code'], data['disability_class']),
                         (datetime.date(2019, 4, 1), 'female', 'e1', True, '6008216', '2'))
        self.assertEqual(data['weekdays'], {'weekday_mon', 'weekday_wed', 'weekday_fri'})
        self.assertEqual(data['guardian']['relation'], 'mother')
        self.assertEqual((data['certificate']['granted_days'], data['certificate']['monthly_cap'], data['certificate']['valid_until']),
                         (10, 4600, datetime.date(2027, 3, 31)))
        with self.assertRaisesMessage(importer.RowError, '生年月日が空です'):
            importer.parse_row({'last_name': '山田', 'first_name': '太郎', 'date_of_birth': ''})
        with self.assertRaisesMessage(importer.RowError, '日付を読み取れません'):
            importer.parse_row({'last_name': '山田', 'first_name': '太郎', 'date_of_birth': '平成31年'})
        with self.assertRaisesMessage(importer.RowError, '性別の書き方が違います'):
            importer.parse_row({'last_name': '山田', 'first_name': '太郎', 'date_of_birth': '2019-04-01', 'gender': '男子'})
        with self.assertRaisesMessage(importer.RowError, '有効期間（終了）'):
            importer.parse_row({'last_name': '山田', 'first_name': '太郎', 'date_of_birth': '2019-04-01', 'certificate_number': '1'})
        with self.assertRaisesMessage(importer.RowError, '利用予定曜日の書き方'):
            importer.parse_row({'last_name': '山田', 'first_name': '太郎', 'date_of_birth': '2019-04-01', 'weekdays': '月・日'})

    def test_preview_then_commit_create_and_update(self):
        import datetime
        existing = Beneficiary.objects.create(facility=self.f, last_name='鈴木', first_name='一郎', date_of_birth=datetime.date(2018, 5, 5), address='前の住所')
        rows = [['姓', '名', '生年月日', '性別', '住所', '利用予定曜日', '保護者 姓', '保護者 名', '保護者 続柄', '受給者証番号', '受給者証 有効期間（終了）', '支給量（日/月）'],
                ['山田', '太郎', '2019-04-01', '男', '京都市', '月・水', '山田', '花子', '母', '2600001234', '2027-03-31', '10'],
                ['鈴木', '一郎', '2018/5/5', '', '', '火', '', '', '', '', '', ''],
                ['', '名無し', '2019-04-01', '', '', '', '', '', '', '', '', ''],
                ['佐藤', '花', 'いつか', '', '', '', '', '', '', '', '', '']]
        res = self.client.post(self.url, {'file': self.xlsx(rows)})
        self.assertContains(res, '2. 確かめる（まだ登録していません）')
        self.assertContains(res, '新規 1')
        self.assertContains(res, '台帳にいる 1')
        self.assertContains(res, '読めない行 2')
        self.assertContains(res, '新しく登録します（保護者・受給者証 も）')
        self.assertContains(res, 'すでに台帳にいます')
        self.assertContains(res, '台帳の空欄に入れる：利用予定曜日')
        self.assertContains(res, '台帳の値を残す')
        self.assertContains(res, '姓が空です')
        self.assertContains(res, '日付を読み取れません')
        self.assertEqual(Beneficiary.objects.filter(facility=self.f).count(), 1)     # まだ登録していない
        res = self.client.post(self.url, {'action': 'commit'}, follow=True)
        self.assertContains(res, '新規 1 名・台帳にいた人 1 名（台帳の値を残し、空欄だけ埋めました）（読めなかった行 2 件は登録していません）')
        self.assertEqual(Beneficiary.objects.filter(facility=self.f).count(), 2)
        taro = Beneficiary.objects.get(facility=self.f, last_name='山田')
        self.assertEqual((taro.gender, taro.address, taro.weekday_mon, taro.weekday_wed, taro.weekday_tue, taro.has_prior_records),
                         ('male', '京都市', True, True, False, True))
        g = taro.guardians.get()
        self.assertEqual((g.first_name, g.relation, g.is_primary), ('花子', 'mother', True))
        c = taro.recipient_certificates.get()
        self.assertEqual((c.certificate_number, c.granted_days, c.valid_until), ('2600001234', 10, datetime.date(2027, 3, 31)))
        existing.refresh_from_db()
        self.assertEqual((existing.address, existing.weekday_tue, existing.weekday_mon), ('前の住所', True, False))   # 空の欄は変えない
        # もう一度同じファイル → 全部「書き換え」、件数は増えない。受給者証も増えない
        res = self.client.post(self.url, {'file': self.xlsx(rows)})
        self.assertContains(res, '新規 0')
        self.assertContains(res, '台帳にいる 2')
        self.client.post(self.url, {'action': 'commit'})
        self.assertEqual(Beneficiary.objects.filter(facility=self.f).count(), 2)
        self.assertEqual(taro.recipient_certificates.count(), 1)
        self.assertEqual(taro.guardians.count(), 1)
        # 内容が無いまま登録
        res = self.client.post(self.url, {'action': 'commit'}, follow=True)
        self.assertContains(res, '取り込む内容がありません')

    def test_csv_shift_jis_and_bad_files(self):
        res = self.client.post(self.url, {'file': self.csv_file('姓,名,生年月日,学年\n高橋,結,2017-04-01,小3\n', enc='cp932')})
        self.assertContains(res, '新規 1')
        self.client.post(self.url, {'action': 'commit'})
        self.assertEqual(Beneficiary.objects.get(facility=self.f, last_name='高橋').grade, 'e3')
        res = self.client.post(self.url, {'file': self.csv_file('氏名,誕生日\n高橋 結,2017\n')})
        self.assertContains(res, '「姓」「名」「生年月日」が見つかりません')
        res = self.client.post(self.url, {})
        self.assertContains(res, 'ファイルを選んでください')
        res = self.client.post(self.url, {'file': self.csv_file('姓,名,生年月日\n')})
        self.assertContains(res, 'データの行がありません')

    def test_only_for_ryoiku_and_own_facility(self):
        import datetime
        from facilities.models import Facility
        g = Facility.objects.create(name='ほか', layout=Facility.LAYOUT_RYOIKU)
        Beneficiary.objects.create(facility=g, last_name='山田', first_name='太郎', date_of_birth=datetime.date(2019, 4, 1), address='よそ')
        self.client.post(self.url, {'file': self.csv_file('姓,名,生年月日\n山田,太郎,2019-04-01\n')})
        res = self.client.post(self.url, {'action': 'commit'}, follow=True)
        self.assertContains(res, '新規 1 名')                       # 別の事業所の同名は書き換えない
        self.assertEqual(Beneficiary.objects.filter(facility=g, address='よそ').count(), 1)
        self.assertContains(self.client.get(reverse('beneficiaries:list')), 'Excel・CSV から取り込む')
        self.f.layout = Facility.LAYOUT_STANDARD
        self.f.save(update_fields=['layout'])
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.assertEqual(self.client.get(reverse('beneficiaries:import_template')).status_code, 404)
        self.assertNotContains(self.client.get(reverse('beneficiaries:list')), 'Excel・CSV から取り込む')


class GuardianListImportTests(TestCase):
    """前のシステムの「保護者一覧」（1行が保護者1人・児童は名前で照合）の取り込み"""

    def setUp(self):
        import datetime
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU)
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.url = reverse('beneficiaries:import')
        self.taro = Beneficiary.objects.create(facility=self.f, last_name='山田', first_name='太郎', date_of_birth=datetime.date(2019, 4, 1))
        self.hana = Beneficiary.objects.create(facility=self.f, last_name='山田', first_name='花', date_of_birth=datetime.date(2021, 4, 1), address='もとの住所')

    @staticmethod
    def xlsx_from_format(rows):
        """依頼者から届いた書式（見出しだけ）に行を足す"""
        import io, os
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        path = os.path.join(os.path.dirname(__file__), 'testdata', 'guardian_list_format.xlsx')
        wb = load_workbook(path); ws = wb.active
        for r in rows:
            ws.append(r)
        buf = io.BytesIO(); wb.save(buf)
        return SimpleUploadedFile('保護者一覧.xlsx', buf.getvalue())

    def test_format_headers_and_parse(self):
        from . import importer
        rows = importer.rows_from_file(self.xlsx_from_format([
            ['山田 花子', 'ヤマダ ハナコ', '母', '山田 太郎、山田 花', '600-8216', '京都府', '京都市下京区', '○○町1-1', 'コーポ101', '送迎は母',
             '母携帯', '090-1111-1111', '', 'hanako@example.com', '父携帯', '090-2222-2222', '075-000-0000', '',
             '口座振替', '保護者', '山田 一郎', '600-0000', '京都府', '京都市', '△△町2-2', '', '075-111-1111', '', '請求は父へ'],
        ]))
        self.assertEqual(importer.format_of(rows), 'guardian')
        r = rows[0]
        self.assertEqual((r['g_name'], r['relation'], r['children'], r['c1_label'], r['c2_phone2'], r['bill_name'], r['bill_note']),
                         ('山田 花子', '母', '山田 太郎、山田 花', '母携帯', '075-000-0000', '山田 一郎', '請求は父へ'))
        data = importer.parse_guardian_row(r)
        g = data['guardian']
        self.assertEqual((g['last_name'], g['first_name'], g['kana'], g['relation'], g['phone'], g['phone2'], g['email'], g['memo']),
                         ('山田', '花子', 'ヤマダ ハナコ', 'mother', '090-1111-1111', '090-2222-2222', 'hanako@example.com', '送迎は母'))
        self.assertEqual(data['children'], ['山田 太郎', '山田 花'])
        self.assertEqual((data['postal_code'], data['address']), ('6008216', '京都府京都市下京区○○町1-1コーポ101'))
        self.assertEqual(g['extra']['連絡先1 名称'], '母携帯')
        self.assertEqual(g['extra']['連絡先2 電話番号2'], '075-000-0000')
        self.assertEqual(g['extra']['請求先 住所'], '京都府京都市△△町2-2')
        self.assertEqual(g['extra']['支払い方法'], '口座振替')
        self.assertNotIn('連絡先1 その他', g['extra'])        # メールは email に入れたので重複させない
        with self.assertRaisesMessage(importer.RowError, '児童が空です'):
            importer.parse_guardian_row({'g_name': '山田 花子', 'children': ''})

    def test_preview_commit_and_missing_child(self):
        f = self.xlsx_from_format([
            ['山田 花子', 'ヤマダ ハナコ', '母', '山田 太郎、山田 花', '600-8216', '京都府', '京都市下京区', '○○町1-1', '', '',
             '母携帯', '090-1111-1111', '', 'hanako@example.com', '', '', '', '', '', '', '', '', '', '', '', '', '', '', ''],
            ['鈴木 一郎', '', '祖父', '鈴木 次郎', '', '', '', '', '', '', '', '090-3333-3333', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', ''],
        ])
        res = self.client.post(self.url, {'file': f})
        self.assertContains(res, '保護者一覧</b>の形')
        self.assertContains(res, '台帳の児童に付ける 1')
        self.assertContains(res, '読めない行 1')
        self.assertContains(res, '台帳にいない児童：鈴木 次郎')
        self.assertContains(res, '保護者を 山田 太郎・山田 花 さんに付けます')
        res = self.client.post(self.url, {'action': 'commit'}, follow=True)
        self.assertContains(res, '保護者 2 件（新しく作った児童 0 名）（読めなかった行 1 件は登録していません）')
        g = self.taro.guardians.get()
        self.assertEqual((g.full_name, g.kana, g.relation, g.phone, g.email, g.is_primary, g.extra['連絡先1 名称']),
                         ('山田 花子', 'ヤマダ ハナコ', 'mother', '090-1111-1111', 'hanako@example.com', True, '母携帯'))
        self.assertEqual(self.hana.guardians.get().full_name, '山田 花子')
        self.taro.refresh_from_db(); self.hana.refresh_from_db()
        self.assertEqual((self.taro.postal_code, self.taro.address), ('6008216', '京都府京都市下京区○○町1-1'))
        self.assertEqual(self.hana.address, 'もとの住所')                          # 入っている住所は変えない
        self.assertFalse(Beneficiary.objects.filter(facility=self.f, last_name='鈴木').exists())
        # 「仮の生年月日で作る」を付けると児童も作る。同じ保護者をもう一度入れても増えない
        f = self.xlsx_from_format([
            ['山田 花子', 'ヤマダ ハナコ', '母', '山田 太郎', '', '', '', '', '', '', '', '090-9999-9999', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', ''],
            ['鈴木 一郎', '', '祖父', '鈴木 次郎', '', '', '', '', '', '', '', '090-3333-3333', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', ''],
        ])
        res = self.client.post(self.url, {'file': f, 'create_children': '1'})
        self.assertContains(res, '児童も新しく作る 1')
        self.assertContains(res, '仮の生年月日 2000-01-01 で新しく作ります')
        res = self.client.post(self.url, {'action': 'commit', 'create_children': '1', 'mode': 'overwrite'}, follow=True)
        self.assertContains(res, '保護者 2 件（新しく作った児童 1 名）')
        jiro = Beneficiary.objects.get(facility=self.f, last_name='鈴木', first_name='次郎')
        self.assertEqual(jiro.date_of_birth.isoformat(), '2000-01-01')
        self.assertIn('仮の値', jiro.notes)
        g2 = jiro.guardians.get()
        self.assertEqual((g2.full_name, g2.relation, g2.extra.get('続柄（原文）'), g2.phone), ('鈴木 一郎', 'other', '祖父', '090-3333-3333'))
        self.assertEqual(self.taro.guardians.count(), 1)
        self.assertEqual(self.taro.guardians.get().phone, '090-9999-9999')            # 書き換え
        # 詳細に「取り込んだ情報」
        res = self.client.get(reverse('beneficiaries:detail', args=[self.taro.pk]))
        self.assertContains(res, '取り込んだ情報')
        self.assertContains(res, '母携帯')
        self.assertContains(res, 'ヤマダ ハナコ')


class CsvVariantsImportTests(TestCase):
    """CSV の読み方：Shift_JIS・UTF-16（Excel の Unicode テキスト）・タブ区切り・表題行つき・拡張子違い"""

    def setUp(self):
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU)
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.url = reverse('beneficiaries:import')

    @staticmethod
    def up(name, data):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(name, data, content_type='application/octet-stream')

    def test_read_rows_variants(self):
        from config.tabular import read_rows, guess_delimiter
        body = '姓,名,生年月日\n高橋,結,2017-04-01\n'
        self.assertEqual(read_rows(body.encode('cp932'), '名簿.csv')[1], ['高橋', '結', '2017-04-01'])
        self.assertEqual(read_rows(('﻿' + body).encode('utf-8'), '名簿.csv')[0], ['姓', '名', '生年月日'])
        tab = '姓\t名\t生年月日\n高橋\t結\t2017/4/1\n'
        self.assertEqual(read_rows(tab.encode('utf-16'), '名簿.txt')[1], ['高橋', '結', '2017/4/1'])       # Excel の Unicode テキスト
        self.assertEqual(read_rows(tab.encode('utf-16-le'), '名簿.tsv')[1], ['高橋', '結', '2017/4/1'])    # BOM 無しの UTF-16
        self.assertEqual(read_rows('姓;名;生年月日\n高橋;結;2017-04-01\n'.encode('utf-8'), 'x.csv')[1][1], '結')
        self.assertEqual(guess_delimiter('a,b\tc\td'), '\t')
        # .csv という名前の xlsx も中身で見分ける
        import io
        from openpyxl import Workbook
        wb = Workbook(); wb.active.append(['姓', '名', '生年月日']); wb.active.append(['高橋', '結', '2017-04-01'])
        buf = io.BytesIO(); wb.save(buf)
        self.assertEqual(read_rows(buf.getvalue(), '名簿.csv')[1], ['高橋', '結', '2017-04-01'])
        with self.assertRaisesMessage(ValueError, '古い Excel 形式'):
            read_rows(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'0' * 16, '名簿.xls')
        with self.assertRaisesMessage(ValueError, 'PDF は読めません'):
            read_rows(b'%PDF-1.4 ...', '名簿.pdf')

    def test_title_row_above_header_and_guardian_csv(self):
        from . import importer
        # 表題行と空行のあとに見出し
        rows = importer.rows_from_file(self.up('名簿.csv', '利用者名簿（2026年度）\n\n姓,名,生年月日,学年\n高橋,結,2017-04-01,小3\n'.encode('cp932')))
        self.assertEqual(rows, [{'last_name': '高橋', 'first_name': '結', 'date_of_birth': '2017-04-01', 'grade': '小3'}])
        # 保護者一覧を Shift_JIS の CSV（前のシステムの書き出し）で
        header = ','.join(importer.GUARDIAN_HEADERS)
        line = ','.join(['山田 花子', 'ヤマダ ハナコ', '母', '山田 太郎', '', '京都府', '京都市', '1-1', '', '', '母携帯', '090-1111-1111'] + [''] * 17)
        rows = importer.rows_from_file(self.up('保護者一覧.csv', f'保護者一覧\n{header}\n{line}\n'.encode('cp932')))
        self.assertEqual(importer.format_of(rows), 'guardian')
        self.assertEqual((rows[0]['g_name'], rows[0]['children'], rows[0]['c1_label']), ('山田 花子', '山田 太郎', '母携帯'))
        # 画面からタブ区切り（UTF-16）で
        import datetime
        Beneficiary.objects.create(facility=self.f, last_name='山田', first_name='太郎', date_of_birth=datetime.date(2019, 4, 1))
        tsv = '\t'.join(importer.GUARDIAN_HEADERS) + '\n' + line.replace(',', '\t') + '\n'
        res = self.client.post(self.url, {'file': self.up('保護者一覧.txt', tsv.encode('utf-16'))})
        self.assertContains(res, '台帳の児童に付ける 1')


class ChildrenListImportTests(TestCase):
    """前のシステムの「児童一覧」（見出しなし・15列・1行が契約1件）の取り込み"""

    ROWS = [
        ['山田 花子', 'ヤマダ ハナコ', '000010', '山田 太郎', 'ヤマダ タロウ', '男', '2018/4/2', '', '退所', '2023/4/1', '2025/3/31', '利用中', '2025/4/1', '', '利用なし'],
        ['山田 花子', 'ヤマダ ハナコ', '000011', '山田 太郎', 'ヤマダ タロウ', '男', '2018/4/2', '', '利用なし', '', '', '退所', '2021/1/1', '2022/3/31', '利用なし'],
        ['山田 花子', 'ヤマダ ハナコ', '000012', '山田 花', 'ヤマダ ハナ', '女', '2020/5/5', '', '退所', '2024/4/1', '2025/3/31', '利用なし', '', '', '利用なし'],
        ['佐藤 花子', 'サトウ ハナコ', '000013', '佐藤 花', 'サトウ ハナ', '女', '2020/5/5', '', '利用中', '2025/4/1', '', '利用なし', '', '', '利用なし'],
    ]

    def setUp(self):
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU)
        self.user = StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.url = reverse('beneficiaries:import')

    @staticmethod
    def csv(rows, enc='cp932', name='児童一覧.csv'):
        from django.core.files.uploadedfile import SimpleUploadedFile
        text = '\r\n'.join(','.join(r) for r in rows) + '\r\n'
        return SimpleUploadedFile(name, text.encode(enc))

    def test_rows_are_grouped_per_child(self):
        from . import importer
        rows = importer.rows_from_file(self.csv(self.ROWS))
        self.assertEqual(importer.format_of(rows), importer.FORMAT_CHILDREN)
        self.assertEqual([(r['last_name'], r['first_name']) for r in rows], [('山田', '太郎'), ('山田', '花'), ('佐藤', '花')])
        taro, hana, sato = rows
        self.assertEqual((taro['status'], taro['admission_date'], taro['discharge_date']), ('在籍中', '2021/1/1', ''))
        self.assertEqual((taro['last_name_kana'], taro['first_name_kana'], taro['guardian_kana']), ('やまだ', 'たろう', 'やまだ はなこ'))
        self.assertEqual((taro['guardian_last_name'], taro['guardian_first_name'], taro['gender']), ('山田', '花子', '男'))
        self.assertIn('管理番号 000010 児童発達支援 退所 2023/4/1〜2025/3/31', taro['notes_append'])
        self.assertIn('管理番号 000010 放課後等デイ 利用中 2025/4/1〜', taro['notes_append'])
        self.assertIn('管理番号 000011 放課後等デイ 退所 2021/1/1〜2022/3/31', taro['notes_append'])
        self.assertEqual((hana['status'], hana['admission_date'], hana['discharge_date']), ('退所', '2024/4/1', '2025/3/31'))
        self.assertIn('姓だけ違う', hana['_detail'])
        self.assertIn('佐藤 花', hana['_detail'])
        self.assertIn('山田 花', sato['_detail'])

    def test_header_row_and_utf8_are_accepted(self):
        from . import importer
        header = ['保護者', 'カナ', '番号', '児童', 'カナ', '性別', '生年月日', '', '児発', '開始', '終了', '放デイ', '開始', '終了', '他']
        rows = importer.rows_from_file(self.csv([header] + self.ROWS, enc='utf-8', name='list.csv'))
        self.assertEqual(importer.format_of(rows), importer.FORMAT_CHILDREN)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]['_line'], 2)

    HEADER17 = ['保護者（名前）', '保護者（カナ）', '管理番号', '児童（名前）', '児童（カナ）', '性別', '生年月日', '備考',
                '児発状態', '児発利用契約日', '児発退所日', '放デイ状態', '放デイ利用契約日', '放デイ退所日', '保訪状態', '保訪利用契約日', '保訪退所日']

    def test_header_17_columns_with_note_and_visit_support(self):
        """見出し付き 17 列（ゆあーずの書式）。保護者（名前）があっても保護者一覧ではなく児童一覧として読む"""
        from . import importer
        rows = [self.HEADER17,
                ['山田 花子', 'ヤマダ ハナコ', '20', '山田 太郎', 'ヤマダ タロウ', '男', '2018/4/2', '卵アレルギー',
                 '', '', '', '利用中', '2025/4/1', '', '利用中', '2025/6/1', ''],
                ['', '', '21', '佐藤 花', 'サトウ ハナ', '女', '2019/5/5', '', '退所', '2023/4/1', '2025/3/31', '', '', '', '', '', ''],
                ['', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '']]
        data = importer.rows_from_file(self.csv(rows, enc='cp932', name='児童.csv'))
        self.assertEqual(importer.format_of(data), importer.FORMAT_CHILDREN)
        self.assertEqual(len(data), 2)
        taro, hana = data
        self.assertEqual((taro['status'], taro['admission_date']), ('在籍中', '2025/4/1'))
        self.assertIn('管理番号 20 保育所等訪問支援 利用中 2025/6/1〜', taro['notes_append'])
        self.assertIn('備考：卵アレルギー', taro['notes_append'])
        self.assertEqual((hana['status'], hana['discharge_date']), ('退所', '2025/3/31'))
        self.assertNotIn('guardian_last_name', hana)
        res = self.client.post(self.url, {'file': self.csv(rows, name='児童.csv')})
        self.assertContains(res, '新規 2')
        self.client.post(self.url, {'action': 'commit'})
        taro = Beneficiary.objects.get(facility=self.f, last_name='山田')
        self.assertIn('備考：卵アレルギー', taro.notes)
        self.assertEqual(taro.guardians.get().full_name, '山田 花子')

    def test_existing_child_keep_or_overwrite(self):
        """台帳にいる子：違う値は「残す」「上書き」を選べる。台帳の空欄はどちらでも埋める"""
        import datetime
        rows = [self.HEADER17,
                ['山田 花子', 'ヤマダ ハナコ', '20', '山田 太郎', 'ヤマダ タロウ', '男', '2018/4/2', '',
                 '', '', '', '利用中', '2025/4/1', '', '', '', '']]
        b = Beneficiary.objects.create(facility=self.f, last_name='山田', first_name='太郎', date_of_birth=datetime.date(2018, 4, 2),
                                       last_name_kana='やまもと', first_name_kana='', gender='male')
        res = self.client.post(self.url, {'file': self.csv(rows)})
        self.assertContains(res, '台帳にいる 1')
        self.assertContains(res, '台帳と違う値')
        self.assertContains(res, 'ふりがな（姓）：台帳「やまもと」／取り込み「やまだ」')
        self.assertContains(res, '台帳の空欄に入れる：ふりがな（名）')
        self.assertContains(res, 'name="mode" value="keep"')
        self.client.post(self.url, {'action': 'commit', 'mode': 'keep'})
        b.refresh_from_db()
        self.assertEqual((b.last_name_kana, b.first_name_kana, b.admission_date), ('やまもと', 'たろう', datetime.date(2025, 4, 1)))
        self.client.post(self.url, {'file': self.csv(rows)})
        self.client.post(self.url, {'action': 'commit', 'mode': 'overwrite'})
        b.refresh_from_db()
        self.assertEqual(b.last_name_kana, 'やまだ')
        self.assertEqual(Beneficiary.objects.filter(facility=self.f).count(), 1)

    def test_preview_and_commit_twice(self):
        import datetime
        res = self.client.post(self.url, {'file': self.csv(self.ROWS)})
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, '児童一覧')
        self.assertContains(res, '放課後等デイ 利用中（2025/4/1〜）')
        self.assertContains(res, '姓だけ違う')
        res = self.client.post(self.url, {'action': 'commit'}, follow=True)
        self.assertContains(res, '新規 3 名')
        taro = Beneficiary.objects.get(facility=self.f, last_name='山田', first_name='太郎')
        self.assertEqual((taro.status, taro.admission_date, taro.discharge_date, taro.gender),
                         ('active', datetime.date(2021, 1, 1), None, 'male'))
        self.assertEqual((taro.last_name_kana, taro.first_name_kana), ('やまだ', 'たろう'))
        self.assertEqual(taro.notes.count('管理番号 000010'), 2)
        self.assertTrue(taro.has_prior_records)
        g = taro.guardians.get()
        self.assertEqual((g.last_name, g.first_name, g.kana, g.relation, g.is_primary), ('山田', '花子', 'やまだ はなこ', 'other', True))
        hana = Beneficiary.objects.get(facility=self.f, last_name='山田', first_name='花')
        self.assertEqual((hana.status, hana.discharge_date), ('inactive', datetime.date(2025, 3, 31)))
        # もう一度読ませても増えない。備考も二重にならない
        self.client.post(self.url, {'file': self.csv(self.ROWS)})
        res = self.client.post(self.url, {'action': 'commit'}, follow=True)
        self.assertContains(res, '台帳にいた人 3 名')
        self.assertEqual(Beneficiary.objects.filter(facility=self.f).count(), 3)
        taro.refresh_from_db()
        self.assertEqual(taro.notes.count('管理番号 000010'), 2)
        self.assertEqual(taro.guardians.count(), 1)

    def test_standard_format_status_and_discharge(self):
        import datetime
        from . import importer
        header = ','.join(['姓', '名', '生年月日', '在籍状況', '退所日', '保護者 姓', '保護者 名', '保護者 ふりがな'])
        f = self.csv([header.split(','), ['山田', '太郎', '2018-04-02', '退所', '2025-03-31', '山田', '花子', 'やまだ はなこ']], enc='utf-8', name='a.csv')
        rows = importer.rows_from_file(f)
        self.assertEqual(importer.format_of(rows), importer.FORMAT_BENEFICIARY)
        importer.apply(self.f, importer.plan(self.f, rows))
        taro = Beneficiary.objects.get(facility=self.f, last_name='山田', first_name='太郎')
        self.assertEqual((taro.status, taro.discharge_date), ('inactive', datetime.date(2025, 3, 31)))
        self.assertEqual(taro.guardians.get().kana, 'やまだ はなこ')
        f = self.csv([['姓', '名', '生年月日', '在籍状況'], ['山田', '太郎', '2018-04-02', '在籍中']], enc='utf-8', name='b.csv')
        importer.apply(self.f, importer.plan(self.f, importer.rows_from_file(f)))
        taro.refresh_from_db()
        self.assertEqual((taro.status, taro.discharge_date), ('active', None))
        f = self.csv([['姓', '名', '生年月日', '在籍状況'], ['山田', '太郎', '2018-04-02', '通所中']], enc='utf-8', name='c.csv')
        planned = importer.plan(self.f, importer.rows_from_file(f))
        self.assertEqual(planned[0]['action'], 'error')
        self.assertIn('在籍状況', planned[0]['detail'])


class KnowledgeTests(TestCase):
    """診断書などを AI で読み取り → 確認して保存・台帳に反映 → 療育記録の AI が参考にする（RAG）"""

    READ = {
        'kind': '診断書', 'title': '診断書', 'doc_date': '2026-04-10', 'issuer': 'こども発達クリニック',
        'name_on_doc': '青木 子', 'birth_on_doc': '2019-04-01', 'diagnosis': '自閉スペクトラム症', 'severe': 'no',
        'summary': '自閉スペクトラム症と診断。聴覚過敏があり、大きな音で混乱しやすい。',
        'support_points': ['大きな音の出る活動では、イヤーマフを使う', '予定の変更は絵カードで前もって伝える'],
        'medical_notes': ['てんかんの薬（朝夕）を服用している'],
        'full_text': '診断名：自閉スペクトラム症\n所見：聴覚過敏が強く、太鼓や掃除機の音で耳をふさぐ。トランポリンなど揺れる遊びは好む。\n'
                     '保険証番号：（書かない）',
        'unreadable': '',
    }

    def setUp(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from accounts.models import StaffAccount
        from facilities.models import Facility
        from .models import BeneficiaryDocument
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU, use_therapy_record=True)
        StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', last_name_kana='あおき',
                                              first_name_kana='こ', date_of_birth=datetime.date(2019, 4, 1))
        self.doc = BeneficiaryDocument.objects.create(beneficiary=self.kid, file_name='診断書.pdf',
                                                      file=SimpleUploadedFile('shindan.pdf', b'%PDF-1.4 test'))

    def _mock_response(self, data, stop='end_turn'):
        from types import SimpleNamespace
        return SimpleNamespace(stop_reason=stop, content=[SimpleNamespace(type='text', text=json.dumps(data, ensure_ascii=False))])

    def _read(self, client_cls, data=None):
        client_cls.return_value.messages.create.return_value = self._mock_response(data or self.READ)
        return self.client.post(reverse('beneficiaries:knowledge_read', args=[self.kid.pk]),
                                {'source': 'document', 'source_pk': self.doc.pk})

    @mock.patch('beneficiaries.knowledge.anthropic.Anthropic')
    @override_settings(ANTHROPIC_API_KEY='test-key')
    def test_read_review_save_and_apply(self, client_cls):
        from therapy.models import TherapyProfile
        from .models import BeneficiaryKnowledge
        TherapyProfile.objects.create(beneficiary=self.kid, cautions='・予定の変更は絵カードで前もって伝える')
        res = self._read(client_cls)
        k = BeneficiaryKnowledge.objects.get()
        self.assertRedirects(res, reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k.pk]))
        kwargs = client_cls.return_value.messages.create.call_args.kwargs
        block = kwargs['messages'][0]['content'][0]
        self.assertEqual((block['type'], block['source']['media_type']), ('document', 'application/pdf'))
        self.assertEqual(kwargs['output_config']['format']['type'], 'json_schema')
        self.assertIn('青木 子さん', kwargs['messages'][0]['content'][1]['text'])
        # 確認待ち：台帳はまだ変わらない。反映案（留意点はすでにある項目を除く）
        self.assertEqual(k.status, 'draft')
        self.assertEqual(k.proposals['disability_type'], {'current': '', 'new': '自閉スペクトラム症', 'checked': True})
        self.assertNotIn('is_severe', k.proposals)
        self.assertIn('てんかんの薬', k.proposals['notes_add']['text'])
        self.assertIn('イヤーマフ', k.proposals['cautions_add']['text'])
        self.assertNotIn('絵カード', k.proposals['cautions_add']['text'])
        self.assertEqual(k.proposals['warnings'], [])
        self.kid.refresh_from_db()
        self.assertEqual(self.kid.disability_type, '')
        page = self.client.get(reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k.pk]))
        self.assertContains(page, '確認待ち')
        self.assertContains(page, '障害種別を書き換える')
        self.assertContains(page, '療育記録の「留意点」に足す')
        self.assertContains(page, 'name="apply_assessment"')
        detail = self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk]))
        self.assertContains(detail, '確認待ち')
        # 保存：印を付けたもの（障害種別・留意点）だけ反映。備考は印なし
        res = self.client.post(reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k.pk]), {
            'action': 'save', 'kind': '診断書', 'title': '診断書', 'doc_date': '2026-04-10', 'issuer': 'こども発達クリニック',
            'summary': k.summary, 'points': k.points, 'text': k.text, 'use_in_ai': '1',
            'apply_disability_type': '1', 'disability_type': '自閉スペクトラム症',
            'notes_add': k.proposals['notes_add']['text'],
            'apply_cautions': '1', 'cautions_add': k.proposals['cautions_add']['text'],
            'apply_assessment': '1',
        }, follow=True)
        self.assertContains(res, '台帳に反映：障害種別・留意点・アセスメント・資料')
        # アセスメント・資料にも 1 件（書類のファイルを写して付ける）
        a = self.kid.assessments.get()
        self.assertEqual((a.kind, a.title, str(a.date), a.file_name), ('medical', '診断書', '2026-04-10', '診断書.pdf'))
        self.assertIn('聴覚過敏', a.content)
        self.assertIn('【支援で気をつけること】', a.content)
        self.assertIn('発行元：こども発達クリニック', a.content)
        self.assertTrue(a.file and a.file.name != self.doc.file.name and a.file.storage.exists(a.file.name))
        self.kid.refresh_from_db()
        self.assertEqual((self.kid.disability_type, self.kid.notes), ('自閉スペクトラム症', ''))
        self.assertIn('イヤーマフ', TherapyProfile.objects.get(beneficiary=self.kid).cautions)
        k.refresh_from_db()
        self.assertEqual((k.status, k.applied, k.use_in_ai), ('saved', ['障害種別', '留意点', 'アセスメント・資料'], True))
        self.assertGreaterEqual(k.chunks.count(), 3)
        detail = self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk]))
        self.assertContains(detail, '書類から分かっていること')
        self.assertContains(detail, '読み取り済み')
        self.assertContains(detail, 'AIの参考')
        # 同じ書類を読み直して保存すると入れ替わる
        self._read(client_cls)
        k2 = BeneficiaryKnowledge.objects.get(status='draft')
        self.client.post(reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k2.pk]),
                         {'action': 'save', 'summary': 'x', 'points': '', 'text': '', 'use_in_ai': '1'})
        self.assertEqual(list(BeneficiaryKnowledge.objects.values_list('pk', flat=True)), [k2.pk])

    @mock.patch('beneficiaries.knowledge.anthropic.Anthropic')
    @override_settings(ANTHROPIC_API_KEY='test-key')
    def test_warnings_discard_and_errors(self, client_cls):
        from .models import BeneficiaryKnowledge
        self._read(client_cls, dict(self.READ, name_on_doc='山田 花子', birth_on_doc='2018-01-01', severe='yes'))
        k = BeneficiaryKnowledge.objects.get()
        self.assertEqual(len(k.proposals['warnings']), 2)
        self.assertEqual(k.proposals['is_severe'], {'checked': True})
        page = self.client.get(reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k.pk]))
        self.assertContains(page, '別のお子さまの書類でないか')
        res = self.client.post(reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k.pk]), {'action': 'discard'}, follow=True)
        self.assertContains(res, '読み取り結果を消しました')
        self.assertFalse(BeneficiaryKnowledge.objects.exists())
        # AI が断った・長すぎた
        client_cls.return_value.messages.create.return_value = self._mock_response({}, stop='refusal')
        res = self.client.post(reverse('beneficiaries:knowledge_read', args=[self.kid.pk]),
                               {'source': 'document', 'source_pk': self.doc.pk}, follow=True)
        self.assertContains(res, 'AI がこの書類を読み取れませんでした')
        client_cls.return_value.messages.create.return_value = self._mock_response({}, stop='max_tokens')
        res = self.client.post(reverse('beneficiaries:knowledge_read', args=[self.kid.pk]),
                               {'source': 'document', 'source_pk': self.doc.pk}, follow=True)
        self.assertContains(res, 'ページを分けて')
        self.assertFalse(BeneficiaryKnowledge.objects.exists())

    def test_readable_types_and_other_facility(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from accounts.models import StaffAccount
        from facilities.models import Facility
        from . import knowledge
        from .models import BeneficiaryDocument
        self.assertTrue(knowledge.is_readable('a.PDF'))
        self.assertTrue(knowledge.is_readable('a.jpeg'))
        self.assertFalse(knowledge.is_readable('a.xlsx'))
        heic = BeneficiaryDocument.objects.create(beneficiary=self.kid, file_name='a.heic', file=SimpleUploadedFile('a.heic', b'x'))
        with override_settings(ANTHROPIC_API_KEY='test-key'):
            res = self.client.post(reverse('beneficiaries:knowledge_read', args=[self.kid.pk]),
                                   {'source': 'document', 'source_pk': heic.pk}, follow=True)
        self.assertContains(res, 'HEIC')
        with override_settings(ANTHROPIC_API_KEY=''):
            res = self.client.post(reverse('beneficiaries:knowledge_read', args=[self.kid.pk]),
                                   {'source': 'document', 'source_pk': self.doc.pk}, follow=True)
        self.assertContains(res, 'ANTHROPIC_API_KEY')
        other = Facility.objects.create(name='ほか', layout=Facility.LAYOUT_RYOIKU)
        StaffAccount.objects.create_user('oth', password='pw12345678', facility=other, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='oth', password='pw12345678')
        res = self.client.post(reverse('beneficiaries:knowledge_read', args=[self.kid.pk]), {'source': 'document', 'source_pk': self.doc.pk})
        self.assertEqual(res.status_code, 404)

    def test_context_for_only_this_child_and_relevant_passages(self):
        from . import knowledge
        from .models import BeneficiaryKnowledge
        other = Beneficiary.objects.create(facility=self.f, last_name='井上', first_name='太郎', date_of_birth=datetime.date(2018, 5, 1))
        k = BeneficiaryKnowledge.objects.create(beneficiary=self.kid, kind='診断書', title='診断書', issuer='こども発達クリニック',
                                                summary='聴覚過敏がある。', points='・大きな音の活動ではイヤーマフを使う',
                                                text=self.READ['full_text'] + '\n' + 'その他の記載。' * 80, status='saved')
        knowledge.reindex(k)
        BeneficiaryKnowledge.objects.create(beneficiary=other, kind='診断書', title='井上の書類', summary='別の子', status='saved')
        BeneficiaryKnowledge.objects.create(beneficiary=self.kid, kind='その他', title='確認待ち', summary='まだ', status='draft')
        text, n = knowledge.context_for(self.kid, '太鼓あそび トランポリン')
        self.assertEqual(n, 1)
        self.assertIn('書類から分かっていること', text)
        self.assertIn('イヤーマフ', text)
        self.assertIn('太鼓や掃除機の音', text)          # 今日の活動に関係しそうな記載
        self.assertNotIn('井上', text)
        self.assertNotIn('確認待ち', text)
        self.assertLessEqual(len(text), knowledge.CONTEXT_MAX + 80)
        k.use_in_ai = False
        k.save()
        self.assertEqual(knowledge.context_for(self.kid, '太鼓'), ('', 0))


@override_settings(ANTHROPIC_API_KEY='test-key')
class KnowledgeInTherapyAiTests(TestCase):
    """療育記録の AI（留意点の要約・記録の文）に、その子の書類から分かっていることが渡る"""

    def setUp(self):
        from accounts.models import StaffAccount
        from facilities.models import Facility
        from . import knowledge
        from .models import BeneficiaryKnowledge
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU, use_therapy_record=True)
        StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2019, 4, 1))
        k = BeneficiaryKnowledge.objects.create(beneficiary=self.kid, kind='診断書', title='診断書', summary='聴覚過敏がある。',
                                                points='・大きな音の活動ではイヤーマフを使う', text='太鼓の音で耳をふさぐ。', status='saved')
        knowledge.reindex(k)

    def _ai(self, client_cls, text):
        from types import SimpleNamespace
        client_cls.return_value.messages.create.return_value = SimpleNamespace(content=[SimpleNamespace(type='text', text=text)])

    @mock.patch('ai_assist.quick.anthropic.Anthropic')
    def test_record_summary_and_cautions_get_reference(self, client_cls):
        self._ai(client_cls, '太鼓あそびでは聴覚過敏との所見があるため、イヤーマフを使って取り組む。' * 3)
        res = self.client.post(reverse('therapy:record_summary', args=[self.kid.pk]),
                               {'cautions': '・音に敏感', 'activity_1': '太鼓あそび', 'date': '2026-10-02'})
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()['references'], 1)
        kwargs = client_cls.return_value.messages.create.call_args.kwargs
        self.assertIn('【書類から分かっていること（参考。職員が確かめて登録したもの）】', kwargs['messages'][0]['content'])
        self.assertIn('イヤーマフ', kwargs['messages'][0]['content'])
        self.assertIn('太鼓の音で耳をふさぐ', kwargs['messages'][0]['content'])
        self.assertIn('医学的な判断を加えたりしない', kwargs['system'])
        # 留意点の「短く要約」も参考にする。「詳しくまとめる」は話したことの整理なので使わない
        self._ai(client_cls, '・音に敏感\n・大きな音の活動ではイヤーマフを使う')
        res = self.client.post(reverse('therapy:cautions_summary', args=[self.kid.pk]), {'text': '音に敏感'})
        self.assertEqual(res.json()['references'], 1)
        self.assertIn('イヤーマフ', client_cls.return_value.messages.create.call_args.kwargs['messages'][0]['content'])
        res = self.client.post(reverse('therapy:cautions_summary', args=[self.kid.pk]), {'text': '音に敏感', 'mode': 'detail'})
        self.assertEqual(res.json()['references'], 0)
        self.assertNotIn('書類から分かっていること', client_cls.return_value.messages.create.call_args.kwargs['messages'][0]['content'])
        page = self.client.get(reverse('therapy:child', args=[self.kid.pk]))
        self.assertContains(page, '書類から分かっていること</a>（1 件')


class RyoikuConflictTests(TestCase):
    """アセスメント・資料と、書類から分かっていることの修正の競合"""

    def setUp(self):
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU)
        StaffAccount.objects.create_user('ryo', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='ryo', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2019, 4, 1))

    def test_assessment_and_knowledge(self):
        from config.concurrency import version_token
        from .models import BeneficiaryAssessment, BeneficiaryKnowledge
        a = BeneficiaryAssessment.objects.create(beneficiary=self.kid, date=datetime.date(2026, 9, 1), title='面談', content='もと')
        opened = version_token(a)
        detail = self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk]))
        self.assertContains(detail, f'data-concurrency="beneficiary_assessment:{a.pk}" data-concurrency-lazy')
        a.content = '他の職員'
        a.save()
        url = reverse('beneficiaries:assessment_update', args=[self.kid.pk, a.pk])
        res = self.client.post(url, {'date': '2026-09-01', 'title': '面談', 'kind': 'assessment', 'content': 'わたし', 'version': opened}, follow=True)
        self.assertContains(res, '他の職員が')
        a.refresh_from_db()
        self.assertEqual(a.content, '他の職員')
        k = BeneficiaryKnowledge.objects.create(beneficiary=self.kid, kind='診断書', title='診断書', summary='もと', status='saved')
        opened = version_token(k)
        page = self.client.get(reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k.pk]))
        self.assertContains(page, f'data-concurrency="knowledge:{k.pk}"')
        k.summary = '他の職員'
        k.save()
        res = self.client.post(reverse('beneficiaries:knowledge_review', args=[self.kid.pk, k.pk]),
                               {'action': 'save', 'summary': 'わたし', 'points': '', 'text': '', 'use_in_ai': '1', 'version': opened}, follow=True)
        self.assertContains(res, '他の職員が')
        k.refresh_from_db()
        self.assertEqual(k.summary, '他の職員')


class ListOrderTests(TestCase):
    """利用者一覧の 50 音順（ふりがなの無い人は最後・行の絞り込み・カタカナはひらがなに）"""

    def setUp(self):
        self.f = Facility.objects.create(name='はぴねす')
        StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f)
        self.client.login(username='st', password='pw12345678')
        d = datetime.date(2018, 1, 1)
        for ln, fn, lk, fk in [('田中', '太郎', 'たなか', 'たろう'), ('佐藤', '花', 'サトウ', 'ハナ'), ('渡辺', '空', '', ''),
                               ('青木', '子', 'あおき', 'こ'), ('高橋', '一', 'たかはし', 'はじめ')]:
            Beneficiary.objects.create(facility=self.f, last_name=ln, first_name=fn, last_name_kana=lk, first_name_kana=fk, date_of_birth=d)

    def test_kana_order_blank_last_and_rows(self):
        b = Beneficiary.objects.get(last_name='佐藤')
        self.assertEqual((b.last_name_kana, b.first_name_kana), ('さとう', 'はな'))     # 保存でひらがなに
        names = [x.last_name for x in self.client.get(reverse('beneficiaries:list')).context['beneficiaries']]
        self.assertEqual(names, ['青木', '佐藤', '高橋', '田中', '渡辺'])              # ふりがな無しは最後
        names = [x.last_name for x in self.client.get(reverse('beneficiaries:list') + '?row=た').context['beneficiaries']]
        self.assertEqual(names, ['高橋', '田中'])
        names = [x.last_name for x in self.client.get(reverse('beneficiaries:list') + '?row=他').context['beneficiaries']]
        self.assertEqual(names, ['渡辺'])
        self.assertContains(self.client.get(reverse('beneficiaries:list')), 'ふりがな無し')


class BeneficiaryDeleteTests(TestCase):
    """退所した利用者の削除（管理者だけ・氏名の確認・関連する記録とファイルも消える）"""

    def setUp(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import BeneficiaryDocument
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU, use_therapy_record=True)
        self.admin = StaffAccount.objects.create_user('adm', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.staff = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f)
        self.b = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', last_name_kana='あおき',
                                            first_name_kana='こ', date_of_birth=datetime.date(2019, 4, 1), status='inactive')
        self.doc = BeneficiaryDocument.objects.create(beneficiary=self.b, file_name='a.pdf', file=SimpleUploadedFile('a.pdf', b'%PDF-1.4 x'))
        from support_plans.models import SupportPlan
        SupportPlan.objects.create(facility=self.f, beneficiary=self.b, title='第1期', created_by=self.admin)
        from therapy.models import TherapyRecord
        TherapyRecord.objects.create(facility=self.f, beneficiary=self.b, date=datetime.date(2026, 9, 1))
        self.url = reverse('beneficiaries:delete', args=[self.b.pk])

    def test_only_admin_and_active_is_warned(self):
        self.client.login(username='st', password='pw12345678')
        res = self.client.post(self.url, {'confirm_name': '青木 子', 'agree': '1'}, follow=True)
        self.assertContains(res, '管理者だけ')
        self.assertTrue(Beneficiary.objects.filter(pk=self.b.pk).exists())
        self.assertNotContains(self.client.get(reverse('beneficiaries:detail', args=[self.b.pk])), 'を削除する</a>')
        self.client.login(username='adm', password='pw12345678')
        self.b.status = 'active'
        self.b.save()
        res = self.client.get(self.url)
        self.assertContains(res, 'この人はいま「在籍中」です')
        res = self.client.post(self.url, {'confirm_name': '青木 子', 'agree': '1'}, follow=True)
        self.assertContains(res, '「青木 子」を削除しました')
        self.assertFalse(Beneficiary.objects.filter(pk=self.b.pk).exists())

    def test_status_tabs_graduated_and_list_actions(self):
        self.client.login(username='adm', password='pw12345678')
        Beneficiary.objects.create(facility=self.f, last_name='井上', first_name='空', last_name_kana='いのうえ',
                                   first_name_kana='そら', date_of_birth=datetime.date(2018, 4, 1))
        res = self.client.post(reverse('beneficiaries:status', args=[self.b.pk]), {'status': 'graduated'}, follow=True)
        self.assertContains(res, '「卒業」にしました')
        self.b.refresh_from_db()
        self.assertEqual(self.b.status, 'graduated')
        self.assertEqual(self.b.discharge_date, datetime.date.today())
        lst = reverse('beneficiaries:list')
        res = self.client.get(lst)
        self.assertContains(res, '井上 空')
        self.assertNotContains(res, '青木 子')
        self.assertContains(res, '卒業 <span')
        self.assertContains(res, reverse('therapy:child', args=[Beneficiary.objects.get(last_name='井上').pk]))
        res = self.client.get(lst + '?status=graduated')
        self.assertContains(res, '青木 子')
        self.assertNotContains(res, '井上 空')
        self.assertContains(res, reverse('beneficiaries:delete', args=[self.b.pk]))
        self.assertNotContains(self.client.get(lst + '?status=inactive'), '青木 子')
        # 在籍中に戻すと退所日は消える
        self.client.post(reverse('beneficiaries:status', args=[self.b.pk]), {'status': 'active'})
        self.b.refresh_from_db()
        self.assertEqual((self.b.status, self.b.discharge_date), ('active', None))
        res = self.client.post(reverse('beneficiaries:status', args=[self.b.pk]), {'status': 'x'}, follow=True)
        self.assertContains(res, '在籍状況を選んでください')

    def test_confirm_and_delete(self):
        from support_plans.models import SupportPlan
        from therapy.models import TherapyRecord
        self.client.login(username='adm', password='pw12345678')
        res = self.client.get(reverse('beneficiaries:detail', args=[self.b.pk]))
        self.assertContains(res, 'この利用者を削除する')
        res = self.client.get(self.url)
        self.assertContains(res, '個別支援計画：1 件')
        self.assertContains(res, '療育記録：1 件')
        self.assertContains(res, '5 年間')
        res = self.client.post(self.url, {'confirm_name': '青木 太郎', 'agree': '1'}, follow=True)
        self.assertContains(res, '削除していません')
        self.assertTrue(Beneficiary.objects.filter(pk=self.b.pk).exists())
        storage, name = self.doc.file.storage, self.doc.file.name
        self.assertTrue(storage.exists(name))
        res = self.client.post(self.url, {'confirm_name': '青木　子', 'agree': '1'}, follow=True)   # 全角スペースでもよい
        self.assertContains(res, '「青木 子」を削除しました')
        self.assertFalse(Beneficiary.objects.filter(pk=self.b.pk).exists())
        self.assertFalse(SupportPlan.objects.exists())
        self.assertFalse(TherapyRecord.objects.exists())
        self.assertFalse(storage.exists(name))


class DevelopmentAssessmentTests(TestCase):
    """5領域アセスメント（評価シート。施設設定で使うにした事業所）"""

    def setUp(self):
        import datetime
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='児童発達支援センター　オウル', layout=Facility.LAYOUT_RYOIKU, use_dev_assessment=True)
        self.user = StaffAccount.objects.create_user('owl', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_STAFF, display_name='大和')
        self.client.login(username='owl', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', last_name_kana='あおき',
                                              first_name_kana='こ', date_of_birth=datetime.date(2019, 4, 1), school_name='南小学校')
        self.detail = reverse('beneficiaries:detail', args=[self.kid.pk])
        self.new_url = reverse('beneficiaries:dev_assessment_create', args=[self.kid.pk])

    def post_data(self, **over):
        data = {'date': '2026-10-07', 'interviewed_with': '保護者（母）',
                'rating_health': '4', 'now_health': '食事は自分で食べる。', 'goal_health': '着替えの手順表を使う',
                'rating_motor': '3', 'rating_cognition': '2', 'now_cognition': '切り替えに時間がかかる',
                'rating_language': '', 'rating_social': '5',
                'strengths': '電車が好き', 'concerns': '大きな音が苦手', 'wishes_child': '友だちと遊びたい',
                'wishes_family': '集団に慣れてほしい', 'summary': '視覚的な手がかりを増やす'}
        data.update(over)
        return data

    def test_disabled_hides_section_and_redirects(self):
        self.f.use_dev_assessment = False
        self.f.save()
        res = self.client.get(self.detail)
        self.assertNotContains(res, 'id="dev-assessments"')
        res = self.client.get(self.new_url)
        self.assertRedirects(res, reverse('facilities:dashboard'))

    def test_create_edit_print_delete(self):
        from .models import DevelopmentAssessment
        res = self.client.get(self.detail)
        self.assertContains(res, 'id="dev-assessments"')
        self.assertContains(res, '新しく作る')
        self.assertNotContains(res, '前回を写して作る')
        res = self.client.get(self.new_url)
        self.assertContains(res, '健康・生活')
        self.assertContains(res, '人間関係・社会性')
        self.assertContains(res, 'まだ保存していません')
        self.assertContains(res, 'data-voice-target="now-health"')
        res = self.client.post(self.new_url, self.post_data())
        a = DevelopmentAssessment.objects.get()
        self.assertRedirects(res, reverse('beneficiaries:dev_assessment_edit', args=[self.kid.pk, a.pk]))
        self.assertEqual((a.assessed_by, a.created_by, a.interviewed_with), (self.user, self.user, '保護者（母）'))
        self.assertEqual(a.rating('health'), 4)
        self.assertIsNone(a.rating('language'))
        self.assertEqual(a.domains['cognition']['now'], '切り替えに時間がかかる')
        self.assertEqual([v for _, v in a.rating_marks], [4, 3, 2, None, 5])
        # 一覧に評価の印が出る。印刷（画面表示）には評価・様子・まとめが出る
        res = self.client.get(self.detail)
        self.assertContains(res, '2026年10月7日')
        self.assertContains(res, '前回を写して作る')
        res = self.client.get(reverse('beneficiaries:dev_assessment_pdf', args=[self.kid.pk, a.pk]) + '?fmt=html')
        self.assertContains(res, '5領域アセスメント')
        self.assertContains(res, '青木 子')
        self.assertContains(res, '（7歳）')          # 2019/4/1 生まれ → 2026/10/7 に 7 歳
        self.assertContains(res, '食事は自分で食べる。')
        self.assertContains(res, '視覚的な手がかりを増やす')
        self.assertContains(res, '<span class="on">4</span>')
        # 直す（評価を変える）。版の確認もする
        edit = reverse('beneficiaries:dev_assessment_edit', args=[self.kid.pk, a.pk])
        res = self.client.get(edit)
        self.assertContains(res, 'value="2026-10-07"')
        self.assertContains(res, 'id="rt-health-4" value="4" checked')
        self.client.post(edit, self.post_data(rating_health='5', summary='更新した'))
        a.refresh_from_db()
        self.assertEqual((a.rating('health'), a.summary), (5, '更新した'))
        res = self.client.post(edit, dict(self.post_data(summary='古い画面から'), version='2000-01-01T00:00:00'), follow=True)
        a.refresh_from_db()
        self.assertEqual(a.summary, '更新した')
        self.assertContains(res, '保存')          # 競合のメッセージ（ほかの人が保存…）が出る
        # 削除
        res = self.client.post(reverse('beneficiaries:dev_assessment_delete', args=[self.kid.pk, a.pk]))
        self.assertRedirects(res, self.detail + '#dev-assessments')
        self.assertFalse(DevelopmentAssessment.objects.exists())

    def test_copy_previous_and_previous_column(self):
        from .models import DevelopmentAssessment
        self.client.post(self.new_url, self.post_data(date='2026-04-01'))
        first = DevelopmentAssessment.objects.get()
        res = self.client.get(self.new_url + '?copy=1')
        self.assertContains(res, '電車が好き')
        self.assertContains(res, 'id="rt-health-4" value="4" checked')
        self.assertContains(res, '前回：<b>4</b>')
        self.client.post(self.new_url, self.post_data(date='2026-10-07', rating_health='5'))
        second = DevelopmentAssessment.objects.exclude(pk=first.pk).get()
        self.assertEqual(second.previous(), first)
        self.assertIsNone(first.previous())
        res = self.client.get(reverse('beneficiaries:dev_assessment_pdf', args=[self.kid.pk, second.pk]) + '?fmt=html')
        self.assertContains(res, '前回：4')
        self.assertContains(res, '2026/4/1')

    def test_bad_date_and_other_facility(self):
        from facilities.models import Facility
        from .models import DevelopmentAssessment
        res = self.client.post(self.new_url, self.post_data(date='bad'), follow=True)
        self.assertContains(res, '実施日を入れてください')
        self.assertFalse(DevelopmentAssessment.objects.exists())
        g = Facility.objects.create(name='ほか', layout=Facility.LAYOUT_RYOIKU, use_dev_assessment=True)
        other = Beneficiary.objects.create(facility=g, last_name='他', first_name='人', date_of_birth=datetime.date(2019, 4, 1))
        res = self.client.get(reverse('beneficiaries:dev_assessment_create', args=[other.pk]))
        self.assertEqual(res.status_code, 404)


class SheetToPlanTests(TestCase):
    """5領域アセスメント → 個別支援計画、レーダーチャート"""

    def setUp(self):
        import datetime
        from accounts.models import StaffAccount
        from facilities.models import Facility
        from .models import DevelopmentAssessment
        self.f = Facility.objects.create(name='児童発達支援センター　オウル', layout=Facility.LAYOUT_RYOIKU, use_dev_assessment=True)
        self.user = StaffAccount.objects.create_user('owl', password='pw12345678', facility=self.f,
                                                     role=StaffAccount.ROLE_STAFF, display_name='大和')
        self.client.login(username='owl', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2019, 4, 1))
        doms = {k: {'rating': i + 1, 'now': f'{label}の様子', 'goal': f'{label}の目標'} for i, (k, label) in enumerate(DevelopmentAssessment.DOMAINS)}
        doms['social']['goal'] = ''
        self.sheet = DevelopmentAssessment.objects.create(beneficiary=self.kid, date=datetime.date(2026, 10, 1), assessed_by=self.user,
                                                          interviewed_with='母', domains=doms, strengths='電車が好き',
                                                          wishes_child='遊びたい', wishes_family='集団に慣れてほしい', summary='見通しを持てるように')

    def test_radar_svg_on_page_and_pdf(self):
        from .sheet_plan import radar_svg
        svg = radar_svg(self.sheet.rows())
        self.assertIn('<polygon', svg)
        self.assertIn('健康・生活', svg)
        res = self.client.get(reverse('beneficiaries:dev_assessment_edit', args=[self.kid.pk, self.sheet.pk]))
        self.assertContains(res, 'class="da-radar"')
        self.assertContains(res, '新しい計画を作って取り込む')
        res = self.client.get(reverse('beneficiaries:dev_assessment_pdf', args=[self.kid.pk, self.sheet.pk]) + '?fmt=html')
        self.assertContains(res, 'class="da-radar"')

    def test_import_creates_plan_fills_assessment_then_goals(self):
        from support_plans.models import PlanGoal, SupportPlan
        url = reverse('beneficiaries:dev_assessment_to_plan', args=[self.kid.pk, self.sheet.pk])
        res = self.client.post(url, follow=True)
        plan = SupportPlan.objects.get(beneficiary=self.kid)
        self.assertEqual((plan.title, plan.current_step, plan.manager), ('第1期 個別支援計画', 1, self.user))
        a = plan.get_step(1)
        self.assertIn('【健康・生活】評価 1（全面的な支援が必要）', a.condition)
        self.assertIn('電車が好き', a.condition)
        self.assertEqual(a.wishes, '本人：遊びたい\n家族：集団に慣れてほしい')
        self.assertEqual((a.interview_date, a.interviewed_with, a.interviewer), (self.sheet.date, '母', self.user))
        self.assertContains(res, '心身の状況・希望する生活・面談日・面談相手 を埋めました')
        self.assertContains(res, '5領域アセスメントから取り込む')        # ステップ画面にも取り込みの箱
        # ステップ2 では目標として入る（同じ文は二重にしない）
        plan.current_step = 2
        plan.save()
        res = self.client.post(url, {'plan': plan.pk}, follow=True)
        goals = list(plan.goals.order_by('order'))
        self.assertEqual(len(goals), 4)                                 # 社会性は目標が空なので入らない
        self.assertEqual(goals[0].content, '健康・生活：健康・生活の目標')
        self.assertEqual(goals[0].goal_type, PlanGoal.TYPE_SHORT)
        self.assertIn('いまの様子（2026/10/1 の5領域アセスメント・評価 1）：健康・生活の様子', goals[0].support_content)
        self.assertEqual(goals[0].target_date, self.sheet.date + datetime.timedelta(days=182))
        self.assertEqual(plan.get_step(2).policy, '見通しを持てるように')
        self.assertContains(res, '短期目標を 4 件足しました')
        res = self.client.post(url, {'plan': plan.pk}, follow=True)
        self.assertEqual(plan.goals.count(), 4)
        self.assertContains(res, '同じ文の目標 4 件はそのままです')
        # ステップ3 以降は取り込めない
        plan.current_step = 3
        plan.save()
        res = self.client.post(url, {'plan': plan.pk}, follow=True)
        self.assertContains(res, 'ステップ3（担当者会議）以降に進んでいる')
        self.assertEqual(plan.goals.count(), 4)


class DevelopmentTestTests(TestCase):
    """発達検査の結果（入力・推移・AI の下書き）"""

    def setUp(self):
        import datetime
        from accounts.models import StaffAccount
        from facilities.models import Facility
        self.f = Facility.objects.create(name='児童発達支援センター　オウル', layout=Facility.LAYOUT_RYOIKU, use_daily_ops=True)
        self.user = StaffAccount.objects.create_user('owl', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_STAFF)
        self.client.login(username='owl', password='pw12345678')
        self.kid = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2021, 4, 1))

    def test_create_edit_chart(self):
        import datetime
        from .models import DevelopmentTest
        url = reverse('beneficiaries:dev_test_create', args=[self.kid.pk])
        res = self.client.get(url)
        self.assertContains(res, '姿勢・運動（P-M）')        # 新版K式の標準の領域
        self.assertContains(res, 'の書類から読み取る')
        res = self.client.post(url, {'test': 'kshiki', 'test_other': '新版K式2020', 'date': '2026-04-01', 'examiner': 'クリニック',
                                     'ca_y': '5', 'ca_m': '0', 'oa_y': '3', 'oa_m': '9', 'overall_quotient': '75',
                                     'r0_label': '姿勢・運動（P-M）', 'r0_y': '4', 'r0_m': '0', 'r0_q': '80',
                                     'r1_label': '言語・社会（L-S）', 'r1_y': '3', 'r1_m': '6', 'r1_q': '70', 'r2_label': '', 'note': '所見'})
        t = DevelopmentTest.objects.get()
        self.assertRedirects(res, reverse('beneficiaries:detail', args=[self.kid.pk]) + '#dev-tests')
        self.assertEqual((t.test_name, t.ca_months, t.overall_age_months, t.overall_quotient, t.ca_label),
                         ('新版K式発達検査', 60, 45, 75, '5歳0か月'))
        self.assertEqual(t.results, [{'label': '姿勢・運動（P-M）', 'age_months': 48, 'quotient': 80},
                                     {'label': '言語・社会（L-S）', 'age_months': 42, 'quotient': 70}])
        DevelopmentTest.objects.create(beneficiary=self.kid, test='kshiki', date=datetime.date(2026, 10, 1), overall_quotient=82,
                                       results=[{'label': '言語・社会（L-S）', 'age_months': 48, 'quotient': 78}])
        res = self.client.get(reverse('beneficiaries:detail', args=[self.kid.pk]))
        self.assertContains(res, 'id="dev-tests"')
        self.assertContains(res, '発達検査の指数の推移')
        self.assertContains(res, '<polyline')
        self.assertContains(res, '全体 3歳9か月・75')
        # 直す・消す
        edit = reverse('beneficiaries:dev_test_edit', args=[self.kid.pk, t.pk])
        res = self.client.get(edit)
        self.assertContains(res, 'value="新版K式2020"')
        self.client.post(edit, {'test': 'kshiki', 'date': '2026-04-02', 'overall_quotient': '76', 'r0_label': 'A', 'r0_q': '1'})
        t.refresh_from_db()
        self.assertEqual((t.date, t.overall_quotient, t.ca_months, t.results), (datetime.date(2026, 4, 2), 76, None, [{'label': 'A', 'age_months': None, 'quotient': 1}]))
        self.client.post(reverse('beneficiaries:dev_test_delete', args=[self.kid.pk, t.pk]))
        self.assertEqual(DevelopmentTest.objects.count(), 1)

    def test_ai_read_prefills_form(self):
        from unittest import mock
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import BeneficiaryDocument
        doc = BeneficiaryDocument.objects.create(beneficiary=self.kid, file=SimpleUploadedFile('kekka.pdf', b'%PDF-1.4'), file_name='kekka.pdf')
        data = {'test': 'enjoji', 'test_other': '', 'date': '2026-09-01', 'examiner': '発達センター', 'ca_months': 65,
                'overall_age_months': None, 'overall_quotient': None, 'note': '所見の要点', 'test_name': '遠城寺式',
                'results': [{'label': '移動運動', 'age_months': 50, 'quotient': None}]}
        with mock.patch('beneficiaries.dev_tests.read_test', return_value=data):
            res = self.client.post(reverse('beneficiaries:dev_test_read', args=[self.kid.pk]), {'source': f'doc:{doc.pk}'}, follow=True)
        self.assertContains(res, 'AI が「書類：kekka.pdf」から読み取った下書きです')
        self.assertContains(res, 'value="2026-09-01"')
        self.assertContains(res, 'value="移動運動"')
        self.assertContains(res, 'value="発達センター"')

    def test_normalize(self):
        from .dev_tests import normalize
        d = normalize({'test': 'x', 'test_name': '独自の検査', 'date': 'bad', 'examiner': '', 'ca_months': -1, 'overall_age_months': 30,
                       'overall_quotient': -1, 'results': [{'label': ' 運動 ', 'age_months': -1, 'quotient': 90}, {'label': '', 'age_months': 1, 'quotient': 1}], 'note': ''})
        self.assertEqual((d['test'], d['test_other'], d['date'], d['ca_months'], d['overall_age_months'], d['overall_quotient']),
                         ('other', '独自の検査', '', None, 30, None))
        self.assertEqual(d['results'], [{'label': '運動', 'age_months': None, 'quotient': 90}])


class BeneficiaryListBulkTests(TestCase):
    """利用者一覧から：1人ずつの在籍状況・削除への導線、まとめて在籍状況を変える・まとめて削除（管理者だけ・確認あり）"""

    def setUp(self):
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU, use_therapy_record=True)
        self.admin = StaffAccount.objects.create_user('adm', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.staff = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f)
        mk = lambda last, kana: Beneficiary.objects.create(facility=self.f, last_name=last, first_name='子', last_name_kana=kana,
                                                            date_of_birth=datetime.date(2019, 4, 1))
        self.a, self.b, self.c = mk('青木', 'あおき'), mk('井上', 'いのうえ'), mk('上田', 'うえだ')
        other = Facility.objects.create(name='ほか', layout=Facility.LAYOUT_RYOIKU)
        self.x = Beneficiary.objects.create(facility=other, last_name='他', first_name='子', date_of_birth=datetime.date(2019, 4, 1))
        self.url = reverse('beneficiaries:bulk')
        self.lst = reverse('beneficiaries:list')

    def test_row_menu_and_bulk_status(self):
        self.client.login(username='st', password='pw12345678')
        res = self.client.get(self.lst)
        self.assertContains(res, 'name="ids" value="%d"' % self.a.pk)
        self.assertContains(res, 'formaction="%s"' % reverse('beneficiaries:status', args=[self.a.pk]))
        self.assertNotContains(res, 'value="delete"')               # 削除は管理者だけ
        # 1人ずつ（一覧から）：一覧に戻る
        res = self.client.post(reverse('beneficiaries:status', args=[self.a.pk]), {'status': 'inactive', 'next': self.lst})
        self.assertRedirects(res, self.lst)
        self.a.refresh_from_db()
        self.assertEqual(self.a.status, 'inactive')
        # まとめて卒業（別の事業所の人は混ぜても変わらない）
        res = self.client.post(self.url, {'action': 'graduated', 'ids': [self.b.pk, self.c.pk, self.x.pk], 'next': self.lst}, follow=True)
        self.assertContains(res, '2 名を「卒業」にしました')
        self.assertEqual(set(Beneficiary.objects.filter(status='graduated').values_list('pk', flat=True)), {self.b.pk, self.c.pk})
        self.b.refresh_from_db()
        self.assertEqual(self.b.discharge_date, datetime.date.today())
        res = self.client.post(self.url, {'action': 'active', 'ids': [self.b.pk]}, follow=True)
        self.b.refresh_from_db()
        self.assertEqual((self.b.status, self.b.discharge_date), ('active', None))
        res = self.client.post(self.url, {'action': 'graduated'}, follow=True)
        self.assertContains(res, '印を付けてから')
        # スタッフはまとめて削除できない
        res = self.client.post(self.url, {'action': 'delete', 'ids': [self.a.pk]}, follow=True)
        self.assertContains(res, '管理者だけ')
        self.assertTrue(Beneficiary.objects.filter(pk=self.a.pk).exists())

    def test_bulk_delete_with_confirmation(self):
        from therapy.models import TherapyRecord
        TherapyRecord.objects.create(facility=self.f, beneficiary=self.a, date=datetime.date(2026, 9, 1))
        self.client.login(username='adm', password='pw12345678')
        self.assertContains(self.client.get(self.lst), 'value="delete"')
        res = self.client.post(self.url, {'action': 'delete', 'ids': [self.a.pk, self.b.pk, self.x.pk]})
        self.assertContains(res, '2 名を削除する')
        self.assertContains(res, '療育記録 1 件')
        self.assertContains(res, '在籍中の人が 2 名')
        self.assertNotContains(res, '他 子')
        res = self.client.post(self.url, {'action': 'delete', 'confirm': '1', 'ids': [self.a.pk, self.b.pk], 'confirm_text': 'さくじょ', 'agree': '1'}, follow=True)
        self.assertContains(res, '削除していません')
        self.assertEqual(Beneficiary.objects.filter(facility=self.f).count(), 3)
        res = self.client.post(self.url, {'action': 'delete', 'confirm': '1', 'ids': [self.a.pk, self.b.pk, self.x.pk],
                                          'confirm_text': '削除', 'agree': '1'}, follow=True)
        self.assertContains(res, '2 名を削除しました')
        self.assertEqual(list(Beneficiary.objects.filter(facility=self.f)), [self.c])
        self.assertTrue(Beneficiary.objects.filter(pk=self.x.pk).exists())
        self.assertFalse(TherapyRecord.objects.exists())


class RecordDigestTests(TestCase):
    """面談資料：期間の日誌・療育記録を集め、数字は計算、文章は AI（記録にあることだけ）。直して保存・印刷"""

    def setUp(self):
        from unittest import mock  # noqa: F401
        self.f = Facility.objects.create(name='児童発達支援センター　オウル', layout=Facility.LAYOUT_RYOIKU, use_therapy_record=True)
        self.user = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f)
        self.client.login(username='st', password='pw12345678')
        self.b = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2019, 4, 1))
        from therapy.models import TherapyRecord
        today = datetime.date.today()
        TherapyRecord.objects.create(facility=self.f, beneficiary=self.b, date=today - datetime.timedelta(days=10),
                                     activities=['ブロック'], body='「できた」と言って最後まで積んだ。')
        TherapyRecord.objects.create(facility=self.f, beneficiary=self.b, date=today - datetime.timedelta(days=200),
                                     activities=['古い'], body='半年より前の記録。')
        self.url = reverse('beneficiaries:digest', args=[self.b.pk])

    def test_collect_and_prompt(self):
        from . import digest
        start, end = digest.period_range('3m')
        data = digest.collect(self.b, start, end)
        self.assertEqual((data['stats']['therapy'], data['stats']['journals']), (1, 0))
        content = digest.build_content(self.b, start, end, data)
        self.assertIn('「できた」と言って最後まで積んだ', content)
        self.assertNotIn('半年より前', content)
        self.assertIn('記録にあることだけ', digest.SYSTEM_PROMPT)
        s6, _ = digest.period_range('6m')
        self.assertLess(s6, start)
        cs, ce = digest.period_range('custom', date_from=datetime.date(2026, 4, 1), date_to=datetime.date(2026, 9, 30))
        self.assertEqual((cs, ce), (datetime.date(2026, 4, 1), datetime.date(2026, 9, 30)))

    def test_create_edit_print(self):
        from unittest import mock
        self.assertContains(self.client.get(reverse('beneficiaries:detail', args=[self.b.pk])), '面談資料（記録のまとめ）')
        res = self.client.get(self.url)
        self.assertContains(res, '療育記録 1 件')
        with mock.patch('beneficiaries.digest.ask_ai', return_value=('## この期間の様子（全体）\nブロックを最後まで積んだ。', None)) as ask:
            res = self.client.post(self.url, {'period': '3m', 'purpose': '10月の面談'}, follow=True)
        self.assertContains(res, '面談資料のまとめを作りました')
        self.assertContains(res, '【この期間の様子（全体）】')
        self.assertIn('「できた」', ask.call_args.args[1])
        item = self.b.digests.get()
        self.assertEqual((item.purpose, item.stats['therapy']), ('10月の面談', 1))
        detail = reverse('beneficiaries:digest_detail', args=[self.b.pk, item.pk])
        self.client.post(detail, {'summary': '直した文', 'purpose': '面談'})
        item.refresh_from_db()
        self.assertEqual(item.summary, '直した文')
        res = self.client.get(detail + '?print=1&fmt=html')
        self.assertContains(res, '直した文')
        # 記録の無い期間はまとめない
        res = self.client.post(self.url, {'period': 'custom', 'date_from': '2020-01-01', 'date_to': '2020-02-01'}, follow=True)
        self.assertContains(res, '日誌・療育記録がありません')
        res = self.client.post(detail, {'action': 'delete'}, follow=True)
        self.assertContains(res, '削除しました')
        self.assertFalse(self.b.digests.exists())

    def test_other_facility_cannot_see(self):
        other = Facility.objects.create(name='ほか')
        StaffAccount.objects.create_user('ot', password='pw12345678', facility=other)
        self.client.login(username='ot', password='pw12345678')
        self.assertEqual(self.client.get(self.url).status_code, 404)


class DobOverwriteAndPartialEditTests(TestCase):
    """生年月日の書き換え（取り込みで同じ子として）と、編集の窓で画面に無い項目を消さないこと"""

    def setUp(self):
        self.f = Facility.objects.create(name='発達支援ルーム　ゆあーず', layout=Facility.LAYOUT_RYOIKU)
        self.user = StaffAccount.objects.create_user('st', password='pw12345678', facility=self.f, role=StaffAccount.ROLE_ADMIN)
        self.client.login(username='st', password='pw12345678')

    def test_edit_window_keeps_hidden_fields(self):
        b = Beneficiary.objects.create(facility=self.f, last_name='青木', first_name='子', date_of_birth=datetime.date(2000, 1, 1),
                                       postal_code='6000000', address='京都市', school_name='南小', grade='e2',
                                       admission_date=datetime.date(2025, 4, 1), has_prior_records=True, weekday_mon=True)
        res = self.client.post(reverse('beneficiaries:update', args=[b.pk]), {
            '_partial': '1', 'last_name': '青木', 'first_name': '子', 'date_of_birth': '2019-04-02',
            'gender': 'male', 'status': 'active', 'weekday_tue': 'on'})
        self.assertEqual(res.status_code, 302)
        b.refresh_from_db()
        self.assertEqual(b.date_of_birth, datetime.date(2019, 4, 2))
        self.assertEqual((b.postal_code, b.address, b.school_name, b.grade, b.admission_date, b.has_prior_records),
                         ('6000000', '京都市', '南小', 'e2', datetime.date(2025, 4, 1), True))
        self.assertEqual((b.weekday_mon, b.weekday_tue), (False, True))      # 窓にある項目は送ったとおり
        self.assertContains(self.client.get(reverse('beneficiaries:detail', args=[b.pk])), 'name="_partial" value="1"')

    def test_import_overwrites_placeholder_dob_when_checked(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        b = Beneficiary.objects.create(facility=self.f, last_name='山田', first_name='太郎', date_of_birth=datetime.date(2000, 1, 1),
                                       notes='生年月日は取り込み時の仮の値（2000-01-01）です。正しい日付に直してください。')
        csv = '姓,名,生年月日,せい（ふりがな）\n山田,太郎,2018/4/2,やまだ\n'
        url = reverse('beneficiaries:import')
        res = self.client.post(url, {'file': SimpleUploadedFile('a.csv', csv.encode('utf-8'))})
        self.assertContains(res, '同じ子として生年月日を書き換える')
        self.assertContains(res, 'name="dob" value="2" form="imp-commit" checked')
        # 印を外すと別の子として作る
        res = self.client.post(url, {'action': 'commit'}, follow=True)
        self.assertEqual(Beneficiary.objects.filter(facility=self.f, last_name='山田').count(), 2)
        Beneficiary.objects.exclude(pk=b.pk).delete()
        self.client.post(url, {'file': SimpleUploadedFile('a.csv', csv.encode('utf-8'))})
        res = self.client.post(url, {'action': 'commit', 'dob': ['2']}, follow=True)
        self.assertContains(res, '台帳にいた人 1 名')
        self.assertEqual(Beneficiary.objects.filter(facility=self.f, last_name='山田').count(), 1)
        b.refresh_from_db()
        self.assertEqual((b.date_of_birth, b.last_name_kana), (datetime.date(2018, 4, 2), 'やまだ'))
        self.assertNotIn('仮の値', b.notes)
